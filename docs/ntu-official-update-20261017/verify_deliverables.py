"""Check the refill workbook against its unchanged historical sheets and source."""
import csv
import hashlib
from itertools import zip_longest
import json
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
old = load_workbook(ROOT / "docs/ntu-official-update-20261015/臺大116-1_282校逐欄回填_20261015.xlsx", read_only=True)
target = OUT / "臺大116-1_逐欄回填_20261017.xlsx"
new = load_workbook(target, read_only=True)
preserved = {}
blank_serialization_changes = 0
for sheet in old:
    title = "10月15日_" + sheet.title.removeprefix("10月15日_")
    count = 0
    for before, after in zip_longest(sheet.values, new[title].values):
        assert before is not None and after is not None, (sheet.title, count)
        normalize = lambda row: tuple("" if value is None else value for value in row)
        assert normalize(before) == normalize(after), (sheet.title, count)
        blank_serialization_changes += sum(a != b for a, b in zip(before, after))
        count += 1
    preserved[sheet.title] = count

with (OUT / "25校十類增量查核_20261017.csv").open(encoding="utf-8-sig") as f:
    rows = list(csv.DictReader(f))
assert len(rows) == 250 and len({r["臺大ID"] for r in rows}) == 25
assert all(r["回填答覆"] == r["核實來源"] == r["核實適用範圍"] == r["原文或頁碼"] == "" for r in rows)
source = ROOT / "sources/ntu/2026-10-17-official-update.json"
assert (OUT / "結構化候選資料.json").read_bytes() == source.read_bytes()
assert new["10月17日_289校正式列表"].max_row == 290
assert new["10月17日_原件基準"].max_row == 38
result = {"workbook_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
          "historical_sheets_all_cell_values_preserved": preserved,
          "blank_string_to_empty_cell_serialization_changes": blank_serialization_changes,
          "incremental_rows": 250, "incremental_schools": 25, "current_formal_list_rows": 289,
          "local_successful_originals": 37, "refill_answer_cells_blank": True,
          "candidate_json_equals_source_snapshot": True}
(OUT / "artifact-validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(result, ensure_ascii=False))
