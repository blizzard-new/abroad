"""Export the scoped review; requires openpyxl only as an authoring dependency."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
PREVIOUS = ROOT / "docs/ntu-official-update-20261015"
SOURCE = ROOT / "sources/ntu/2026-10-17-official-update.json"
data = json.loads(SOURCE.read_text())
refs = {r["portal_id"]: r for r in data["receiving_references"]}
fields = [
    ("quota", "臺大本期甄選名額", None), ("eligibility", "接待校申請資格", "申請資格"),
    ("language", "接待校語言要求", "申請資格"), ("documents", "最新年度接待校文件", None),
    ("deadlines", "接待校提名與申請期限", None), ("calendar", "接待校學期與行事曆", "學校年曆"),
    ("housing", "住宿資訊", "住宿資訊"), ("cost", "費用與生活預算", "生活費與保險資訊"),
    ("courses", "課程與選課限制", "注意事項"), ("insurance_visa", "保險與簽證", "生活費與保險資訊"),
]
headers = ["臺大ID", "學校", "國家", "欄位", "本輪狀態", "待確認警語", "核實內容與限制",
           "原文短引", "頁碼或段落", "適用年／學期", "官方URL", "查核日期", "原件SHA256",
           "正式名額", "全期已核實", "回填答覆", "核實來源", "核實適用範圍", "原文或頁碼"]
review = []
for school in data["schools"]:
    sections = {s["heading"]: s for s in school["sections"]}
    ref = refs.get(school["portal_id"])
    for key, label, section in fields:
        facts = [f for f in ref["facts"] if f["field"] == key] if ref else []
        url, checked, sha = school["url"], school["checked_at"], school["sha256"]
        status = "not_independently_checked_this_batch"
        content = "本輪未獨立補核此欄，沿用10月15日逐欄查核結論；不能解讀為官方未公告。"
        excerpt, location, period = "", "", "未推定116-1"
        if section in sections:
            content = school["note_zh"] + "\n\n臺大內頁參考原文：\n" + sections[section]["text"]
            status = "NTU_portal_reference_pending" if school["pending_warning_present"] else "NTU_portal_reference_field_year_unstated"
            location = "臺大詳細頁／" + section
            if key == "calendar" and school["calendar_academic_year"]:
                period = school["calendar_academic_year"]
                status = "explicit_field_period_reference_only"
        if facts:
            content += "\n\n接待校要項：\n" + "\n".join(f["text_zh"] for f in facts)
            excerpt = " | ".join(f["source_excerpt"] for f in facts)
            location += " | " + "; ".join("PDF第%s頁" % f["page"] if f["page"] else f["section"] for f in facts)
            period = " | ".join(sorted({f["academic_year"] for f in facts}))
            url += "\n" + ref["source_url"]
            checked += "\n" + ref["checked_at"]
            sha += "\n" + ref["source_sha256"]
            status = "receiving_reference_in_own_period"
            if school["portal_id"] == "1283" and key == "cost": status = "unresolved_source_conflict"
        if key == "quota":
            content = "116-1／2027-2028正式列表此校名額空白。詳細頁人數只作參考，不代替臺大甄選名額。"
            status, location, period = "formal_selection_list_blank", "列表：名額／人數上限", "116-1／2027-2028"
            url, checked, sha = (data["list_receipt"][k] for k in ("url", "checked_at", "sha256"))
        review.append([school["portal_id"], school["school_name"], school["country"], label, status,
                       "有" if school["pending_warning_present"] else "無；仍逐欄核對年度",
                       content, excerpt, location, period, url, checked, sha, "空白", "否", "", "", "", ""])

with (OUT / "25校十類增量查核_20261017.csv").open("w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f, lineterminator="\n"); writer.writerow(headers); writer.writerows(review)

with (PREVIOUS / "原282校十類查核_20261015.csv").open(encoding="utf-8-sig") as f:
    old = list(csv.DictReader(f))
old_ids = {r["臺大ID"] for r in old}
current_ids = {r["portal_id"] for r in data["list_receipt"]["rows"]}
coverage = {"legacy_282_schools": len(old_ids), "legacy_rows": len(old),
            "legacy_ids_still_in_current_list": len(old_ids & current_ids),
            "legacy_ids_absent": sorted(old_ids - current_ids),
            "current_roster_extra_ids": sorted(current_ids - old_ids),
            "updated_25_overlap_legacy": len(old_ids & {s["portal_id"] for s in data["schools"]}),
            "legacy_rows_reverified_as_oct17": False, "new_scoped_rows": len(review)}
(OUT / "覆蓋對照.json").write_text(json.dumps(coverage, ensure_ascii=False, indent=2) + "\n")

book = load_workbook(PREVIOUS / "臺大116-1_282校逐欄回填_20261015.xlsx")
for sheet in book:
    sheet.title = "10月15日_" + sheet.title.removeprefix("10月15日_")

def table(name, columns, rows, first=False):
    sheet = book.create_sheet(name, 0 if first else None)
    sheet.append(columns)
    for row in rows: sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for c in sheet[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="184B5A")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i in range(1, len(columns) + 1): sheet.column_dimensions[get_column_letter(i)].width = 26 if i < 7 else 55
    return sheet

table("10月17日_25校十類", headers, review, first=True)
table("10月17日_289校正式列表", ["臺大ID", "學校", "國家", "正式名額原文", "人數上限原文", "查核日期", "來源"],
      [[r["portal_id"], r["name"], r["country"], r["quota"] or "空白", r["headcount"] or "空白",
        data["list_receipt"]["checked_at"], data["list_receipt"]["url"]] for r in data["list_receipt"]["rows"]])
table("10月17日_原件基準", ["來源ID", "來源角色", "官方URL", "查核日期", "SHA256", "原件路徑"],
      [[k, r["role"], r["url"], r["checked_at"], r["sha256"], r["local_file"]] for k, r in data["source_manifest"].items()])
table("10月17日_6個未知原件", ["學校", "原始URL", "父任務查核結果", "本輪重試", "原件恢復"],
      [[a["school"], a["url"], a["parent_reported_status"], "否；尊重網站限制", "否"] for a in data["unknown_original_attachments"]])
table("10月17日_活動", ["學校", "活動", "開始", "結束", "時區", "地點", "官方首頁", "報名表", "截止日", "性質"],
      [[data["event"][k] for k in ("school_name", "title", "start", "end", "timezone", "location", "homepage_url", "registration_url")]
       + ["表單未見；不補造", "資訊活動；不是甄選名額、規則或申請期限"]])
table("10月17日_使用說明", ["項目", "說明"], [
    ["範圍", "本日25校×10類250列；舊282校×10類2820列原值保留於10月15日工作表，沒有全數改標今天查核。"],
    ["正式列表", "289校：20正名額、262空白、7零名額；43學年名額／最多81人次；本輪未改名額。"],
    ["年份", "接收文件依自己年度；UdeM保險仍2026/27；UVic與Uniandes仍2026/27。四英國頁無警語也不表示每欄116-1核實。"],
    ["差額與覆蓋", json.dumps(coverage, ensure_ascii=False)],
    ["待補", "空白回填欄供後續官方證據補核；未找到、未核實、明示待更新及原件無法讀取分開處理。"],
    ["候選與部署", "此檔供PR審查；沒有形成真人核准、發布、部署或申請提醒。"],
], first=True)
target = OUT / "臺大116-1_逐欄回填_20261017.xlsx"
book.save(target)
check = load_workbook(target, read_only=True)
assert check["10月17日_25校十類"].max_row == 251
assert len(old) == 2820 and len(old_ids) == 282 and not coverage["legacy_ids_absent"]
shutil.copyfile(SOURCE, OUT / "結構化候選資料.json")
print(json.dumps({"workbook": target.name, "bytes": target.stat().st_size, "sha256": hashlib.sha256(target.read_bytes()).hexdigest(), "coverage": coverage}, ensure_ascii=False))
