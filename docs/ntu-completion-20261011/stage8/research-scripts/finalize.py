from pathlib import Path
from collections import Counter
from urllib.parse import urlsplit
import json,hashlib,datetime,shutil,zipfile,re,posixpath
from openpyxl import load_workbook
R=Path('/workspace/abroad-ntu-completion-20261011');D=R/'docs/ntu-completion-20261011';S=D/'stage8';V=S/'validation';H=Path(__file__).parent
read=lambda p:json.loads(p.read_text());sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();write=lambda p,x:p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
snap=read(R/'sources/ntu/completion-2026-10-11.json');summary=read(S/'review-summary.json');matrix=read(S/'all-school-review.json');progress=read(R/'sources/ntu/priority-review-progress-current-2026-10-11.json');m=read(R/'sources/ntu/reviewed-source-bodies-current-manifest.json');manual=read(R/'sources/ntu/semantic-review-current-2026-10-11.json');ledger=read(R/'sources/ntu/full-body-review-current-2026-10-11.json');delta=read(S/'本批新增參考增量.json');run=read(V/'run-receipt.json');count=sum(len(s['observations']) for s in snap['schools']);source_sha=sha(R/'sources/ntu/completion-2026-10-11.json');assert count==700+len(delta)==summary['new_observations']==run['observations'];assert len(matrix)==4032
assert source_sha==sha(V/'tested-source-snapshot.json')==run['source_sha256'];assert run['source_code_tests_and_materialized_bytes_unchanged_during_tests'];assert all(sha(R/p)==s for p,s in run['test_files_sha256'].items());assert all(sha(R/p)==s for p,s in run['code_files_sha256'].items());assert all(sha(R/p)==s for p,s in run['materialized_files_sha256'].items());assert all(v['returncode']==0 for v in run['jobs'].values());assert all(not o['selection_rule_confirmed'] for s in snap['schools'] for o in s['observations']);assert progress['fixed_priority_field_denominator']==501 and progress['fixed_priority_school_denominator']==117 and progress['completed_original_cached_body_sets']==501
root_checked=set(read(H/'root-reviewed-observation-ids.json'));assert {o['observation_id'] for o in delta}<=root_checked
results={}
for name in run['jobs']:
 text=(V/(name+'.log')).read_text();last=next((line for line in reversed(text.splitlines()) if re.search(r'\d+ passed',line)),None);assert last and 'failed' not in last;results[name]={'final_line':last,'source_sha256':source_sha,'log':name+'.log','command':run['jobs'][name]['command']}
book=load_workbook(S/'288校逐欄查核_可回填.xlsx',read_only=True);sheets={s.title:s.max_row-1 for s in book};assert len(sheets)==7 and sheets['288校逐欄回填']==4032 and sheets['固定501項正文進度']==501 and sheets['逐源全文閱讀紀錄']==len(ledger) and sheets['來源原文索引']==len(snap['public_source_reviews'])+1;assert not any(c.data_type=='f' for s in book for row in s for c in row);book.close()
sheet_table='\n'.join('| '+title+' | '+str(rows)+' |' for title,rows in sheets.items())
archive=read(V/'repo-archive-verifier.log');assert archive['current_bodies']==m['source_count'] and archive['reviewed_citations']==len(snap['public_source_reviews']) and archive['all_quote_pages_match'];old=read(H/'prior-554-manifest.json');assert all(e in m['entries'] for e in old['entries']) and all(p in m['archive_files'] for p in old['archive_files']);write(S/'archive-verification.json',{**archive,'prior554_versions_and19_parts_immutable':True,'source_snapshot_sha256':source_sha,'incremental_observation_quote_raw_hash_and_pdf_page_checks':len(delta),'incremental_sources_checked_before_import':True})
write(V/'summary.json',{'tested_source_sha256':source_sha,'summary_adapter_sha256':run['summary_adapter_sha256'],'observations':count,'tests':results,'all_references_text_scope_and_URL_visible_in_first10_API_summary_points':True,'api_visibility_test':'test_all_288_schools_accept_offline_reference_assets_and_guarded_workflows','excel':{'sheet_rows':sheets,'formula_cells':0,'GUI_visual_workbook_review_performed':False},'scope':read(V/'scope-and-rebuild.json'),'archive':archive,'pr8_intersection':read(H/'pr8-intersection.json'),'isolated_harness_correction':read(V/'isolated-test-harness-fix.json'),'superseded_failed_runs':run.get('superseded_failed_runs',[]),'life_count_fixture_refresh':read(V/'life-count-fixture-refresh.json'),'original_ten_field_fully_verified_rows':0,'merged_or_deployed':False})
pending={r['portal_id'] for r in matrix if r['field']=='quota' and not r['target_requirement_confirmed']};tenfields={'quota','eligibility','language','documents','deadlines','calendar','housing','cost','courses','insurance_visa'};ten=[r for r in matrix if r['portal_id'] in pending and r['field'] in tenfields];assert len(ten)==2820 and not any(r['target_requirement_confirmed'] for r in ten)
coverage={'school_denominator':282,'field_denominator':10,'rows':len(ten),'fully_verified_rows':0,'old_workbook_retrieved':False,'old_exact_membership_comparison_available':False,'publication_statuses':dict(Counter(r['publication_status'] for r in ten)),'readable_unread_rows':sum(bool(r['unread_candidate_body_count']) for r in ten),'exclusive_external_blocker_rows':sum(r['external_blocker_exclusive'] for r in ten)};write(S/'282-school-ten-field-coverage.json',coverage)
# Same original46 membership and same non-NTU domain criterion, with hosted original host PDFs distinguished.
by={s['portal_id']:s for s in snap['schools']};materialized_life={x['school_name']:x for line in (R/'src/exchange_info/official_data/ntu/student_life.jsonl').read_text().splitlines() if line.strip() for x in [json.loads(line)]};origin=read(D/'stage7/46-life-gap-source-origin-comparison.json')
for x in origin:
 school=by[x['portal_id']];life=materialized_life[school['school_name']]['evidence'];urls=list(dict.fromkeys(l['url'] for e in life for l in e['sources']));direct=[u for u in urls if 'ntu.edu.tw' not in urlsplit(u).netloc];was=x['independent_host_gap_remains'];x['stage7_reported_independent_host_gap_remains']=was;x['stage7_reported_classification_corrected']=x['portal_id']=='819';x.update(current_sources=urls,current_scope=' | '.join(dict.fromkeys(e.get('scope','') for e in life)),any_scoped_official_reference_now=bool(life),independent_host_gap_remains=not bool(direct),stage8_new_independent_host_source=was and bool(direct),independent_origin_definition='same original criterion: official life evidence hosted outside NTU domains; hosted original host PDF listed separately',ntu_hosted_original_host_document_present=any('studyabroad.ntu.edu.tw' in u and '.pdf' in u.lower() for u in urls),all_life_facets_or_target_cycle_verified=False)
 if direct:x['origin_class']='independent_host_domain_scoped_reference'
 elif urls:x['origin_class']='NTU_hosted_reference_only'
 else:x['origin_class']='no_scoped_life_text'
write(S/'46-life-gap-source-origin-comparison.json',origin);write(S/'remaining-independent-host-life-research.json',[x for x in origin if x['independent_host_gap_remains']]);life_gap=sum(not x['evidence'] for x in materialized_life.values());life_paragraphs=sum(len(x['evidence']) for x in materialized_life.values());independent_gap=sum(x['independent_host_gap_remains'] for x in origin)
base=read(D/'stage7/daily-baseline-2026-10-11.json');base.update(stage='stage8_original501_priority_sets_read_and_references_materialized',research_checkpoint_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),source_snapshot_sha256=source_sha,reviewed_observations=count,runtime_validated_observations=count,incremental_imported_observations=len(delta),incremental_unimported_observations=0,reviewed_source_quotes=len(snap['public_source_reviews']),archived_body_count=m['source_count'],manual_semantic_decisions=len(manual),unread_public_body_rows=summary['readable_unread_rows'],fixed_priority_progress={k:v for k,v in progress.items() if k!='fields'},new_host_retrieval_window_utc=['2026-10-11T16:28:43.599530+00:00'],timestamp_note='新增香港語言PDF實際訪問UTC16:28；其餘本批主要核讀既有原件。臺大正式3入口實際最後讀取UTC10:42，288詳情08:31–08:37；本文判讀／重建／測試日期不是全網重新查核日期。');write(S/'daily-baseline-2026-10-11.json',base)
shutil.copyfile(R/'sources/ntu/completion-2026-10-11.json',S/'結構化候選資料.json');shutil.copyfile(R/'sources/ntu/reviewed-source-bodies-current-manifest.json',S/'reviewed-source-bodies-manifest.json');shutil.copyfile(R/'sources/ntu/priority-review-progress-current-2026-10-11.json',S/'固定501組最終覆蓋.json')
report=f'''# 第八批：固定501組原候選本文完成；{count}則參考已匯入Draft PR9

本批新增{len(delta)}則（700→{count}），已重建合作校申請參考與目的地學生指南，並以來源SHA **{source_sha}** 完成API、手機／桌面與全套回歸驗證。所有參考保留原文片段、官方URL、頁碼、適用學年／學期／校區／program／申請階段及來源實際查核日期；每則`selection_rule_confirmed=false`。只更新同一Draft PR9，未merge／deploy，不能稱正式平台已上線。

## 本輪閱讀與核對的結案範圍

接續上一批359組，固定**501校欄／117校／890原候選本文SHA**已全部完成原cached本文集合閱讀（501完成、0待讀），沒有用不斷增加候選改分母。root先完整讀5校，獨立分批核讀72校，按原sourceID＋rawSHA核對、去除確切raw+text同版別名。root逐筆覆核本批全部{len(delta)}新增中文事實、原scope與年份，並驗最小引用、原件SHA及每筆PDF指定頁碼。[整合收據](獨立批次整合收據.json)、[原501組逐項覆蓋](固定501組最終覆蓋.json)、Excel逐源閱讀表保留接續證據。

原集合逐校欄源版本完整讀取{progress['distinct_original_school_field_body_review_pairs']}列，另循公開來源{progress['additional_followup_school_field_body_pairs']}列，共{len(ledger)}列；人工語意{len(manual)}列、引用來源{len(snap['public_source_reviews'])}件。完整閱讀是本批已取得優先候選的適用性查核，**不是當期所有子項核實，也不是所有未取公開連結或282校所有資訊已查完**。錯校、赴外、學位招生、僱員／研究員、特定院級與不符入學期的候選均拒套一般incoming交換。

全288×14欄仍4032列；當日282待定校×原10類仍2820列，完整當期核實仍**0列／0校**。新年度單項、舊年、有條件免證明、版本差異、官方明示某部分待發布、已讀本文未找到、提名／錄取後個人入口與訪問受限各自分類；未找到不等於確實未公告。全矩陣仍有{summary['readable_unread_rows']}個欄位帶可讀候選未完成該欄語意核對（原十類{coverage['readable_unread_rows']}），只限固定原5類優先集合的0待讀不能替代其他欄／新公開連結進度。純外部限制{summary['exclusive_external_blocker_rows']}列（原十類{coverage['exclusive_external_blocker_rows']}）；可讀待判讀是工作，不能稱網站阻礙。

## 有用發現與已處理的誤配

- 筑波AY2027–28明列2027/10修課交換：母校2027年2月提名、學生與CoE 3/5截止；研究交換別表不能混用。金澤KUEP Autumn2027明列提名2/5、申請2/26；2027起指定領域須指導教師簽Provisional Acceptance Agreement，email不能替代。金澤JLCP另一方案的2027時程TBD與KUEP已公告分開。
- 岡山已有2027–28交換時程；orientation／classes的TBA保留為該單項。原傳單物理頁March2019與月生活估值不改為2027新版。父端另查11校結果尚未輸入，本批以原來源自主查核，不等待或偽造父端結果。
- KAIST2027_v2已具新期資格、語言量尺、學期與宿舍健康文件；截止列只列月日不自行補錯年份。一般最低學分可由母校例外允許，不能硬化。明治AY2027通則、August2027英文課表待公布與Spring2027 COE分開；root覆核後把只適用2027春的兩筆改列舊批次，青山Spring2027住房也改舊批次，不因含2027字樣就當116-1。
- 蒙大拿2027–28表與Registrar同年學期結束日差一天，保留版本矛盾；蒙特雷檔名2027–28但正文只列2027 terms，僅Fall2027可對本期，不能推2028春。匹茲堡檔名／封面／頁尾年度及校曆scope分開，原費用年度不改標。
- 慶應AY2027課表約3月中、神戶April2027住民資訊約January、SFU Fall2027住房重要日期Spring2027啟用，都是官方明示特定部分待提供；並非「整校未公告」。原2026 DET暫准不當Fall2027可用。
- 渥太華語言效期起算、滑鐵盧冬提名、Maryland春開放日、Illinois三份申請日期／英語要求、Cornell DET、Temple簽證身分與fee適用、蒙特婁截止／保費、UConn費用年份殘留、UBC兩校區日曆／費率等版本差異均保留。母校推薦／提名／GPA量尺、minimum／recommended、財力證明與預算／帳單、初申與錄取後程序分開。
- UNC原PDF欄位文字抽取順序曾造成春／秋日期誤配，root目視原第1頁後依欄位復原：Spring母校9/1、學生9/15；Fall母校2/15、學生3/1，沒有把抽取錯誤當版本矛盾。[目視頁碼收據](evidence/unc-visual-page-review.json)保留原PDF及render SHA。UNC健康頁J-1不能買StudentBlue，partner頁2026又要求全體該保險，這項實際版本差異另保留。
- 原臺大116-1名額仍282空白／6零；UC與CSU依系統分配校區，未拿總部、Irvine或任選campus補正式接待位置、照片、財力額或日曆，host capacity沒有替代臺大名額。

個別引用請用[新增增量JSON](本批新增參考增量.json)、Excel來源原文及回填欄，而不是僅據本段摘要。簽證／保險只記校方有範圍程序，未替臺灣個人作法律裁定；沒有發信、登入或繞過robots／封鎖。

## 封存、生活與資產缺口

{m['source_count']}本文版本／{m['distinct_source_id_count']}sourceID以{len(m['archive_files'])}件ZIP存原件，原554版本／19分件逐項不變，新增{m['source_count']-554}版本。全部raw與archive SHA、來源引用及PDF頁碼、人工語意來源、歷史59件均經離線封存工具實跑；影像式熊本原文依手抄＋原PDF與12張render SHA核，不假裝有可自動抽取全文。本批主要讀既有本文，唯一新增公開取件是香港大學仍連結的July2021語言PDF；URL2024/08不當文件更新年。

任一官方生活段落現{288-life_gap}校／{life_paragraphs}段、仍缺{life_gap}校；原46校「獨立接待校網域生活來源」仍缺{independent_gap}校，保持原46成員與原non-NTU網域判定，NTU代存接待校原PDF另作註記，沒有改分母消缺口。[46校來源對照](46-life-gap-source-origin-comparison.json)逐校可核。任一生活來源不代表住宿／費用／文化等所有面向或當期全部核實。上一批蘭州手冊原URL已存臺大網域，卻被算作獨立接待校網域，這一處分類有誤：按原non-NTU網域標準，上一批與本批actual runtime均缺39校，原報告38更正為39；原46成員不變且1963段舊生活證據全保留。NTU代存接待校原作者手冊仍另註，沒有否認原作者。[基準分類更正收據](46-life-origin-baseline-correction.json)列兩批逐URL比對。

照片仍缺UC(9)、CSU(12)、特魯瓦(67)；座標仍缺UC、CSU、蒙斯(1302)。特魯瓦官方legal要求再製許可，Commons尚無可同時驗原圖／作者／授權的同校區成品；蒙斯Place du Parc20與其他候選地址／校區不一致或缺可核座標。沿用[具體資產收據](../stage4/evidence/asset-gap-review.json)，沒有發信或冒填。先前核可19照片與16代表座標不變，代表點不當個人指定交換校區。

## 守門、日期與交付

116-1時程與正式共通簡章已公告，報名臺灣2026/11/3 10:00–11/9 16:00。臺大三入口最後實讀UTC10:42、288詳情08:31–08:37；香港新取件UTC16:28，其餘原來源保持自己的checked_at，本文判讀／重建／測試時間不是全網重查時間。[每日基準](daily-baseline-2026-10-11.json)保存比較值，本任務沒有重複建每日排程。

完整測試與隔離固定PR8組合結果：{'；'.join(name+' '+value['final_line'] for name,value in results.items())}。全288 API核全部{count}則正文、scope與URL均在首10顯示項，quota仍None／0且workflow無提醒資格；正式partner資格／臺大名額／申請期限／狀態／身份與main逐項相同，兩步重建byte-idempotent、其他母校catalog與NTU workflow byte一致。PR8固定9927ec…只在隔離樹組合驗證；main59a20f…及PR8未改、未merge。PR8隔離測試最初33通過／1失敗，原因是既有Node部分JavaScript擷取測試漏取PR8新增視覺辨識輔助函式；只修隔離測試擷取依賴，原斷言與PR8產品檔byte不變，重跑34項全通過。失敗log、修前測試與SHA、修後實際命令均保留，PR9測試／產品檔不受該修補影響。完整非E2E首次2184通過／1失敗是舊生活段數快照仍預期1963，已新增161段後為2124；只更新該固定數量斷言，原錯誤頁排除與全部資料驗證不變，來源與產品程式不改。修前log、測試SHA與整套重跑結果均保留，既有瀏覽器source／runtime／E2E測試byte一致而沿用已通過結果。實際未mock的CGU兩正式sheet及77校含Lincoln另留before/during/after收據。

Excel7表、0公式皆核，未宣稱Excel GUI全表目視。

| 工作表 | 資料列 |
| --- | ---: |
{sheet_table}

來源／程式／全部測試及materialized檔SHA、每次命令與結果在[validation/summary.json](validation/summary.json)與run-receipt。全部本批語意結果、原引用／頁碼、封存manifest與基準放回填ZIP；原件分件位於同repo sources/ntu，可按manifest下載驗SHA。

原2026-10-07舊Excel未取得，只以當日臺大288portalIDs重建282待定集合，不能宣稱舊282逐校成員／逐格完全對齊。Library先前401且父端停止重試，沒有Library檔案ID；固定GitHub commit作接續取得方式。環境斷線通知後2026/10/11 UTC16:26:24實際shell/git檢查exit0，未提交資料已byte一致備份，沒有重建或覆寫工作區；見[環境收據](環境斷線檢查_162624.json)。此批完成的是已取得的固定原優先本文集合與新增可用參考，剩餘當期未定、非優先欄位與未取公開連結仍可續查，不能把本輪收集完成當全部缺口解決。
'''
(S/'問題解決報告.md').write_text(report)
for name in ['問題解決報告.md','review-summary.json','288校逐欄查核_可回填.xlsx','結構化候選資料.json','all-school-review.csv','all-school-review.json','daily-baseline-2026-10-11.json','282-school-ten-field-coverage.json','46-life-gap-source-origin-comparison.json','remaining-independent-host-life-research.json']:shutil.copyfile(S/name,D/name)
root_report=re.sub(r'(?<=\]\()([^()]+)(?=\))',lambda match: match.group(0) if '://' in match.group(0) else posixpath.normpath('stage8/'+match.group(0)),report)
(D/'問題解決報告.md').write_text(root_report)
for p in H.glob('*.py'):
 dst=S/'research-scripts'/p.name;dst.parent.mkdir(exist_ok=True);shutil.copyfile(p,dst)
for name in ['cgu-before','cgu-during','cgu-after']:
 target=V/name;target.mkdir(exist_ok=True)
 for p in (H/name).glob('*.json'):shutil.copyfile(p,target/p.name)
for p in (H/'workers').glob('*/*.json'):
 target=S/'independent-source-fitness'/p.parent.name;target.mkdir(parents=True,exist_ok=True);shutil.copyfile(p,target/p.name)
a=S/'臺大288校候選查核_可回填.zip'
with zipfile.ZipFile(a,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in S.rglob('*'):
  if p.is_file() and p!=a and p.suffix!='.zip' and p.name!='tested-source-snapshot.json':z.write(p,p.relative_to(S))
 z.write(R/'sources/ntu/priority-original-cached-texts-2026-10-11.zip','接續原候選文字.zip')
 z.writestr('原始本文說明.txt',f'{len(m["archive_files"])}件原件分件在同repo sources/ntu，以本包manifest逐項驗SHA。歷史PDF目視證據在stage4/5/7 evidence；沒有移除或取代。\n')
shutil.copyfile(a,D/a.name);write(S/'deliverables-manifest.json',[{'path':str(p.relative_to(R)),'size_bytes':p.stat().st_size,'sha256':sha(p)} for p in [S/'288校逐欄查核_可回填.xlsx',S/'問題解決報告.md',S/'結構化候選資料.json',a]])
(D/'README.md').write_text(f'# 臺大288校研究與Draft PR9資料交付\n\n最新為[第八批報告](stage8/問題解決報告.md)、[Excel](stage8/288校逐欄查核_可回填.xlsx)、[JSON](stage8/結構化候選資料.json)及[回填ZIP](stage8/臺大288校候選查核_可回填.zip)。根目錄同名交付保持最新內容，報告附件連結已轉到stage8；前7批均保留歷史狀態。\n\n本批新增{len(delta)}則，700→{count}已匯入且重驗，未merge/deploy。固定501組原候選本文集合全部完成閱讀，當期2820欄完整核實仍0。照片與座標各缺3；任一生活來源缺{life_gap}校、原46校獨立接待校網域來源缺{independent_gap}校。\n\n{m["source_count"]}原件版本由{len(m["archive_files"])}件ZIP按SHA封存。Library401且父端停止重試，無LibraryID；固定GitHubcommit可續取。原件閱讀完成不等所有公開後續連結／年度子項或全平台缺口均已解。\n')
print(json.dumps({'observations':count,'delta':len(delta),'source_sha256':source_sha,'sheets':sheets,'life_gap':life_gap,'independent46_gap':independent_gap,'body_versions':m['source_count'],'parts':len(m['archive_files']),'tests':results},ensure_ascii=False))
