from pathlib import Path
import os,json,hashlib,base64,zipfile,subprocess,datetime
from urllib.request import Request,urlopen
from openpyxl import load_workbook

ROOT=Path('/workspace/abroad-ntu-completion-20261011');HERE=Path(__file__).parent
publication=json.loads((HERE/'publication.json').read_text());COMMIT=publication['remote_commit'];TREE=publication['remote_tree']
headers={'Authorization':'Bearer '+os.environ['GH_TOKEN'],'Accept':'application/vnd.github+json','X-GitHub-Api-Version':'2022-11-28'}
def get(path):
    with urlopen(Request('https://api.github.com/repos/blizzard-new/abroad/'+path,headers=headers),timeout=240) as response:
        return json.load(response)
commit=get('git/commits/'+COMMIT)
assert commit['tree']['sha']==TREE and commit['parents'][0]['sha']==publication['expected_previous_remote_commit']
tree=get('git/trees/'+TREE+'?recursive=1');assert not tree['truncated']
paths={r['path']:r for r in tree['tree'] if r['type']=='blob'}
manifest=json.loads((HERE/'publish-manifest.json').read_text())
assert all(paths[r['path']]['sha']==r['blob_sha'] for r in manifest)
run=json.loads((ROOT/'docs/ntu-completion-20261011/stage9/validation/run-receipt.json').read_text())
bound={**run['test_files_sha256'],**run['code_files_sha256'],**run['materialized_files_sha256'],**run['source_files_sha256']}
bound['sources/ntu/completion-2026-10-11.json']=run['source_sha256']
for path in ['src/exchange_info/static/app.js','src/exchange_info/static/assets/school-images.json']:
    bound[path]=hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
tested_blobs=[]
for path,sha in bound.items():
    p=ROOT/path;assert hashlib.sha256(p.read_bytes()).hexdigest()==sha
    oid=subprocess.check_output(['git','hash-object',str(p)],cwd=ROOT,text=True).strip()
    assert paths[path]['sha']==oid
    tested_blobs.append({'path':path,'sha256':sha,'git_blob_oid':oid})
(HERE/'tested-remote-blob-receipts.json').write_text(json.dumps(tested_blobs,ensure_ascii=False,indent=2)+'\n')
destination=HERE/'remote-downloads';destination.mkdir(exist_ok=True);downloads=[]
for item in json.loads((ROOT/'docs/ntu-completion-20261011/stage9/deliverables-manifest.json').read_text()):
    path=item['path'];blob=get('git/blobs/'+paths[path]['sha']);assert blob['encoding']=='base64'
    data=base64.b64decode(blob['content']);assert data==(ROOT/path).read_bytes()
    assert len(data)==item['size_bytes'] and hashlib.sha256(data).hexdigest()==item['sha256']
    target=destination/Path(path).name;target.write_bytes(data)
    downloads.append({**item,'blob_sha':paths[path]['sha'],'byte_identical_to_tested_local_commit':True})
book=load_workbook(destination/'288校逐欄查核_可回填.xlsx',read_only=True)
rows={s.title:s.max_row-1 for s in book};assert len(rows)==8 and rows['288校逐欄回填']==4032 and rows['429欄最終結論']==429 and rows['固定501項正文進度']==501
book.close()
source=json.loads((destination/'結構化候選資料.json').read_text());assert sum(len(s['observations']) for s in source['schools'])==run['observations']
with zipfile.ZipFile(destination/'臺大288校候選查核_可回填.zip') as z:
    assert z.testzip() is None
    assert z.read('結構化候選資料.json')==(destination/'結構化候選資料.json').read_bytes()
    assert 'validation/fixed-pr8-test-extraction-dependencies.patch' in z.namelist()
body_manifest=json.loads((ROOT/'sources/ntu/reviewed-source-bodies-current-manifest.json').read_text())
assert body_manifest['source_count']>=1228 and len(body_manifest['archive_files'])>=42
pr9=get('pulls/9');pr8=get('pulls/8');main=get('git/ref/heads/main')
assert pr9['head']['sha']==COMMIT and pr9['draft'] and not pr9['merged'] and pr9['state']=='open'
assert pr8['head']['sha']=='9927ec1d8765a9772671bbaa68e7f920995f1258' and pr8['draft'] and not pr8['merged']
assert main['object']['sha']=='59a20f27fcee00fb597abc79e75ec11778317422'
receipt={'checked_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'remote_commit':COMMIT,'remote_tree':TREE,'source_sha256':run['source_sha256'],'all_incremental_blob_shas_match':True,'incremental_blob_count':len(manifest),'all_tested_source_code_tests_and_materialized_blob_shas_match':True,'tested_blob_count':len(tested_blobs),'verified_downloads':downloads,'downloaded_excel_sheet_rows':rows,'downloaded_zip_integrity_verified':True,'archived_body_versions':body_manifest['source_count'],'archive_parts':len(body_manifest['archive_files']),'PR9_draft_unmerged':True,'PR8_unchanged_draft_unmerged':True,'main_unchanged':True,'library_file_ids':[]}
(HERE/'remote-publication-verification.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(receipt,ensure_ascii=False))
