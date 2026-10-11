from pathlib import Path
from collections import Counter
from urllib.parse import urlsplit
from openpyxl import load_workbook
import json,hashlib,shutil,re,posixpath,zipfile,datetime
R=Path('/workspace/abroad-ntu-completion-20261011');D=R/'docs/ntu-completion-20261011';S=D/'stage9';V=S/'validation';H=Path(__file__).parent
read=lambda p:json.loads(p.read_text());write=lambda p,x:p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n');sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
snap=read(R/'sources/ntu/completion-2026-10-11.json');summary=read(S/'review-summary.json');matrix=read(S/'all-school-review.json');closed=read(S/'429欄最終結論.json');delta=read(S/'本批新增參考增量.json');m=read(R/'sources/ntu/reviewed-source-bodies-current-manifest.json');ledger=read(R/'sources/ntu/full-body-review-current-2026-10-11.json');run=read(V/'run-receipt.json');count=sum(len(s['observations']) for s in snap['schools']);source_sha=sha(R/'sources/ntu/completion-2026-10-11.json')
assert count==1373+len(delta)==summary['new_observations']==run['observations'];assert len(matrix)==4032 and len(closed)==429 and summary['fixed429_final_conclusion_rows']==429 and summary['readable_unread_rows']==0
assert source_sha==run['source_sha256']==sha(V/'tested-source-snapshot.json') and run['source_code_tests_and_materialized_bytes_unchanged_during_tests'];assert all(j['returncode']==0 for j in run['jobs'].values())
for key in ['test_files_sha256','code_files_sha256','materialized_files_sha256','source_files_sha256']:assert all(sha(R/path)==digest for path,digest in run[key].items())
checked=set(read(H/'root-reviewed-observation-ids.json'));assert {o['observation_id'] for o in delta}<=checked
old=read(H/'base-1373.json');oldobs={o['observation_id']:o for s in old['schools'] for o in s['observations']};newobs={o['observation_id']:o for s in snap['schools'] for o in s['observations']};assert all(newobs[k]==v for k,v in oldobs.items());assert all(e in next(s for s in snap['schools'] if s['portal_id']==prior['portal_id'])['life_evidence'] for prior in old['schools'] for e in prior['life_evidence'])
results={}
for name,job in run['jobs'].items():
 text=(V/(name+'.log')).read_text();last=next((line for line in reversed(text.splitlines()) if re.search(r'\d+ passed',line)),None);assert last and 'failed' not in last;results[name]={'final_line':last,'log':name+'.log','command':job['command']}
book=load_workbook(S/'288校逐欄查核_可回填.xlsx',read_only=True);sheets={s.title:s.max_row-1 for s in book};assert len(sheets)==8 and sheets['288校逐欄回填']==4032 and sheets['429欄最終結論']==429 and sheets['固定501項正文進度']==501;assert not any(c.data_type=='f' for s in book for row in s for c in row);book.close()
archive=read(V/'repo-archive-verifier.log');assert archive['all_quote_pages_match'];priorarchive=read(H/'base-body-manifest1228.json');assert all(e in m['entries'] for e in priorarchive['entries']) and all(p in m['archive_files'] for p in priorarchive['archive_files'])
write(S/'archive-verification.json',{**archive,'prior1228_versions_and42_parts_immutable':True,'source_snapshot_sha256':source_sha,'new_fetches':0})
readrows=lambda p:[json.loads(line) for line in p.read_text().splitlines() if line.strip()]
life={r['school_name']:r for r in readrows(R/'src/exchange_info/official_data/ntu/student_life.jsonl')};oldlife=readrows(H/'prior-runtime/student_life.jsonl');assert sum(len(r['evidence']) for r in oldlife)==2124
assert all(e in life[s['school_name']]['evidence'] for s in oldlife for e in s['evidence']);life_gap=sum(not s['evidence'] for s in life.values());life_count=sum(len(s['evidence']) for s in life.values())
origin=read(D/'stage8/46-life-gap-source-origin-comparison.json');pids={s['portal_id']:s for s in snap['schools']}
for row in origin:
 s=life[pids[row['portal_id']]['school_name']];urls=list(dict.fromkeys(u['url'] for e in s['evidence'] for u in e['sources']));direct=[u for u in urls if 'ntu.edu.tw' not in urlsplit(u).netloc];before=row['independent_host_gap_remains'];row.update(current_sources=urls,current_scope=' | '.join(dict.fromkeys(e.get('scope','') for e in s['evidence'])),any_scoped_official_reference_now=bool(s['evidence']),independent_host_gap_remains=not bool(direct),stage9_new_independent_host_source=before and bool(direct),all_life_facets_or_target_cycle_verified=False)
independent_gap=sum(r['independent_host_gap_remains'] for r in origin);write(S/'46-life-gap-source-origin-comparison.json',origin);write(S/'remaining-independent-host-life-research.json',[r for r in origin if r['independent_host_gap_remains']]);shutil.copyfile(D/'stage8/46-life-origin-baseline-correction.json',S/'46-life-origin-baseline-correction.json')
pending={r['portal_id'] for r in matrix if r['field']=='quota' and not r['target_requirement_confirmed']};assert len(pending)==282
ten=[r for r in matrix if r['portal_id'] in pending and r['field'] in {'quota','eligibility','language','documents','deadlines','calendar','housing','cost','courses','insurance_visa'}];assert len(ten)==2820 and not any(r['target_requirement_confirmed'] for r in ten)
write(S/'282-school-ten-field-coverage.json',{'school_denominator':282,'field_denominator':10,'rows':2820,'fully_verified_rows':0,'old_workbook_retrieved':False,'old_exact_membership_comparison_available':False,'readable_unread_rows':sum(bool(r['unread_candidate_body_count']) for r in ten),'fixed429_concluded_rows':sum(r['bounded_field_conclusion_complete'] for r in ten)})
baseline=read(D/'stage8/daily-baseline-2026-10-11.json');baseline.update(stage='stage9_fixed429_terminal_field_conclusions',research_checkpoint_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),source_snapshot_sha256=source_sha,reviewed_observations=count,runtime_validated_observations=count,incremental_imported_observations=len(delta),incremental_unimported_observations=0,archived_body_count=m['source_count'],unread_public_body_rows=0,fixed429_final_conclusion_rows=429,new_host_retrieval_window_utc=[],timestamp_note='本批0取件，僅既有快取的欄位判讀；沿用來源實際checked_at，不把判讀/重建/測試時間冒充全網重查時間。');write(S/'daily-baseline-2026-10-11.json',baseline)
necessary=[{'portal_id':r['portal_id'],'school_name':r['school_name'],'field':r['field'],'link':link} for r in closed for link in r.get('necessary_public_links_not_fetched',[])];write(S/'少數必要官方連結_未新增取件.json',necessary)
write(S/'結構化候選資料.json',snap);shutil.copyfile(R/'sources/ntu/reviewed-source-bodies-current-manifest.json',S/'reviewed-source-bodies-manifest.json')
write(V/'summary.json',{'tested_source_sha256':source_sha,'observations':count,'tests':results,'excel':{'sheet_rows':sheets,'formula_cells':0,'GUI_visual_workbook_review_performed':False},'scope':read(V/'scope-and-rebuild.json'),'archive':archive,'pr8_intersection':read(H/'pr8-intersection.json'),'prior1373_observations_and2124_runtime_life_evidence_unchanged':True,'fixed429_all_terminal_conclusions_complete':True,'whole_current_cycle_fields_verified':0,'new_fetches':0,'merged_or_deployed':False})
outcomes=Counter(r['outcome'] for r in closed);outtable='\n'.join('| '+key+' | '+str(value)+' |' for key,value in sorted(outcomes.items()));sheettable='\n'.join('| '+key+' | '+str(value)+' |' for key,value in sheets.items())
report=f'''# 臺大交換資料：既有429欄最終判讀與可取得範圍驗收

本輪固定429欄／121校已全部留下終局結論，每欄所有已取得候選均有source ID、原件SHA、實際取件日與具體適用/拒用理由；**0欄再以同一份已讀正文列待判**。沒有擴張搜尋或新增取件。新增{len(delta)}則可採參考，1373→{count}，已落同Draft PR9的資料及指南；每校仍需確認的學年、校區、身分、費率、個人入口等具體子項一併展示。未merge/deploy。

[429欄結論](429欄最終結論.json)及Excel第8表可逐欄回填；[完整4032列](all-school-review.json)維持全288校×14類，[282校原10類對照](282-school-ten-field-coverage.json)維持2820列。282待定校×原10類的完整116-1／2027–28核實仍0列，終局判讀完成不等當期規定全部已知。

| 結論代碼 | 欄位數 |
| --- | ---: |
{outtable}

`verified_scoped_reference`僅已核原文件範圍參考；`not_applicable_to_exchange_scope`為錯校/方向/方案/欄位；`insufficient_existing_source`為現有來源未給必要子項；`version_conflict`保留原版本差異；`officially_unannounced`只在原文有肯定待公告聲明才使用；`external_information_required`具名提名/錄取/校區分配/個人或官方外部資料需求。未找到不是確實未公告。拒用來源仍封存，未偽稱只看身分/相關頁的本文已全文讀完；[整合收據](獨立批次整合收據.json)記錄這些有限判讀，舊全文紀錄保持。

固定501原優先本文集合在第8批已讀完，501／117校／890 raw SHA原分母與0剩餘保持。本輪429為其後全矩陣的欄位適用性收斂，包含其他欄及同原件不同欄位；兩者不能互相替代。歷史未取seed仍保留在原來源清單作追溯，已結案欄不再被它無限重開。另有4個固定429以外的歷史標記：創價10頁Faculty Guide只核封面/27字抽取，未核其內頁生活文字；釜山documents只有公開資料夾檔名、SNUcost只有宿舍介紹、HSGinsurance只有入口。它們是原件/索引不足的具體限制，舊判讀時間及原文保留，未冒稱全文已讀或當期已核；也不由通用可讀首頁另生992個「正文待讀」工作項。[少數具體必要官方連結](少數必要官方連結_未新增取件.json)共有{len(necessary)}筆欄位連結，不建無界候選。

## 已採與拒用

[本輪增量](本批新增參考增量.json)保存新增原文、PDF物理頁、原學年/學期/校區/方案、原URL及查核日。例：慶北2026正式日曆3/3開課，不能用未標年3/2概述覆写或改標2027；HKU Spring2027屬2026/27；UVic2026/27春季多TBD。慶北研9與舊PDF10學分、web自動申寮與舊PDF另申、保險估價差別全部保留。南卡保險數字明示2024/25而非該2025/26要項年度，未冒充2027價格；一般outgoing頁/教職員協定頁拒作incoming規定。

秋田大學候選混入Akita International University、Mānoa課程候選是Summer2026註冊日曆等均附SID-specific理由，不將錯校/錯欄正文再留待判。所有新增觀察selection_rule_confirmed=False，海外接收名額不代臺大甄選名額，簽證/保險只記校方有範圍資料而不替個人作法律/醫療裁定。

## 正式基準與固定資產限制

臺大116-1共通簡章/時程已公告，報名臺灣2026/11/3 10:00–11/9 16:00。288校學校專屬臺大名額仍282空白、6零；正式資格/名額/期限/狀態/身分與main一致。[每日基準](daily-baseline-2026-10-11.json)沿用三入口實際UTC10:42、288詳情08:31–08:37及個別來源checked_at；本輪判讀日期不是新取件日期。父端負責每日台灣10:00排程，本輪未建重複排程，未發信/登入/繞封鎖。

UC與CSU交換校區未指定，照片與座標不能以任意分校填滿。仍缺照片UC(9)、CSU(12)、特魯瓦(67)，座標UC、CSU、蒙斯(1302)。特魯瓦原圖/作者/授權未齊且官方再製許可未取得；蒙斯Place du Parc20與候選地址/校區不一致或缺可核座標。[資產收據](../stage4/evidence/asset-gap-review.json)正式保留，19核可照片/16代表點未變，不繼續無界查。

生活原2124段全保留，現{288-life_gap}校/{life_count}段，任一生活来源仍缺{life_gap}校。原46校獨立接待校網域生活來源按同成員/同non-NTU口徑仍缺{independent_gap}校。[逐校URL對照](46-life-gap-source-origin-comparison.json)區分NTU代存原校PDF。第7批報告38誤把蘭州(819)臺大代存手冊算獨立海外網域；第7/8批實際runtime都39，本輪沒有刪掉來源或改分母。[更正收據](46-life-origin-baseline-correction.json)保留，不能把單段生活来源當所有面向核實。

## 驗收、封存及接續

來源快照SHA256：`{source_sha}`。全部新增原文/頁碼/原件SHA離線核對，{m['source_count']}原件版本/{len(m['archive_files'])}ZIP分件封存，舊1228版本/42分件逐件不變。拒用本文亦有封存定位；本包含manifest與判讀，原件分件在同repo sources/ntu。[封存驗證](archive-verification.json)。

驗證：{'；'.join(name+' '+v['final_line'] for name,v in results.items())}。全部{count}參考正文/scope/URL保留在288校API首10摘要項，新增429結論與缺項於指南/API可讀；桌面/手機實跑。正式partner欄、工作流/截止提醒、其他母校及共享app.js守門，兩步重建byte-idempotent，見[驗收收據](validation/summary.json)。固定PR8(9927ec…)只在隔離樹组合，坐標修改记录无交集。沿用第8批已公开的隔离Node擷取依賴補丁，产品PR8 bytes未變；若将来合并PR8，需同时应用此测试补丁，不能声称未经补丁的组合测试全通过。

本輪首次隔離focused為36通過/1失敗，查核JSON值相同而key順序不穩，已以pop/reinsert結論欄修重放程式，原idempotence斷言未改；首次瀏覽器17通過/2失敗是新案例漏展開既有官方指南，补上展開操作後保留全部可見性/內容斷言。首次全套430通過/7skip後為修已重現問題而中止、組合瀏覽器亦中止且0完成；未把它們稱完整通過。失敗/中止log、修前程式/測試、原SHA及診斷在[initial-failed-run](validation/initial-failed-run/run-receipt.json)，修補與整輪重跑證據在[corrections](validation/corrections.json)。引用來源JSON及全部原件SHA未變。

其後完整回歸為2189通過/1失敗/13跳過：60個新verification欄在首次建置時尚無source_urls，結論key與重放順序不同；JSON值完全相同。已將結論寫入移至觀察迴圈建立來源與狀態欄之後，原單次重建及歷史重放byte斷言均不改，兩项针对修正驗證通過，再整輪重跑上述四組。該失敗整輪log/輸入SHA、修前程式與runtime、單次建置產物及逐欄key診斷在[replay-build-failed-run](validation/replay-build-failed-run/run-receipt.json)，[修正收據](validation/replay-build-correction.json)確認来源SHA、所有資料值及原斷言未變。

Excel8表/0公式已核，不宣稱全表GUI目視。

| 工作表 | 資料列 |
| --- | ---: |
{sheettable}

原2026-10-07 Excel未取得，已按當日臺大288portal ID重建282待定集合；逐校/逐格與舊Excel精確重疊未知，未要求用戶逐校補。[報告](問題解決報告.md)、[Excel](288校逐欄查核_可回填.xlsx)、[候選JSON](結構化候選資料.json)、[回填ZIP](臺大288校候選查核_可回填.zip)同Draft PR9固定commit可取。Library先前401且父端停止重試，沒有Library檔案ID。這是可取得既有範圍的最終驗收；仍保留具名當期/錄取/個人/權限缺項，未宣稱所有282校資料缺口已解。
'''
(S/'問題解決報告.md').write_text(report)
for worker in ['a1','a2','a3','a4','a5','a6','root']:
 dest=S/'independent-field-conclusions'/worker;dest.mkdir(parents=True,exist_ok=True)
 for name in ['assigned-fields.json','field-review-bundle.json']:shutil.copyfile(H/'workers'/worker/name,dest/name)
 for p in (H/'workers'/worker).glob('*.json'):
  if p.name not in ['assigned-fields.json','field-review-bundle.json']:shutil.copyfile(p,dest/p.name)
 for p in (H/'workers'/worker).rglob('*.png'):
  target=dest/p.relative_to(H/'workers'/worker);target.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(p,target)
for p in H.glob('*.py'):
 dest=S/'research-scripts'/p.name;dest.parent.mkdir(exist_ok=True);shutil.copyfile(p,dest)
shutil.copyfile(H.parent/'stage8/fixed-pr8-test-extraction-dependencies.patch',V/'fixed-pr8-test-extraction-dependencies.patch')
archivepath=S/'臺大288校候選查核_可回填.zip'
with zipfile.ZipFile(archivepath,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in S.rglob('*'):
  if p.is_file() and p!=archivepath and p.suffix!='.zip' and p.name!='tested-source-snapshot.json':z.write(p,p.relative_to(S))
 z.writestr('原始本文說明.txt',f'全部{len(m["archive_files"])}件原件分件在同repo sources/ntu，按manifest下載驗SHA；本輪0新取件。\n')
for name in ['review-summary.json','288校逐欄查核_可回填.xlsx','結構化候選資料.json','all-school-review.csv','all-school-review.json','daily-baseline-2026-10-11.json','282-school-ten-field-coverage.json','46-life-gap-source-origin-comparison.json','remaining-independent-host-life-research.json',archivepath.name]:shutil.copyfile(S/name,D/name)
(D/'問題解決報告.md').write_text(re.sub(r'(?<=\]\()([^()]+)(?=\))',lambda x:x.group(0) if '://' in x.group(0) else posixpath.normpath('stage9/'+x.group(0)),report))
write(S/'deliverables-manifest.json',[{'path':str(p.relative_to(R)),'size_bytes':p.stat().st_size,'sha256':sha(p)} for p in [S/'288校逐欄查核_可回填.xlsx',S/'問題解決報告.md',S/'結構化候選資料.json',archivepath]])
(D/'README.md').write_text(f'# 臺大288校資料與Draft PR9\n\n最新為[第9批429欄最終判讀報告](stage9/問題解決報告.md)、[Excel](stage9/288校逐欄查核_可回填.xlsx)、[候選JSON](stage9/結構化候選資料.json)、[回填ZIP](stage9/臺大288校候選查核_可回填.zip)。根目錄同名交付同步最新，前8批歷史保留。\n\n固定429欄已結論、0現有可讀候選欄待判；新增{len(delta)}參考，現{count}。2820當期欄完整核實仍0；具體未知項保留。照片/座標各缺3，任一生活來源缺{life_gap}，原46獨立網域口徑缺{independent_gap}；蘭州38→39分類更正收據保留。\n\n0新增取件，原{m["source_count"]}本文版本由{len(m["archive_files"])}ZIP按SHA封存。Library401且父端停止重試，無Library ID；固定GitHubcommit可續取。未merge/deploy，未重建每日排程。\n')
print(json.dumps({'observations':count,'delta':len(delta),'source_sha256':source_sha,'conclusions':dict(outcomes),'life_gap':life_gap,'life_count':life_count,'independent46_gap':independent_gap,'sheets':sheets,'necessary_links':len(necessary),'archive_versions':m['source_count'],'archive_parts':len(m['archive_files']),'tests':results},ensure_ascii=False))
