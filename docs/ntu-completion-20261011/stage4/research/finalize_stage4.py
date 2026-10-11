from pathlib import Path
from collections import Counter
import json,hashlib,shutil,zipfile,datetime,re
from openpyxl import load_workbook
R=Path('/workspace/abroad-ntu-completion-20261011');D=R/'docs/ntu-completion-20261011';S=D/'stage4';V=S/'validation';H=Path(__file__).parent
read=lambda p:json.loads(p.read_text())
def write(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
snap=read(R/'sources/ntu/completion-2026-10-11.json');matrix=read(S/'all-school-review.json');summary=read(S/'review-summary.json');manual=read(R/'sources/ntu/semantic-review-current-2026-10-11.json');progress=read(H/'priority-review-progress.json');archive=read(R/'sources/ntu/reviewed-source-bodies-current-manifest.json');ledger=read(H/'full-body-field-review-ledger.json')
assert len(matrix)==4032 and sum(len(x['observations']) for x in snap['schools'])==608
assert progress['completed_original_cached_body_sets']==100 and progress['remaining_original_cached_body_sets']==401
assert sha(V/'tested-source-snapshot.json')==sha(R/'sources/ntu/completion-2026-10-11.json')
by={x['portal_id']:x for x in snap['schools']};current={(x['portal_id'],x['field']):x for x in matrix}
pending={x['portal_id'] for x in matrix if x['field']=='quota' and not x['target_requirement_confirmed']};fields10={'quota','eligibility','language','documents','deadlines','calendar','housing','cost','courses','insurance_visa'};ten=[x for x in matrix if x['portal_id'] in pending and x['field'] in fields10]
assert len(ten)==2820 and all(not x['target_requirement_confirmed'] for x in ten)
coverage={'school_denominator':282,'field_denominator':10,'rows':2820,'fully_verified_rows':0,'readable_unread_rows':sum(bool(x['unread_candidate_body_count']) for x in ten),'exclusive_external_blocker_rows':sum(x['external_blocker_exclusive'] for x in ten),'publication_statuses':dict(Counter(x['publication_status'] for x in ten)),'review_progresses':dict(Counter(x['review_progress'] for x in ten)),'field_counts':{f:{'rows':282,'reference_present':sum(x['reference_present'] for x in ten if x['field']==f),'target_requirement_confirmed':0,'host_target_year_partial_fact_rows':sum(x['host_target_year_fact_present'] for x in ten if x['field']==f)} for f in sorted(fields10)},'old_workbook_retrieved':False,'old_exact_membership_comparison_available':False,'coverage_note':'以2026-10-11官方288校portal ID重建當日282待定校×10欄，不宣稱等於未取得之舊282校逐格名單。'};write(S/'282-school-ten-field-coverage.json',coverage)
unread=[x for x in matrix if x['unread_candidate_body_count']];write(S/'readable-body-work-queue.json',unread)
oldlife=read(D/'stage3/46-life-gap-source-origin-comparison.json')
for x in oldlife:
 if x['independent_host_gap_remains']:assert not any('ntu.edu.tw' not in l['url'] for e in by[x['portal_id']]['life_evidence'] for l in e['sources'])
 x['any_scoped_official_reference_now']=current[x['portal_id'],'life_sources']['reference_present'];x['stage4_new_independent_host_source']=False
assert sum(x['independent_host_gap_remains'] for x in oldlife)==38
write(S/'46-life-gap-source-origin-comparison.json',oldlife);write(S/'remaining-independent-host-life-research.json',[x for x in oldlife if x['independent_host_gap_remains']])
exceptions=[x for x in matrix if x['field'] in {'photos','coordinates','life_sources'} and not x['reference_present']];assert len(exceptions)==22
write(S/'remaining-data-exceptions.json',exceptions)
lifefile=R/'src/exchange_info/official_data/ntu/student_life.jsonl';life=[json.loads(x) for x in lifefile.read_text().splitlines()];assert sum(len(x['evidence']) for x in life)==1963
base=read(D/'stage3/daily-baseline-2026-10-11.json');base.update(stage='stage4',research_checkpoint_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),reviewed_observations=608,reviewed_source_quotes=322,archived_body_count=478,manual_semantic_decisions=len(manual),unread_public_body_rows=len(unread),source_snapshot_sha256=sha(R/'sources/ntu/completion-2026-10-11.json'),fixed_priority_progress={k:v for k,v in progress.items() if k!='fields'},stage4_access_receipts=[read(p) for p in sorted(H.glob('followups*-receipts.json'))],timestamp_note='臺大3個正式入口最後實際讀取UTC10:42；288詳情UTC08:31–08:37。第四批接待校新來源自12:55至13:45各自保留checked_at；整理時間不是重查臺大時間。未建立第二個排程。')
base['coverage'].update(life_excerpts=1963,photos=285,locations=285,schools_with_any_scoped_official_life_reference=272,remaining_independent_gaps_among_prior_46=38)
write(S/'daily-baseline-2026-10-11.json',base)
book=load_workbook(S/'288校逐欄查核_可回填.xlsx',read_only=True,data_only=False);sheets={w.title:w.max_row-1 for w in book};assert sheets=={'288校逐欄回填':4032,'來源原文索引':323,'當日臺大來源':3,'臺大116-1正式共通規則':11,'候選段落語意分類':267,'固定501項正文進度':501,'逐源全文閱讀紀錄':376};assert not any(c.data_type=='f' for w in book for row in w for c in row);book.close()
def test(log,passed,skip=0):
 t=(V/log).read_text();assert 'FAILED' not in t and f'{passed} passed' in t
 if skip:assert f'{skip} skipped' in t
 return {'passed':passed,'failed':0,'skipped':skip,'log':log}
validation={'source_snapshot_sha256':sha(R/'sources/ntu/completion-2026-10-11.json'),'base_main_sha':snap['base_main_sha'],'focused_tests':test('focused-final.log',96,1),'feature_browser_tests':test('browser-final.log',13),'browser_widths':[390,1440],'scope_and_rebuild':read(V/'scope-and-rebuild.json'),'archive_verification':read(S/'archive-verification.json'),'excel':{'sheet_rows':sheets,'formula_cells':0},'prior_full_suite':{'stage':'stage3','passed':2182,'skipped':13,'failed':0,'tested_source_sha256':'6ed2605c4ffb5e43a0e62db0553f186e3a228f6dca439ef0302cd94d7042d636','applies_to_current_stage4_facts':False,'log':'../../stage3/validation/full-regression-final.log'},'current_pr8_combined_focused':test('pr8-visual-locations-content.log',23),'additional_migration_guard':test('pr8-combined-focused.log',7,1),'current_pr8_intersection':read(V/'pr8-intersection.json'),'prior_pr8_browser':{'stage':'stage3','fixed_PR8_head':'9927ec1d8765a9772671bbaa68e7f920995f1258','passed':17,'current_stage4_retest':False,'log':'../../stage3/validation/pr8-combined-browser.log'},'all_current_cycle_fields_verified_schools':0,'original_2820_current_cycle_rows_fully_verified':0,'deploy_performed':False,'library_file_ids':[],'library_blocker':'Prior HTTP401; parent stopped retries','parent_parallel_11_school_results_received':False,'first_stage4_failures_preserved':'focused.log','failure_resolution':'更新第三批固定筆數斷言、同步來源metadata後重建；4筆PDF事實补頁碼並核對6個引用索引頁。未改mock或fixture迴避。'};write(V/'summary.json',validation)
shutil.copyfile(R/'sources/ntu/priority-original-cached-texts-manifest.json',S/'evidence/priority-original-cached-texts-manifest.json');shutil.copyfile(R/'sources/ntu/completion-2026-10-11.json',S/'結構化候選資料.json');shutil.copyfile(R/'sources/ntu/reviewed-source-bodies-current-manifest.json',S/'reviewed-source-bodies-current-manifest.json')
report=f'''# 臺大 116-1／第四批逐來源正文核實（2026-10-11）

照片仍缺 **3校**：加州大學（portal9）、加州州立大學（12）、特魯瓦科技大學（67）。座標仍缺 **3校**：加州大學（9）、加州州立大學（12）、蒙斯大學（1302）。本批沒有採用新資產；同一Draft PR先前三批核可的照片19筆、代表座標16筆維持。

|缺口|既有核對與本批處理|仍待取得|
|---|---|---|
|UC／CSU照片及座標|未取得臺大正式分校分配；系統總部或任選校區不當交換目的校區|正式指定校區，再核對可重用校園照片及同一校區點位|
|特魯瓦科技大學照片|正式legal頁需再製許可；沒有發信。EUt+logo拒絕；CommonsAPI robots讀取不可用，未取得本體、作者與授權皆可驗證的同一校區照片|具清楚授權且完成影像核對的校園照片|
|蒙斯大學座標|比利時官方StudyinBelgium已核Place du Parc20,7000Mons。官方校圖robots讀取不可用；Charleroi、BdDolez31與Warocque17索引點的地址／校區不同或缺当前交叉核實，未混填|能交叉核實同一代表校區的地圖與座標，不能以郵址或城市中心替代|

具體URL、訪問收據、拒用線索保存在 [asset-gap-review.json](evidence/asset-gap-review.json)。HTTP200、地址或索引片段各自不等校區／授權已通過。

## 本批實際增加與固定分母

第三批433則候選增加到**608則（新增175）**，來源原文索引322個接待校來源，加上1個臺大116-1正式簡章。已完整閱讀的逐校逐欄逐來源／版本紀錄376列；同URL有新版本時以SHA分開，原始本文478版本、477個sourceID、16個ZIP分件。第三批345版本與12分件不變。本批新增133版本及4分件；22張PDF核對圖均驗SHA。

固定優先正文分母仍為**501項校×欄／117校／890個原候選不同本文SHA**。已完成原候選全部本文集合**100/501**、涉及30校；仍有**401項**未完成。247列為原候選校欄來源閱讀，129列為另循公開連結補讀，不增加原分母。此完成數只代表原cached本文全部讀過，不代表完整當期規定已核實，也不包括尚未跟到的新公開文件。全部來源是人工語意判讀；掃描頁另渲染目視。熊本12頁PDF文字原抽取33字，普通HTTPS取得45.9MB後已讀所有頁，沒有把掃描難解析當訪問阻礙。

新資料例子（完整scope／URL／原文索引／頁碼／查核時間在候選JSON與工作簿）：

- **御茶水FY2027**：官方英日要項4頁已公開且讀完；2027秋申請4/16，只有研究型須3/18事前登錄。GPA2.5/4是推薦值、低於者可個案，不變硬最低值。月寮65300日圓加入居31680、保證4000，保留舊2025版61900及gas除外的差別。只涵蓋2027秋，不套2028春入學。
- **韓國外大2027秋**：官網Spring2027–Fall2027標題明示秋提名與申請2/1–4/30、授課9/1–12/21；迎新與入寮仍TBC。英文GPA/語言／文件／財力與非OECD認證分階段核實。財力8M/16M KRW是最低證明情境、非一律接待校收費。
- **神奈川2027**：募集要項已連公開Box，robots收據不可用造成正文未取到，不能說未公告。官網明确2027英語科目3月、syllabus4月公布；2026獎學金及未標年費表不改標2027。
- **熊本舊版期限**：2025/26 PDF原3/7已刪線改紅字3/14，原先兩日期待選已解。仍保留舊brochure大三與2025/26大二的版本差異，不任選改成2027。掃描brochure12頁費用、課程與宿型全部讀核。
- **梨花**：OIA1165的GPA3.0/4.3與iBT72是梨花學生OUTGOING；臺大incoming對應1142，male/female皆可、2.5/4.0（75/100）。已循正式Application讀到4項初申文件、4/15–5/15／10/15–11/15窗口，以及錄取後另交航班、保險、visa。費用免學費／application與Visiting收費、optionalhealthcenter均分開。
- **UW–Madison**：臺大明列IAPGeneral；新語測band4.5、舊95、IELTS6.5、DET125，不採Engineering90。初申3/1／10/1與錄取後April/November選課分開。SHIP、宿舍lottery／夏季退寮與商學優先各有scope。
- **其他本批**：名長大普通交換與JLCP、近畿ProgramA/B、OU本科／研、USP本科／研與full-degree、Macausemester/research、Hilo公開登入殼、京都KUINEP/GEA/GESR、順天D2/D4、Bern私人寮／州保險、Northwestern季度費率等已逐來源讀核，保留原年份、身分、價格期間及不保證事項。

## 覆蓋與仍未完成

工作簿**7張sheet**：288校×14欄=4032列、來源323列、當日臺大來源3列、臺大正式共通規則11列、人工語意267列、固定正文進度501列、逐源閱讀376列；沒有公式cell，回填欄與篩選可用。

原282校×10欄**2820列完整保留、完整当期核實0列**；學校完整当期核實仍0。部分具年host事實已有，在publication欄另列`official_host_target_cycle_partial_fact`，不因一條具年事實把整欄／整校清空為已完成。原2026-10-07工作簿未取得；以當日臺大288個portal ID重建，282待定+6零名額，不能宣稱舊282逐校／逐格集合完全相同。

任一具範圍官方生活文字**272校／1963段，仍缺16校**；新增50段不等16校缺口已解。另就原46校獨立接待校生活來源缺口，第三批8校有獨立來源、22校只有臺大profile補充、**獨立缺口仍38校**；本批逐一對照38校沒有新獨立生活來源，兩種分母繼續分開。

全矩陣可讀候選未讀**{len(unread)}列**，其中原十類**{coverage['readable_unread_rows']}列**；純外部訪問限制**230列**，其中原十類**{coverage['exclusive_external_blocker_rows']}列**。可讀未讀是一項工作，不是外部阻礙。未在特定頁找到、舊年參考、source身分誤配、官方明說待發布、正式文件已發布但讀取失敗、母校／錄取後才得細節與真正版本差異均分開。即使一個來源失敗，若另有可讀正文／既有参考，不稱整欄外部阻擋。

## 臺大現況基準與後續接續

116-1時程與正式甄選簡章已公告，臺灣報名**2026/11/3 10:00–11/9 16:00**。本期名單當日282格名額空白、6校明示0；不把空白填0或把接待校總床位／接收capacity變臺大名額。臺大最後實際3入口讀取UTC10:42、288詳情UTC08:31–08:37；接待校新來源各用自己的`checked_at`，後續整理不是全網重新查核。

父端平行11校本批沒有輸入，沒有等待或憑空引用。已保存原優先roster、目前逐源ledger及401未完成集合。另封存932個原sourceID／890個本文SHA的抽取文字與937筆校源欄對照，2.92MB可攜ZIP供下一任務繼續閱讀；這是未核實候選文字，掃描／缺字仍須目視，不宣稱未讀原檔全部歸檔。，可按文件／期限／資格／語言／費用接續讀。可讀安全來源尚未耗盡，不能把此批完成當整體缺口完成。日檢查排程由父端負責，本任務未另建。

## 驗證、交付与限制

当前**608則／SHA45a9d975e8beae64bad8d75de4069e76abbec493e700563eb7d3ae13acb87d8d**通過96 focused tests（1 skip）、13 browser tests（390/1440）、全288校受保護欄位与兩步重建byte-idempotence；校內名額／資格／截止、workflow、其他母校資料、shared app.js維持基準。固定PR8 9927ec...組合視覺／座標／校名23項測試通過（另7項guard與1skip）；重算交集只有global座標文件，NTU海外16筆與PR8臺灣4筆record交集空。未merge/deploy，也沒有改PR8。

第三批全套2182/13skip與PR8browser17只列先前證據，不冒充此次608則重跑全套。最初兩個測試失敗紀錄保留，已修第三批固定數斷言及metadata重建；另外補4筆PDF頁码并纠正6个sourcequote索引頁，沒有新增fixture或mock掩盖。

Library先前HTTP401且父端停止重試，本批沒有LibraryID。改用同DraftPR9的固定commit提供Excel、Markdown、候選JSON、ZIP與原始本文分件。ZIP是回填與核對包：包括矩陣、候選、ledger、原清單、22頁核對圖与body manifest；16个rawbodyZIP另在`sources/ntu`，不重複包入，避免45.9MB扫描原件重复。所有分件由manifest SHA可驗证；前3批交付保留。
'''
(S/'問題解決報告.md').write_text(report)
for name in ['問題解決報告.md','review-summary.json','288校逐欄查核_可回填.xlsx','all-school-review.csv','all-school-review.json','daily-baseline-2026-10-11.json','remaining-data-exceptions.json','46-life-gap-source-origin-comparison.json','remaining-independent-host-life-research.json','readable-body-work-queue.json','282-school-ten-field-coverage.json']:shutil.copyfile(S/name,D/name)
archivepath=S/'臺大288校候選查核_可回填.zip'
with zipfile.ZipFile(archivepath,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in S.rglob('*'):
  if p.is_file() and p!=archivepath and p.suffix!='.zip' and 'browser' not in p.parts and p.name not in ['tested-source-snapshot.json','focused.log','deliverable-manifest.json']:
   z.write(p,p.relative_to(S))
 z.write(R/'sources/ntu/priority-original-cached-texts-2026-10-11.zip','接續原候選文字/priority-original-cached-texts-2026-10-11.zip')
 z.writestr('原始本文取得說明.txt','16個原始本文分件保存在同repo sources/ntu/，請依reviewed-source-bodies-current-manifest.json下載並驗SHA。ZIP未重複包含原始本文分件，含22張核對圖。\n')
shutil.copyfile(archivepath,D/archivepath.name)
delivery=[]
for p in [S/'288校逐欄查核_可回填.xlsx',S/'問題解決報告.md',S/'結構化候選資料.json',archivepath]:delivery.append({'path':str(p.relative_to(R)),'size_bytes':p.stat().st_size,'sha256':sha(p)})
write(S/'deliverable-manifest.json',delivery);write(D/'deliverables-manifest.json',delivery);print(json.dumps({'deliverables':delivery,'sheet_rows':sheets,'fixed_priority_complete':100,'life_any_gap':16,'independent_life_gap':38},ensure_ascii=False))
