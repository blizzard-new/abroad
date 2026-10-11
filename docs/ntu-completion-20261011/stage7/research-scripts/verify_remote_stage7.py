from pathlib import Path
import os,json,hashlib,base64,zipfile
from urllib.request import Request,urlopen
from openpyxl import load_workbook

ROOT=Path('/workspace/abroad-ntu-completion-20261011');HERE=Path(__file__).parent
publication=json.loads((HERE/'publication.json').read_text());COMMIT=publication['remote_commit'];TREE=publication['remote_tree']
headers={'Authorization':'Bearer '+os.environ['GH_TOKEN'],'Accept':'application/vnd.github+json','X-GitHub-Api-Version':'2022-11-28'}
def get(path):
    request=Request('https://api.github.com/repos/blizzard-new/abroad/'+path,headers=headers)
    with urlopen(request,timeout=240) as response:
        return json.load(response)
commit=get('git/commits/'+COMMIT);assert commit['tree']['sha']==TREE and commit['parents'][0]['sha']=='6417aaa38fb1172aa3ad95f8cbc0c7ea62a02d3a'
tree=get('git/trees/'+TREE+'?recursive=1');assert not tree['truncated']
paths={r['path']:r for r in tree['tree'] if r['type']=='blob'}
manifest=json.loads((HERE/'publish-manifest.json').read_text());assert all(paths[r['path']]['sha']==r['blob_sha'] for r in manifest)
assert {'src/exchange_info/ntu_partner_summary.py','tests/test_ntu_completion.py','tests/e2e/test_ntu_reference_scope.py'} <= {r['path'] for r in manifest}
tested_blobs=json.loads((ROOT/'docs/ntu-completion-20261011/stage7/validation/materialized-byte-receipts.json').read_text())
assert all(paths[r['path']]['sha']==r['git_blob_oid'] for r in tested_blobs)
destination=HERE/'remote-downloads';destination.mkdir(exist_ok=True);downloads=[]
for path in ['docs/ntu-completion-20261011/stage7/288校逐欄查核_可回填.xlsx','docs/ntu-completion-20261011/stage7/結構化候選資料.json','docs/ntu-completion-20261011/stage7/臺大288校候選查核_可回填.zip','docs/ntu-completion-20261011/stage7/問題解決報告.md']:
    blob=get('git/blobs/'+paths[path]['sha']);assert blob['encoding']=='base64';data=base64.b64decode(blob['content']);local=(ROOT/path).read_bytes();assert data==local
    target=destination/Path(path).name;target.write_bytes(data)
    downloads.append({'path':path,'blob_sha':paths[path]['sha'],'size_bytes':len(data),'sha256':hashlib.sha256(data).hexdigest(),'byte_identical_to_local_commit':True})
book=load_workbook(destination/'288校逐欄查核_可回填.xlsx',read_only=True);assert len(book.sheetnames)==7 and book['288校逐欄回填'].max_row==4033 and book['來源原文索引'].max_row==349;book.close()
with zipfile.ZipFile(destination/'臺大288校候選查核_可回填.zip') as z:
    assert z.testzip() is None
    assert json.loads(z.read('結構化候選資料.json'))['schools']==json.loads((destination/'結構化候選資料.json').read_text())['schools']
pr9=get('pulls/9');pr8=get('pulls/8');main=get('git/ref/heads/main')
assert pr9['head']['sha']==COMMIT and pr9['draft'] and not pr9['merged']
assert pr8['head']['sha']=='9927ec1d8765a9772671bbaa68e7f920995f1258' and pr8['draft'] and not pr8['merged']
assert main['object']['sha']=='59a20f27fcee00fb597abc79e75ec11778317422'
receipt={'remote_commit':COMMIT,'remote_tree':TREE,'all_incremental_blob_shas_match':True,'all_tested_code_and_materialized_blob_shas_match':True,'tested_code_and_materialized_blob_count':len(tested_blobs),'new_summary_adapter_and_regression_tests_published':True,'verified_downloads':downloads,'downloaded_excel_sheets':7,'downloaded_excel_rows':4032,'downloaded_zip_integrity_verified':True,'PR9_draft_unmerged':True,'PR8_unchanged_draft_unmerged':True,'main_unchanged':True,'library_file_ids':[]}
(HERE/'remote-publication-verification.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(receipt,ensure_ascii=False))
