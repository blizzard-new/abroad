from pathlib import Path
from collections import defaultdict,Counter
import json,hashlib,datetime,zipfile,shutil
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
H=Path(__file__).parent;R=Path('/workspace/abroad-ntu-completion-20261011');D=R/'docs/ntu-completion-20261011/stage5';B=D.parent/'stage4'
read=lambda p:json.loads(p.read_text())
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
write=lambda p,x:p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
delta=read(D/'候選增量.json');combined=read(D/'結構化候選資料.json');progress=read(D/'priority-review-progress.json');manifest=read(D/'已讀原文manifest.json');ledger=read(D/'第五批完整閱讀增量.json');roster=read(D/'priority-source-roster.json')
shutil.copyfile(H/'export_deliverables.py',D/'research-scripts/export_deliverables.py')
assert len(delta['observations'])==51 and len(combined['schools'])==288 and sum(len(s['observations']) for s in combined['schools'])==659
assert sha(R/'sources/ntu/completion-2026-10-11.json')==delta['base_stage4_source_sha256']
dispositions=[]
notes={
 '1268':{'eligibility':('not_found_in_read_public_sources','OIA正文多為Bennett赴外與短期InboundImmersion，未列incomingsemester學術門檻。沒有證明全校確實未公告，其他公開入口仍可找。'),'cost':('public_followup_access_unavailable_with_readable_reference','宿舍頁可讀，但無交換費率；其26/27Policy PDF CDN robots不可用。不能將整欄當外部阻擋，也不能把滿額新生公告當2027交換規則。')},
 '89':{'eligibility':('undated_general_rule_and_old_year_reference','母校提名一般incoming與26/27相近學科條件；OUTGOING不是申請者身分。'),'deadlines':('undated_general_rule','分母校提名及學生初申日月，ECOL、SwissMobility、ISP各别；未確定2027–28。'),'documents':('undated_initial_application_documents','一般incoming三項初申文件已列；登入後的個人tasks不可推成公开页完全没文件。'),'cost':('older_year_reference_and_version_conflict','26/27費表有清楚月單位；春宿期同年兩PDF相差5/31與6/30，不任選。')},
 '809':{'eligibility':('prior_entry_cycle_reference','2027Spring要項已公告並有GPA3/4；不是Fall2027初申。JASSO2.3/3不是交換硬門檻。'),'deadlines':('prior_entry_cycle_dates_published','2027Spring的2026/11/2初申確定；不將舊年資料或全年Fallcolumn改成Fall2027初申。'),'documents':('published_prior_cycle_and_after_admission_details','22頁已讀；初申文件、後續COEwebform、錄取後Health與housing清楚分階段。'),'cost':('announced_fee_effective_from_spring2027_with_unreconciled_estimate','50000SupportFee明說從Spring2027起；包含保險與未標年生活表可重疊，未核完整2027–28總預算。')},
 '30':{'eligibility':('program_specific_general_rules','JTW通常成功GPA3.2不是硬最低；JLCC原則非大一本科日本主修、研究生需個案，不採full-degree。'),'language':('program_specific_general_rules','JTW英79/6、JLCC日N1望ましい、一般學部日N1或有足量英授時英79/6不可混用。'),'deadlines':('undated_general_and_home_coordinator_details','JTW秋2/10春10/10未標年；JLCC明說母校案內日，公共住宅／commonfacility日不是交換初申。'),'documents':('partly_published_with_unread_public_links','JLCC成績日語／研修計畫推薦等列出，JTWessay/project列出；基本表與Guide新連結尚未跟讀，未當外部阻礙。'),'cost':('program_specific_waivers_and_historical_rates','JTW學分費明標2019，JLCC免三項費用依有效協定；生活總額頁未寫當期與期間不補單位。')},
 '391':{'eligibility':('undated_partner_nomination_rule','母校提名才申請；課程相近／degreecycle2學期各別，沒有找到具年GPA硬線。'),'language':('undated_general_rule_with_department_exceptions','申請B1入學B2、部分初申B2入學C1；只英授不用德語，接受考試型態與立命館不同。'),'deadlines':('undated_student_deadline_nomination_followup_unread','初申冬5/15夏11/15與宿舍勾選同截止；提名另一頁尚待跟，不冒充初申日為提名日。'),'documents':('undated_initial_vs_after_admission_documents','六項初申與LA签署階段；到校保險、地址登記、銀行文件不同用途。'),'cost':('old_rate_and_unknown_target_rate','Registrar2026夏97EUR、舊85保留；2027–28新費率未在此讀頁列，不足以證明全站未公告。')}
}
for pid,school in roster.items():
 if pid not in notes:continue
 for field,sids in school['fields'].items():
  outcome,note=notes[pid][field];dispositions.append({'portal_id':pid,'school_name':school['name'],'field':field,'outcome':outcome,'note_zh':note,'read_original_source_ids':sids,'all_original_cached_body_texts_read':True,'full_current_cycle_requirement_verified':False})
assert len(dispositions)==20
write(D/'第五批逐項語意待解.json',dispositions)
byfield=defaultdict(list)
for o in delta['observations']:byfield[o['portal_id'],o['field']].append(o)
progressmap={(p['portal_id'],p['field']):p for p in progress['fields']}
book=load_workbook(B/'288校逐欄查核_可回填.xlsx');main=book['288校逐欄回填'];headers=[c.value for c in main[1]];header={name:i+1 for i,name in enumerate(headers)}
for name in ['reference_present','target_requirement_confirmed','status','reviewed_observations','publication_status','field_access_status','review_progress','manual_review_outcome','unread_candidate_source_ids','unread_candidate_body_count','external_blocker_exclusive']:
 main.cell(1,header[name]).comment=Comment('此欄保留已驗證的第四批608則基準；第五批51則候選另見新增4欄與專屬sheet。原集合完整閱讀進度欄已更新至120/501，但尚未重跑runtime分類器。','Uniway research')
newheaders=['第五批候選事實（未匯入runtime）','第五批官方URL','第五批適用範圍與頁碼','第五批語意判讀']
for i,name in enumerate(newheaders,len(headers)+1):main.cell(1,i,name)
dispositionmap={(d['portal_id'],d['field']):d for d in dispositions}
for row in main.iter_rows(min_row=2):
 pid=str(row[0].value);field=row[3].value;key=pid,field
 if p:=progressmap.get(key):
  for name,val in [('fixed_priority_original_body_set_complete',p['all_original_cached_candidates_read']),('original_cached_candidate_count',p['original_cached_candidate_body_count']),('original_cached_read_count',p['read_count']),('original_cached_remaining_source_ids',json.dumps(p['remaining_source_ids']))]:main.cell(row[0].row,header[name],val)
 if os:=byfield.get(key):
  vals=['\n'.join(o['text_zh'] for o in os),'\n'.join(dict.fromkeys(o['source_url'] for o in os)),'\n'.join(o['scope']+(' / p'+str(o['pdf_page']) if o['pdf_page'] else '') for o in os),dispositionmap.get(key,{}).get('note_zh','新公開候選補充；不代表完整當期欄位核實')]
  for i,val in enumerate(vals,len(headers)+1):main.cell(row[0].row,i,val)
fixed=book['固定501項正文進度'];ph=[c.value for c in fixed[1]]
for row in fixed.iter_rows(min_row=2):
 p=progressmap[str(row[0].value),row[2].value]
 for i,key in enumerate(ph,1):fixed.cell(row[0].row,i,json.dumps(p[key],ensure_ascii=False) if isinstance(p[key],(list,dict)) else p[key])
def add_sheet(title,rows,keys):
 s=book.create_sheet(title);s.append(keys)
 for r in rows:s.append([json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(dict,list)) else r.get(k) for k in keys])
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='173B54')
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=25
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
add_sheet('第五批51則候選',delta['observations'],['portal_id','school_name','field','text_zh','source_url','scope','applicability_status','pdf_page','evidence_locator','original_quote','checked_at','raw_body_sha256','selection_rule_confirmed'])
add_sheet('第五批40個已讀來源',manifest['source_bodies'],['source_id','url','final_url','title','checked_at','http_status','document_type','pdf_page_count','sha256','archive_body_path','text_chars','candidate_current_cycle_confirmed'])
add_sheet('第五批20項待解',dispositions,['portal_id','school_name','field','outcome','note_zh','read_original_source_ids','all_original_cached_body_texts_read','full_current_cycle_requirement_verified'])
add_sheet('第五批100則閱讀',ledger,['portal_id','school_name','field','source_id','source_url','body_sha256','source_checked_at','reviewed_at','full_extracted_text_read_for_field','pdf_page_count','note_zh','target_cycle_requirement_verified'])
book.save(D/'288校逐欄查核_可回填.xlsx');book.close()
check=load_workbook(D/'288校逐欄查核_可回填.xlsx',read_only=True,data_only=False);sheetrows={s.title:s.max_row-1 for s in check};assert sheetrows['288校逐欄回填']==4032 and len(sheetrows)==11
assert not any(c.data_type=='f' for s in check for row in s for c in row)
assert sum(row[5] is True for row in check['固定501項正文進度'].iter_rows(min_row=2,values_only=True))==120
basebook=load_workbook(B/'288校逐欄查核_可回填.xlsx',read_only=True,data_only=False)
baseflags={(str(row[0]),row[3]):row[6] for row in basebook['288校逐欄回填'].iter_rows(min_row=2,values_only=True)}
newflags={(str(row[0]),row[3]):row[6] for row in check['288校逐欄回填'].iter_rows(min_row=2,values_only=True)}
assert newflags==baseflags
basebook.close()
check.close()
base=read(B/'daily-baseline-2026-10-11.json');base.update(stage='stage5_candidate_only',research_checkpoint_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),runtime_validated_observations=608,combined_candidate_observations=659,incremental_candidate_observations=51,fixed_priority_progress={k:v for k,v in progress.items() if k!='fields'},new_host_retrieval_window_utc=['2026-10-11T14:15:59+00:00','2026-10-11T14:20:12+00:00'],timestamp_note='臺大正式入口實際最後讀取仍UTC10:42，288詳情08:31–08:37；第五批只新讀接待校公開來源及原cached全文，不將整理日當全站重查。')
write(D/'daily-baseline-2026-10-11.json',base)
report='''# 第五批：五校51則可回填候選（2026-10-11）

照片與座標各仍缺3校。照片：UC(9)、CSU(12)、特魯瓦(67)；座標：UC(9)、CSU(12)、蒙斯(1302)。UC/CSU未取得正式分校分配；特魯瓦照片仍缺可驗證重用授權；蒙斯Place du Parc20官方地址已核，尚無同代表校區可交叉核實的座標。未新增資產，具體來源與拒用理由沿用[第四批資產證據](../stage4/evidence/asset-gap-review.json)。

## 本批完成與適用界線

繼第四批175新增後，本批實際讀完貝內特、聖加倫、立命館、九州與慕尼黑的全部原優先cached正文，新增51則事實、覆蓋32個校欄；另循公開連結新增11次抓取，10可讀、1CDN robots不可用。來源16個被引用（其中10是新來源ID）；合併候選為659則／332接待校sourcequote。40個已讀來源版本（37不同本文SHA）原檔與抽取文字全部封存並驗SHA，逐校欄閱讀增加100列。

固定原優先分母仍501校欄／117校／890本文SHA，原集合完整閱讀由100增至**120/501**，餘**381**，完成涉及35校；300列原校欄來源、176列另循連結，全文ledger共476列。這是完整讀過原cached集合，不等完整當期規定，也不包括尚未跟讀的新文件連結。

第四批608則已重建並通過runtime測試；本批51則只在docs候選交付，**沒有匯入runtime、修改workflow、merge或deploy**。目前runtime與validatedsource SHA仍45a9d975…。合併659則JSON標示candidate_only；不能把第四批96+13測試冒充659則的runtime驗證。

## 五校具體解決與待解

|學校|本批核實|仍待逐項處理|
|---|---|---|
|貝內特|2026–27宿舍滿額公告限定7/25/2026後新生，不當2027交換無房；OIA多赴外／短期immersion|semesterincoming硬資格未在已讀頁列、交換費率未取得；HostelPolicy PDF CDN robots不可用，但另有可讀頁，非整欄訪問阻擋|
|聖加倫|incoming秋提名4/15初申5/15、春9/15與9/30分開；三項初申文件；26/27英德B2不須正式證書，ISP89/6.5不套一般；月預算與課程16–40ECTS|同為26/27的兩個官方PDF春住宿結束5/31與6/30矛盾（各p5）；公開housing春滿額未標哪年；2027–28完整規則未核|
|立命館|2027Spring22頁要項完整讀：GPA3/4、RCJ/IJL/RCE、初申與COE/Health/housing階段、10/2–11/2/2026JST申請；Spring2027起50000SupportFee含指定保險、RUSSES結束|不是Fall2027初申要項；Fall2027日期欄限定Spring2027開始全年者；SupportFee保險與未標年生活估算可能重疊，尚未核總預算；後續COE表明說另提供|
|九州|JTW通常成功GPA3.2非硬最低、研可申；JLCC原則本科非大一日本主修、N1望ましい；JTW秋2/10春10/10日月、JLCC母校案內期限；JTW2019學分費與JLCC協定互免分開|basicform/guide及admittedacademicguide公開新連結尚待讀；JLCC具年期限由母校案內；generalcost未標費期／目標年，不將公共住宅截止當exchange|
|慕尼黑|初申冬5/15夏11/15；六項文件與LA簽署階段；初申B1入學B2及部門C1例外、可母校languageconfirmation；60%major/minor、30ECTS非硬最低；宿舍400–600EUR/月、約半數非個人保障|Nomination另一頁尚待讀；2027–28費率未核，Registrar只有2026夏97EUR等舊表；德法定醫保證明要求已核，但臺灣visa例外與本期保費未完成|

每則事實的URL、短原文、PDF頁／locator、年／學期／校區／身分scope、實際source checked_at及rawSHA均在[候選增量](候選增量.json)與工作簿。「未在已讀頁找到」、「確定另母校／錄取後才給」、「已發布但特定CDN讀不到」、「未讀公開新連結」、「同年文件矛盾」分開，沒有證據時不說確實未公告。

## 工作簿與基準

11張sheet含全288校×14欄4032列與原282校×10類2820列，保留可回填欄；主表新增第五批事實／URL／scope頁碼／語意4欄，原正式核實旗標不變。501進度表與主表原集合閱讀4欄更新至120，額外51候選、40已讀source、20原優先欄待解、100閱讀增量各一sheet。原主表publication與未讀矩陣統計仍是第四批基準，**不能把843列当第五批重新計算的剩餘數**；本批未重新執行runtime分類器。

原2820完整當期核實仍0，舊2026/10/7工作簿仍未取得，當日重建的282集合與舊集合等同與否不可證明。任一官方生活來源缺16校、原46校獨立接待校來源缺38校，是第四批runtime基準；本批沒有匯入life或新增照片座標，分母不混算。臺大116-1已公布，報名臺灣時間2026/11/3 10:00至11/9 16:00；本期282空白名額與6零名額不變，不以接待容量代填。

臺大最後實際入口讀取UTC10:42，詳情08:31–08:37；本批新host抓取UTC14:15:59–14:20:12，原cached來源保留早上checked_at。日檢查基準另存，未重複建排程、未發校方郵件；父端11校沒有新輸入，未等待或臆造。

## 驗證與交還

51/51原文引用均可在封存文本找到；PDF引用逐頁吻合。40/40原檔SHA驗證、ZIP完整性、288/659合併候選、固定501原分母／120完成（另逐一比原rawSHA、非僅同sourceID）、工作簿4032列及11sheet、0公式均通過。5頁关键表另渲染目視（立命館p2/p8/p19、聖加倫兩版p5）；工作簿未以ExcelGUI逐表目視，結構與所有cell值已核對，長文可展開cell閱讀。正式source、runtime及其他檔案維持第四批，不需要為docs候選冒稱已重跑runtime測試。第四批遠端4交付已下載byte-identical，165檔Gitblob SHA吻合、7sheet4032列與ZIP完整，PR9仍Draft，main59a20f及PR8 9927ec未動。

Library先前HTTP401，父端停止重試，**沒有Library檔案ID**。本批以同DraftPR9固定commit與四檔交付供接續，不宣稱全282校十類已解。381組仍有安全可讀候選，真正限制作逐項記錄；下一批繼續文件、期限、資格、語言、費用及剛發現的公開指南。
'''
(D/'問題解決報告.md').write_text(report)
validation={'candidate_only':True,'new_observations':51,'combined_observations':659,'schools':288,'source_quote_ids':332,'original_priority_done':120,'original_priority_total':501,'remaining_original_priority_sets':381,'new_full_body_field_rows':100,'total_full_body_field_rows':476,'raw_source_versions':40,'raw_distinct_hashes':37,'excel_sheet_rows':sheetrows,'formula_cells':0,'runtime_source_unchanged_sha256':sha(R/'sources/ntu/completion-2026-10-11.json'),'runtime_tested_observations':608,'runtime_tests_rerun_for_659':False,'quality_assessment':'Share with caveats — scoped candidates only','pdf_table_pages_visually_reviewed':5,'excel_GUI_render_review_performed':False,'deploy_performed':False,'library_file_ids':[]}
write(D/'validation.json',validation)
with zipfile.ZipFile(D/'臺大288校候選查核_可回填.zip','w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in sorted(D.rglob('*')):
  if p.is_file() and p.name not in ['臺大288校候選查核_可回填.zip','deliverable-manifest.json']:z.write(p,p.relative_to(D))
deliver=[{'path':str(p.relative_to(R)),'size_bytes':p.stat().st_size,'sha256':sha(p)} for p in [D/'288校逐欄查核_可回填.xlsx',D/'問題解決報告.md',D/'結構化候選資料.json',D/'臺大288校候選查核_可回填.zip']]
write(D/'deliverable-manifest.json',deliver);print(json.dumps({'deliverables':deliver,'validation':validation},ensure_ascii=False))
