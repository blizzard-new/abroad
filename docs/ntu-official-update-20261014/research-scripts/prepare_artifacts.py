from pathlib import Path
import json,hashlib,shutil,csv
from openpyxl import Workbook
H=Path(__file__).parent;R=Path('/workspace/abroad-ntu-official-20261014');D=R/'docs/ntu-official-update-20261014';D.mkdir(exist_ok=True);source=json.loads((R/'sources/ntu/2026-10-14-official-update.json').read_text());rows=[]
for q in source['quota_updates']:
 rows.append({'portal_id':q['portal_id'],'school_name':q['school_name'],'category':'臺大正式名額','scope':'116-1／2027/2028；學年名額／最多人次','value':q['quota_display']+'；'+q['exchange_duration_original'],'source_url':source['list_receipt']['url'],'detail_url':q['detail_url'],'pdf_page':None,'checked_at':q['checked_at'],'source_sha256':source['list_receipt']['sha256'],'conclusion':'名額數值已核實；選填系統及其他條件仍待確認','legal_rule_independently_verified':False,'first_fee_discovery':False})
for v in source['visa_reminders']:
 rows.append({'portal_id':v['portal_id'],'school_name':v['school_name'],'category':'臺大新增簽證提醒','scope':'臺大澳洲交換頁；非獨立法律核實','value':v['text_zh'],'source_url':v['source_url'],'detail_url':v['official_external_url'],'pdf_page':None,'checked_at':v['checked_at'],'source_sha256':source['source_manifest'][v['source_id']]['sha256'],'conclusion':'臺大提醒已核原件；個人法規與簽證資格另核','legal_rule_independently_verified':False,'first_fee_discovery':False})
u=source['ubc_update']
for f in u['facts']:
 rows.append({'portal_id':'1','school_name':u['school_name'],'category':'UBC接收參考：'+f['field'],'scope':u['scope'],'value':f['text_zh'],'source_url':u['source_url'],'detail_url':u['replaces_source_url'],'pdf_page':f['pdf_page'],'checked_at':u['checked_at'],'source_sha256':u['source_sha256'],'conclusion':'新年度／原版本差異已核；非臺大甄選期限或名額','legal_rule_independently_verified':False,'first_fee_discovery':False})
for p in source['reference_updates']:
 rows.append({'portal_id':p['portal_id'],'school_name':p['school_name'],'category':'臺大參考文字／連結更新','scope':'僅核臺大頁面指向；非獨立重查接待校規定','value':p['note_zh'],'source_url':p['source_url'],'detail_url':'\n'.join(l['url'] for s in p['sections'] for l in s['links']),'pdf_page':None,'checked_at':p['checked_at'],'source_sha256':source['source_manifest'][p['source_id']]['sha256'],'conclusion':'原资格／最低語言門檻未改','legal_rule_independently_verified':False,'first_fee_discovery':False})
assert len(rows)==31
(D/'本批候選與正式更新.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n')
shutil.copyfile(R/'sources/ntu/2026-10-14-official-update.json',D/'結構化更新快照.json')
book=Workbook();sheet=book.active;sheet.title='本批31項更新';keys=list(rows[0]);sheet.append(keys)
for row in rows:sheet.append([row[k] for k in keys])
sheet.freeze_panes='C2';sheet.auto_filter.ref=sheet.dimensions
for name,width in [('A',14),('B',34),('C',30),('D',56),('E',100),('F',75),('G',75)]:sheet.column_dimensions[name].width=width
from openpyxl.styles import Alignment,Font
for c in sheet[1]:c.font=Font(bold=True)
for row in sheet.iter_rows(min_row=2):
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
s=book.create_sheet('22來源原件索引');s.append(['source_id','url','checked_at','sha256','archive_path','http_status'])
for sid,m in source['source_manifest'].items():s.append([sid,m['url'],m['checked_at'],m['sha256'],m['local_file'],m['http_status']])
s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;book.save(D/'臺大116-1本批更新_可回填.xlsx')
for name in ['online-before-receipt.json','online-before-ntu-dashboard.json','live-before-health.txt']:
 if (H/name).exists():shutil.copyfile(H/name,D/'validation'/name)
(D/'每日來源基準.json').write_text(json.dumps({'source_snapshot_sha256':hashlib.sha256((R/'sources/ntu/2026-10-14-official-update.json').read_bytes()).hexdigest(),'list_receipt':source['list_receipt'],'source_manifest':source['source_manifest'],'counts':source['counts'],'all_current_cycle_fields_verified':False,'application_open_confirmed':False,'parent_attachment_tracking_report':source['parent_attachment_tracking_report'],'timestamp_note':'原件checked_at是實際取件；未以本批閱讀、測試或編譯日期替代來源日期。父端每日任務維持，本任務未新增排程。'},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'delta_rows':len(rows),'archived_sources':len(source['source_manifest']),'excel_sheets':book.sheetnames},ensure_ascii=False))
