import base64
import hashlib
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from pypdf import PdfReader, PdfWriter

from exchange_info.api import create_app
from exchange_info.models import (
    AccountSessionRecord,
    ProductAccessTier,
    UserAccount,
    utc_now,
)
from exchange_info.storage import (
    ensure_data_dirs,
    save_account_sessions,
    save_user_accounts,
)


def _authenticated_client(
    data_dir: Path,
    *,
    product_tier: ProductAccessTier = ProductAccessTier.FREE,
    home_school: str = "長庚大學",
) -> TestClient:
    ensure_data_dirs(data_dir)
    account = UserAccount(
        display_name="申請文件測試學生",
        product_tier=product_tier,
        student_status_verified=True,
        school_email_verified=True,
        verified_university=home_school,
        verified_department="資訊工程學系",
    )
    token = f"application-document-{product_tier.value}-token"
    save_user_accounts([account], data_dir)
    save_account_sessions(
        [
            AccountSessionRecord(
                account_id=account.id,
                token_hash=hashlib.sha256(token.encode("utf-8")).hexdigest(),
                expires_at=utc_now() + timedelta(days=1),
            )
        ],
        data_dir,
    )
    client = TestClient(create_app(data_dir))
    client.headers.update({"Authorization": f"Bearer {token}"})
    return client


def _study_plan_payload(copies: int = 1) -> dict[str, object]:
    return {
        "context": {
            "school_name": "長庚大學",
            "target_school": "林肯大學 Lincoln University",
            "target_country": "紐西蘭",
            "program_type": "exchange",
            "academic_year": "116",
            "semester": "第一學期",
            "flags": [],
        },
        "profile": {
            "student_name_zh": "王小明",
            "student_id": "B1234567",
            "department": "資訊管理學系",
        },
        "answers": {
            "study-plan": {
                "motivation": "希望透過交換培養跨文化合作能力。",
                "school_reason": "課程與我的學習方向相符。",
                "learning_goals": "完成專業課程並提升英文溝通能力。",
                "course_plan": "修習資料分析與跨文化管理課程。",
                "expected_outcomes": "完成課程並分享交換經驗。",
                "return_plan": "返國後完成學分採認與成果分享。",
            }
        },
        "selections": [
            {
                "template_id": "study-plan",
                "copies": copies,
                "reviewed_version": "2026-07-16",
            }
        ],
    }


def _review_current_content(
    client: TestClient,
    payload: dict[str, object],
) -> dict[str, object]:
    manifest = client.post("/application-center/manifest", json=payload)
    assert manifest.status_code == 200
    items = {item["template_id"]: item for item in manifest.json()["items"]}
    for selection in payload["selections"]:  # type: ignore[index]
        selection["reviewed_content_sha256"] = items[selection["template_id"]][  # type: ignore[index]
            "review_content_sha256"
        ]
    return manifest.json()


def _png_data_url() -> str:
    image = BytesIO()
    PILImage.new("RGB", (8, 8), "white").save(image, format="PNG")
    return "data:image/png;base64," + base64.b64encode(image.getvalue()).decode("ascii")


def test_application_center_resolves_conditional_documents(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)

    base = client.get(
        "/application-center/templates",
        params={"program_type": "exchange"},
    ).json()
    subsidized = client.get(
        "/application-center/templates",
        params={"program_type": "exchange", "flags": "subsidy,military_service"},
    ).json()

    assert any(row["id"] == "study-plan" for row in base)
    assert all(row["id"] != "subsidy-receipt" for row in base)
    assert any(row["id"] == "subsidy-receipt" for row in subsidized)
    assert any(row["id"] == "military-service-overseas" for row in subsidized)


def test_verified_free_students_can_read_official_forms_while_generation_requires_subscription(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setenv("EXCHANGE_PRODUCT_RELEASE_MODE", "public")
    payload = _study_plan_payload()
    guest = TestClient(create_app(tmp_path / "guest"))

    forms = guest.get(
        "/application-center/templates",
        params={"program_type": "short_term_credit"},
    )

    assert forms.status_code == 200, "official forms are public; only generating documents needs an account"
    assert guest.post("/application-center/manifest", json=payload).status_code == 401
    assert guest.post("/application-center/pdf", json=payload).status_code == 401

    free = _authenticated_client(tmp_path / "free")
    forms = free.get("/application-center/templates", params={"program_type":"short_term_credit"})
    assert forms.status_code == 200
    assert all(row["official_download_url"].startswith("https://www.cgu.edu.tw/") for row in forms.json())
    assert free.get("/application-center/templates", params={"school":"國立臺北大學"}).status_code == 200
    manifest = free.post("/application-center/manifest", json=payload)
    pdf = free.post("/application-center/pdf", json=payload)
    assert manifest.status_code == 403
    assert pdf.status_code == 403
    assert "進階功能" in manifest.json()["detail"]

    subscription = _authenticated_client(
        tmp_path / "subscription",
        product_tier=ProductAccessTier.SUBSCRIPTION,
    )
    allowed = subscription.post("/application-center/manifest", json=payload)
    assert allowed.status_code == 200
    assert allowed.headers["cache-control"] == "private, no-store, max-age=0"
    assert allowed.headers["x-content-type-options"] == "nosniff"


def test_application_manifest_reports_missing_fields_and_print_copies(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)
    payload = _study_plan_payload(copies=3)
    first_preview = _review_current_content(client, payload)
    assert first_preview["items"][0]["reviewed"] is False  # type: ignore[index]

    complete = client.post("/application-center/manifest", json=payload)
    assert complete.status_code == 200
    assert complete.json()["ready_to_generate"] is True
    assert complete.json()["total_printed_copies"] == 3
    item = complete.json()["items"][0]
    assert item["reviewed"] is True
    assert len(item["review_content_sha256"]) == 64
    assert next(
        field for field in item["preview_fields"] if field["field_id"] == "motivation"
    )["display_value"] == "希望透過交換培養跨文化合作能力。"
    optional_english = next(
        field for field in item["preview_fields"] if field["field_id"] == "english_version"
    )
    assert optional_english["required"] is False
    assert optional_english["is_blank"] is True
    assert optional_english["display_value"] == ""
    assert item["paper_size"] == "A4"
    assert item["print_sides"] == "single_sided"
    assert "A4 單面列印" in item["print_note"]
    assert item["attachment_order"] == []
    assert item["source_url"].startswith("https://www.cgu.edu.tw/")

    payload["answers"]["study-plan"].pop("motivation")  # type: ignore[index]
    incomplete = client.post("/application-center/manifest", json=payload)
    assert incomplete.status_code == 200
    assert incomplete.json()["ready_to_generate"] is False
    assert "申請動機" in incomplete.json()["items"][0]["missing_required_fields"]
    assert incomplete.json()["items"][0]["reviewed"] is False
    motivation = next(
        field
        for field in incomplete.json()["items"][0]["preview_fields"]
        if field["field_id"] == "motivation"
    )
    assert motivation["is_blank"] is True
    assert motivation["display_value"] == ""


def test_application_manifest_and_pdf_require_current_template_review(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)
    payload = _study_plan_payload()

    payload["selections"][0].pop("reviewed_version")  # type: ignore[index]
    unreviewed = client.post("/application-center/manifest", json=payload)
    assert unreviewed.status_code == 200
    assert unreviewed.json()["ready_to_generate"] is False
    assert unreviewed.json()["items"][0]["reviewed"] is False
    assert unreviewed.json()["items"][0]["template_version"] == "2026-07-16"

    rejected = client.post("/application-center/pdf", json=payload)
    assert rejected.status_code == 422
    assert "目前內容的「我已核對」" in rejected.json()["detail"]

    payload["selections"][0]["reviewed_version"] = "outdated-version"  # type: ignore[index]
    outdated = client.post("/application-center/manifest", json=payload)
    assert outdated.status_code == 200
    assert outdated.json()["ready_to_generate"] is False
    assert outdated.json()["items"][0]["reviewed"] is False

    payload = _study_plan_payload()
    _review_current_content(client, payload)
    reviewed = client.post("/application-center/manifest", json=payload)
    assert reviewed.json()["items"][0]["reviewed"] is True
    payload["answers"]["study-plan"]["motivation"] = "核對後修改的內容"  # type: ignore[index]
    changed = client.post("/application-center/manifest", json=payload)
    assert changed.json()["items"][0]["reviewed"] is False
    assert changed.json()["items"][0]["review_content_sha256"] != payload["selections"][0]["reviewed_content_sha256"]  # type: ignore[index]
    assert client.post("/application-center/pdf", json=payload).status_code == 422

    payload = _study_plan_payload(copies=1)
    _review_current_content(client, payload)
    payload["selections"][0]["copies"] = 2  # type: ignore[index]
    copies_changed = client.post("/application-center/manifest", json=payload)
    assert copies_changed.json()["items"][0]["reviewed"] is False


def test_application_center_generates_a_printable_pdf_packet(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)
    files_before = {
        path.relative_to(data_dir): (path.stat().st_size, path.stat().st_mtime_ns)
        for path in data_dir.rglob("*")
        if path.is_file()
    }

    payload = _study_plan_payload(copies=2)
    _review_current_content(client, payload)
    response = client.post("/application-center/pdf", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["cache-control"] == "private, no-store, max-age=0"
    assert response.headers["pragma"] == "no-cache"
    assert response.headers["x-application-storage"] == "request-only"
    assert response.headers["x-application-packet-documents"] == "1"
    assert response.headers["x-application-packet-copies"] == "2"
    assert response.headers["x-application-packet-attachments"] == "0"
    assert response.content.startswith(b"%PDF")
    reader = PdfReader(BytesIO(response.content))
    assert len(reader.pages) >= 3
    assert response.headers["x-application-packet-pages"] == str(len(reader.pages))
    assert reader.metadata.title == "交換申請列印包"
    cover_text = reader.pages[0].extract_text() or ""
    assert "交換申請列印包" in cover_text
    assert "列印後待辦" in cover_text
    assert "伺服器不會永久保存" in cover_text
    assert "A4" in cover_text
    files_after = {
        path.relative_to(data_dir): (path.stat().st_size, path.stat().st_mtime_ns)
        for path in data_dir.rglob("*")
        if path.is_file()
    }
    assert files_after == files_before


def test_short_term_templates_are_tied_to_official_form_numbers(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)

    templates = client.get(
        "/application-center/templates",
        params={"program_type": "short_term_credit"},
    ).json()

    official = {row["official_form_number"] for row in templates if row["official_layout"]}
    assert official == {"2B0001001", "2B0001002", "2B0001003", "2B0001004"}
    assert all(row["source_url"].startswith("https://www.cgu.edu.tw/") for row in templates)


def test_required_uploads_need_real_file_payloads_and_minimum_photo_count(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)

    templates = client.get(
        "/application-center/templates",
        params={"program_type": "short_term_credit"},
    ).json()
    report = next(row for row in templates if row["id"] == "short-term-report-2b0001004")
    photo_field = next(row for row in report["fields"] if row["id"] == "activity_photos")
    assert photo_field["min_files"] == 4
    assert photo_field["max_files"] == 4

    profile: dict[str, object] = {}
    answers: dict[str, object] = {}
    png = _png_data_url()
    for field in report["fields"]:
        if not field["required"]:
            continue
        if field["field_type"] == "files":
            value: object = [{"name": f"photo-{index}.png", "type": "image/png", "data_url": png} for index in range(3)]
        elif field["field_type"] == "date":
            value = "2026-08-01"
        elif field["field_type"] == "checkbox":
            value = True
        elif field["field_type"] == "number":
            value = 1
        elif field["field_type"] == "select":
            value = field["options"][0]
        else:
            value = "測試內容"
        if field["reusable_profile_key"]:
            profile[field["reusable_profile_key"]] = value
        else:
            answers[field["id"]] = value

    payload = {
        "context": {"school_name": "長庚大學", "program_type": "short_term_credit", "flags": []},
        "profile": profile,
        "answers": {report["id"]: answers},
        "selections": [
            {
                "template_id": report["id"],
                "copies": 1,
                "reviewed_version": "2025-10-28",
            }
        ],
    }
    incomplete = client.post("/application-center/manifest", json=payload)
    assert incomplete.status_code == 200
    assert "至少 4 張照片與圖說" in incomplete.json()["items"][0]["missing_required_fields"]

    answers["activity_photos"].append({"name": "photo-4.png", "type": "image/png", "data_url": png})  # type: ignore[union-attr]
    _review_current_content(client, payload)
    complete = client.post("/application-center/manifest", json=payload)
    assert complete.status_code == 200
    assert complete.json()["ready_to_generate"] is True


def test_corrupt_optional_attachment_blocks_packet_instead_of_being_silently_dropped(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)
    corrupt = "data:image/png;base64," + base64.b64encode(b"not-a-real-image").decode("ascii")
    payload = {
        "context": {
            "school_name": "長庚大學",
            "target_school": "Partner University",
            "program_type": "exchange",
            "flags": [],
        },
        "profile": {
            "student_name_zh": "王小明",
            "student_id": "B1234567",
            "department": "資訊管理學系",
        },
        "answers": {
            "return-report": {
                "overseas_institution": "Partner University",
                "study_period": "2026 秋季班",
                "courses_summary": "完成三門專業課程。",
                "life_summary": "已完成住宿與生活適應。",
                "reflection": "返國後將分享交換經驗。",
                "activity_photos": [
                    {"name": "broken.png", "type": "image/png", "data_url": corrupt}
                ],
            }
        },
        "selections": [
            {
                "template_id": "return-report",
                "copies": 1,
                "reviewed_version": "2026-07-16",
            }
        ],
    }

    _review_current_content(client, payload)
    manifest = client.post("/application-center/manifest", json=payload).json()

    assert manifest["items"][0]["missing_required_fields"] == []
    assert manifest["items"][0]["invalid_attachment_fields"] == [
        "照片與圖說：含有無法讀取或格式不符的檔案，請重新選擇 PDF、JPG 或 PNG。"
    ]
    preview = next(
        item
        for item in manifest["items"][0]["preview_fields"]
        if item["field_id"] == "activity_photos"
    )
    assert preview["validation_error"]
    assert manifest["ready_to_generate"] is False
    rejected = client.post("/application-center/pdf", json=payload)
    assert rejected.status_code == 422
    assert "附件無法加入列印包" in rejected.json()["detail"]


def test_uploaded_pdf_is_appended_to_each_printed_document(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)
    attachment = BytesIO()
    attachment_writer = PdfWriter()
    attachment_writer.add_blank_page(width=72, height=72)
    attachment_writer.write(attachment)
    attachment_data_url = "data:application/pdf;base64," + base64.b64encode(attachment.getvalue()).decode("ascii")
    payload = {
        "context": {"school_name": "長庚大學", "program_type": "dual_degree", "flags": []},
        "profile": {
            "student_name_zh": "王小明",
            "student_id": "B1234567",
            "college": "管理學院",
            "department": "資訊管理學系",
            "degree_level": "大學部",
        },
        "answers": {
            "dual-enrollment": {
                "overseas_institution": "Partner University",
                "dual_degree_program": "Dual Degree Program",
                "planned_start_date": "2026-09-01",
                "planned_end_date": "2027-06-30",
                "admission_notice": [{"name": "admission.pdf", "type": "application/pdf", "data_url": attachment_data_url}],
            }
        },
        "selections": [
            {
                "template_id": "dual-enrollment",
                "copies": 2,
                "reviewed_version": "2026-07-16",
            }
        ],
    }

    _review_current_content(client, payload)
    response = client.post("/application-center/pdf", json=payload)

    assert response.status_code == 200
    assert len(PdfReader(BytesIO(response.content)).pages) >= 5

    payload["answers"]["dual-enrollment"]["admission_notice"] = [{"name": "admission.pdf", "type": "application/pdf"}]  # type: ignore[index]
    manifest = client.post("/application-center/manifest", json=payload)
    assert manifest.json()["ready_to_generate"] is False


def test_uploaded_pdf_attachments_follow_the_manifest_order(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)

    def attachment_data_url(width: float) -> str:
        attachment = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=width, height=72)
        writer.write(attachment)
        return "data:application/pdf;base64," + base64.b64encode(
            attachment.getvalue()
        ).decode("ascii")

    payload = {
        "context": {
            "school_name": "長庚大學",
            "target_school": "Partner University",
            "program_type": "exchange",
            "flags": [],
        },
        "profile": {
            "student_name_zh": "王小明",
            "student_id": "B1234567",
            "college": "管理學院",
            "department": "資訊管理學系",
            "degree_level": "學士班",
        },
        "answers": {
            "credit-transfer": {
                "overseas_institution": "Partner University",
                "completed_courses": "IM101｜Information Management｜3｜A",
                "requested_mapping": "IM101 → 專業選修 3 學分",
                "transcript_attachment": [
                    {
                        "name": "01-transcript.pdf",
                        "type": "application/pdf",
                        "data_url": attachment_data_url(72),
                    }
                ],
                "syllabus_attachment": [
                    {
                        "name": "02-syllabus.pdf",
                        "type": "application/pdf",
                        "data_url": attachment_data_url(144),
                    }
                ],
            }
        },
        "selections": [
            {
                "template_id": "credit-transfer",
                "copies": 1,
                "reviewed_version": "2026-07-16",
            }
        ],
    }

    manifest = client.post("/application-center/manifest", json=payload)
    assert manifest.status_code == 200
    manifest_item = manifest.json()["items"][0]
    assert manifest_item["attachment_order"] == [
        "海外正式成績單",
        "課綱／授課內容",
    ]
    assert manifest_item["invalid_attachment_fields"] == []
    assert [item["file_name"] for item in manifest_item["attachments"]] == [
        "01-transcript.pdf",
        "02-syllabus.pdf",
    ]
    assert [item["page_count"] for item in manifest_item["attachments"]] == [1, 1]
    assert all(item["placement"] == "after_document" for item in manifest_item["attachments"])
    assert manifest_item["signatures"] == ["學生親筆簽名"]
    assert manifest_item["stamps"] == ["系所審核", "教務單位核章"]
    assert manifest_item["destinations"] == ["系所／學程辦公室", "教務處註冊組或研教組"]

    payload["selections"][0]["reviewed_content_sha256"] = manifest_item["review_content_sha256"]  # type: ignore[index]
    response = client.post("/application-center/pdf", json=payload)
    assert response.status_code == 200
    assert response.headers["x-application-packet-attachments"] == "2"
    pages = PdfReader(BytesIO(response.content)).pages
    assert [float(page.mediabox.width) for page in pages[-2:]] == [72.0, 144.0]


def _word_text(content: bytes) -> str:
    document = WordDocument(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_application_center_generates_an_editable_word_document(tmp_path) -> None:
    data_dir = tmp_path / "data"
    ensure_data_dirs(data_dir)
    client = _authenticated_client(data_dir)
    files_before = {path.relative_to(data_dir) for path in data_dir.rglob("*") if path.is_file()}

    payload = _study_plan_payload(copies=2)
    _review_current_content(client, payload)
    response = client.post("/application-center/docx", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.headers["cache-control"] == "private, no-store, max-age=0"
    assert response.headers["x-application-storage"] == "request-only"
    assert response.headers["x-application-packet-documents"] == "1"
    assert response.headers["x-application-packet-copies"] == "2"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=exchange-application-documents.docx;")
    assert unquote(disposition.split("filename*=UTF-8''", 1)[1]).endswith(".docx")
    assert response.content.startswith(b"PK")
    text = _word_text(response.content)
    assert "希望透過交換培養跨文化合作能力。" in text
    assert "王小明" in text
    assert "預計列印 2 份" in text
    assert "不代表已送件" in text
    assert {path.relative_to(data_dir) for path in data_dir.rglob("*") if path.is_file()} == files_before


def test_word_export_keeps_the_same_account_completeness_and_review_gates_as_pdf(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setenv("EXCHANGE_PRODUCT_RELEASE_MODE", "public")
    payload = _study_plan_payload()
    guest = TestClient(create_app(tmp_path / "guest"))
    assert guest.post("/application-center/docx", json=payload).status_code == 401
    free = _authenticated_client(tmp_path / "free")
    assert free.post("/application-center/docx", json=payload).status_code == 403
    monkeypatch.delenv("EXCHANGE_PRODUCT_RELEASE_MODE")

    client = _authenticated_client(tmp_path / "subscription", product_tier=ProductAccessTier.SUBSCRIPTION)
    unreviewed = client.post("/application-center/docx", json=payload)
    assert unreviewed.status_code == 422
    assert "目前內容的「我已核對」" in unreviewed.json()["detail"]

    incomplete = _study_plan_payload()
    incomplete["answers"]["study-plan"].pop("motivation")  # type: ignore[index]
    _review_current_content(client, incomplete)
    missing = client.post("/application-center/docx", json=incomplete)
    assert missing.status_code == 422
    assert "必填欄位未完成" in missing.json()["detail"]

    reviewed = _study_plan_payload()
    _review_current_content(client, reviewed)
    reviewed["answers"]["study-plan"]["motivation"] = "核對後才修改"  # type: ignore[index]
    assert client.post("/application-center/docx", json=reviewed).status_code == 422

    empty = _study_plan_payload()
    empty["selections"] = []
    assert client.post("/application-center/docx", json=empty).status_code == 422


def test_word_export_of_several_documents_is_a_zip_with_one_docx_each(tmp_path) -> None:
    client = _authenticated_client(tmp_path / "data")
    payload = _study_plan_payload(copies=1)
    payload["answers"]["return-report"] = {  # type: ignore[index]
        "overseas_institution": "Lincoln University",
        "study_period": "2026/09-2027/01",
        "courses_summary": "修習三門專業課程。",
        "life_summary": "適應當地生活。",
        "reflection": "收穫很多。",
    }
    payload["selections"].append(  # type: ignore[union-attr]
        {"template_id": "return-report", "copies": 3, "reviewed_version": "2026-07-16"}
    )
    _review_current_content(client, payload)

    response = client.post("/application-center/docx", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert response.headers["x-application-packet-documents"] == "2"
    assert response.headers["x-application-packet-copies"] == "4"
    assert "filename=exchange-application-documents.zip" in response.headers["content-disposition"]
    with ZipFile(BytesIO(response.content)) as bundle:
        names = bundle.namelist()
        assert len(names) == 2
        assert names[0].startswith("01-") and names[0].endswith("-列印1份.docx")
        assert names[1].startswith("02-") and names[1].endswith("-列印3份.docx")
        assert "希望透過交換培養跨文化合作能力。" in _word_text(bundle.read(names[0]))
        assert "修習三門專業課程。" in _word_text(bundle.read(names[1]))


def test_word_export_embeds_images_and_lists_pdf_attachments_it_cannot_embed(tmp_path) -> None:
    client = _authenticated_client(tmp_path / "data")
    attachment = BytesIO()
    attachment_writer = PdfWriter()
    attachment_writer.add_blank_page(width=72, height=72)
    attachment_writer.write(attachment)
    pdf_data_url = "data:application/pdf;base64," + base64.b64encode(attachment.getvalue()).decode("ascii")
    payload = {
        "context": {"school_name": "長庚大學", "program_type": "dual_degree", "flags": []},
        "profile": {
            "student_name_zh": "王小明",
            "student_id": "B1234567",
            "college": "管理學院",
            "department": "資訊管理學系",
            "degree_level": "大學部",
        },
        "answers": {
            "dual-enrollment": {
                "overseas_institution": "Partner University",
                "dual_degree_program": "Dual Degree Program",
                "planned_start_date": "2026-09-01",
                "planned_end_date": "2027-06-30",
                "admission_notice": [
                    {"name": "admission.pdf", "type": "application/pdf", "data_url": pdf_data_url},
                    {"name": "scan.png", "type": "image/png", "data_url": _png_data_url()},
                ],
            }
        },
        "selections": [{"template_id": "dual-enrollment", "copies": 1, "reviewed_version": "2026-07-16"}],
    }
    _review_current_content(client, payload)

    response = client.post("/application-center/docx", json=payload)

    assert response.status_code == 200
    assert response.headers["x-application-packet-attachments"] == "2"
    document = WordDocument(BytesIO(response.content))
    assert len(document.inline_shapes) == 1
    text = _word_text(response.content)
    assert "請另行附上的 PDF 附件" in text
    assert "admission.pdf（1 頁）" in text
