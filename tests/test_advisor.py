from __future__ import annotations

import hashlib
from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from exchange_info.advisor import _dashboard_deadlines, answer_question, build_advisor_plan
from exchange_info.api import create_app
from exchange_info.document_checker import check_document
from exchange_info.models import (
    AccountSessionRecord,
    ApplicationPeriodAudience,
    ApplicationPeriodDateStatus,
    ApplicationPeriodItem,
    ApplicationPeriodStatus,
    AdvisorProfile,
    PageCategory,
    PartnerStudentSummary,
    PartnerUniversity,
    ProcessedPage,
    SourceReference,
    StudentDashboard,
    StudentInfoItem,
    UserAccount,
    utc_now,
)
from exchange_info.storage import save_account_sessions, save_user_accounts


SOURCE_URL = "https://www.cgu.edu.tw/oia-ch/Subject/Detail/123"


def _authenticated_client(data_dir) -> TestClient:
    account = UserAccount(display_name="文件檢查測試學生", student_status_verified=True,
                          verified_university="長庚大學", school_email_verified=True)
    token = "advisor-document-check-token"
    save_user_accounts([account], data_dir)
    save_account_sessions(
        [
            AccountSessionRecord(
                account_id=account.id,
                token_hash=hashlib.sha256(token.encode("utf-8")).hexdigest(),
                expires_at=utc_now() + timedelta(days=1),
            )
        ],
        data_dir,
    )
    client = TestClient(create_app(data_dir))
    client.headers.update({"Authorization": f"Bearer {token}"})
    return client


def _dashboard() -> StudentDashboard:
    now = datetime.now(timezone.utc)
    deadline = now + timedelta(days=21)
    return StudentDashboard(
        school_name="長庚大學",
        academic_score_minimum=70,
        language_score_minimums={"toeic": 700},
        application_periods=[
            ApplicationPeriodItem(
                title="交換學生校內申請",
                details=[f"截止日期：{deadline.year}年{deadline.month}月{deadline.day}日"],
                source=SourceReference(title="交換公告", source_url=SOURCE_URL),
                scope="home_application",
                program_types=["exchange"],
                academic_year="116",
                application_round="第一梯次",
                audience=ApplicationPeriodAudience.ALL_STUDENTS,
                opens_at=now - timedelta(days=1),
                deadline_at=deadline,
                date_status=ApplicationPeriodDateStatus.VERIFIED_EXACT,
                status=ApplicationPeriodStatus.OPEN,
                status_note="本梯次正在受理。",
                deadline_note="官方校內申請截止時間。",
                reminder_eligible=True,
                checked_at=now,
            )
        ],
        eligibility=[
            StudentInfoItem(
                title="校內申請資格",
                details=["學業平均成績須達 70 分以上。"],
                source=SourceReference(title="交換公告", source_url=SOURCE_URL),
            )
        ],
        language_requirements=[
            StudentInfoItem(
                title="語言門檻",
                details=["TOEIC 700 分以上。"],
                source=SourceReference(title="交換公告", source_url=SOURCE_URL),
            )
        ],
        required_documents=[
            StudentInfoItem(
                title="申請文件",
                details=["申請表、成績單、讀書計畫與語言證明。"],
                source=SourceReference(title="交換公告", source_url=SOURCE_URL),
            )
        ],
        partner_universities=[
            PartnerUniversity(
                program_type="交換學生",
                country="日本",
                school_name="測試日本大學",
                colleges="全校",
                degree_levels="大學部／研究所",
                eligibility="TOEIC 700",
                application_deadline=f"{deadline.month}月{deadline.day}日",
                source_url=SOURCE_URL,
                image_url="http://testserver/static/assets/schools/test.jpg",
                student_summary_zh=PartnerStudentSummary(
                    eligibility_points=["TOEIC 700 分以上"],
                ),
            ),
            PartnerUniversity(
                program_type="交換學生",
                country="法國",
                school_name="測試法國大學",
                colleges="醫學院",
                degree_levels="大學部",
                eligibility="IELTS 7.0",
                source_url="https://example.edu/france",
            ),
        ],
    )


def _page() -> ProcessedPage:
    return ProcessedPage(
        title="交換學生申請公告",
        source_url=SOURCE_URL,
        text="申請文件包含申請表、成績單、讀書計畫與語言證明。",
        summary="交換申請文件與截止日期整理。",
        category=PageCategory.EXCHANGE_APPLICATION,
        content_hash="advisor-test",
    )


def test_advisor_returns_grounded_match_tasks_and_no_fake_admission_probability(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    plan = build_advisor_plan(
        AdvisorProfile(
            college="工學院",
            degree_level="大學部",
            academic_score=82,
            language_test="toeic",
            language_score=850,
            preferred_country="日本",
            target_term="2027-fall",
            question="交換申請需要哪些文件？",
        ),
        _dashboard(),
        [_page()],
        [],
        [],
    )

    assert [item.school_name for item in plan.recommendations] == ["測試日本大學"], "only the country the student chose"
    assert plan.recommendations[0].country == "日本"
    assert plan.answer and plan.answer.sources[0].source_url == SOURCE_URL
    assert any(task.due_date for task in plan.tasks)
    assert "admission_probability" not in plan.model_dump()
    assert "不是錄取機率" in plan.disclaimer


def test_advisor_does_not_turn_partner_reference_text_into_a_future_deadline() -> None:
    dashboard = StudentDashboard(
        school_name="長庚大學",
        partner_universities=[
            PartnerUniversity(
                program_type="交換學生",
                country="法國",
                school_name="參考時程測試大學",
                application_deadline="秋季 10 月 30 日；適用年度待確認",
                source_url=SOURCE_URL,
            )
        ],
    )

    assert _dashboard_deadlines(dashboard) == []


def test_advisor_unknown_scope_and_threshold_are_not_passed_qualification():
    dashboard = _dashboard()
    unknown = dashboard.partner_universities[0].model_copy(update={
        "colleges": "依當期校際協議確認", "degree_levels": None,
    })
    dashboard = dashboard.model_copy(update={"partner_universities": [unknown], "academic_score_minimum": None})
    plan = build_advisor_plan(AdvisorProfile(college="工學院", degree_level="大學部", academic_score=88,
                                           language_test="toeic", language_score=900), dashboard, [], [], [])
    match = plan.recommendations[0]
    assert match.eligibility_status == "needs_confirmation"
    assert match.blocking_reasons == []
    assert any("學院範圍" in item for item in match.confirmation_items)
    assert "符合學院範圍" not in match.reasons
    assert not any("達長庚" in item for item in match.reasons)
    assert any("門檻待確認" in item for item in plan.qualification)


def test_advisor_misplaced_degree_in_college_field_is_unknown_not_mismatch():
    dashboard = _dashboard()
    partner = dashboard.partner_universities[0].model_copy(update={"colleges":"大學部", "degree_levels":"大學部"})
    plan = build_advisor_plan(AdvisorProfile(college="法律學院", degree_level="大學部", academic_score=88,
                                           language_test="toeic", language_score=900),
                              dashboard.model_copy(update={"partner_universities":[partner]}), [], [], [])
    assert plan.recommendations[0].eligibility_status == "needs_confirmation"
    assert plan.recommendations[0].blocking_reasons == []


def test_advisor_undergraduate_is_not_a_graduate_match():
    dashboard = _dashboard()
    partner = dashboard.partner_universities[0].model_copy(update={"degree_levels":"Undergraduate"})
    plan = build_advisor_plan(AdvisorProfile(college="工學院", degree_level="研究所", academic_score=88,
                                           language_test="toeic", language_score=900),
                              dashboard.model_copy(update={"partner_universities":[partner]}), [], [], [])
    assert plan.recommendations[0].eligibility_status == "does_not_match"
    assert any("學制" in item for item in plan.recommendations[0].blocking_reasons)


def test_advisor_eligibility_precedes_preference_score_and_keeps_blocking_reasons():
    dashboard = _dashboard()
    allowed = dashboard.partner_universities[0]
    blocked = allowed.model_copy(update={"school_name":"不適用但偏好較高的大學", "colleges":"醫學院", "country":"法國"})
    both = dashboard.model_copy(update={"partner_universities":[blocked,allowed]})
    profile = dict(college="工學院", degree_level="大學部", academic_score=88, language_test="toeic", language_score=900,
                   target_school=blocked.school_name)
    # With no country chosen, a school the student cannot apply to never outranks one they can.
    plan = build_advisor_plan(AdvisorProfile(**profile), both, [], [], [])
    assert plan.recommendations[0].school_name == allowed.school_name
    assert plan.recommendations[0].eligibility_status == "matches_known"
    assert plan.target_match.school_name == blocked.school_name
    assert plan.target_match.eligibility_status == "does_not_match"
    assert any("你的學院" in item for item in plan.target_match.blocking_reasons)
    # With a country chosen, only that country is listed, still marked as not applicable rather than promoted.
    plan = build_advisor_plan(AdvisorProfile(**profile, preferred_country="法國"), both, [], [], [])
    assert [item.school_name for item in plan.recommendations] == [blocked.school_name]
    assert plan.recommendations[0].eligibility_status == "does_not_match"
    assert any("你的學院" in item for item in plan.recommendations[0].blocking_reasons)
    assert plan.target_match.school_name == blocked.school_name


def test_advisor_lists_only_the_country_the_student_chose_and_says_so():
    dashboard = _dashboard()
    japan = dashboard.partner_universities[0]
    germany = japan.model_copy(update={"school_name": "測試德國大學", "country": "德國"})
    china = japan.model_copy(update={"school_name": "測試大陸大學", "country": "大陸地區"})
    everything = dashboard.model_copy(update={"partner_universities": [germany, japan, china]})
    profile = dict(college="工學院", degree_level="大學部", academic_score=88, language_test="toeic", language_score=900)

    everywhere = build_advisor_plan(AdvisorProfile(**profile), everything, [], [], [])
    assert {item.country for item in everywhere.recommendations} == {"德國", "日本", "大陸地區"}
    assert "只列出" not in everywhere.summary

    chosen = build_advisor_plan(AdvisorProfile(**profile, preferred_country="大陸地區"), everything, [], [], [])
    assert [item.school_name for item in chosen.recommendations] == ["測試大陸大學"], "no German school for a mainland request"
    assert chosen.summary.startswith("已依你選的「大陸地區」，只列出該地區的 1 所開放合作校")
    assert "不限國家" in chosen.summary

    absent = build_advisor_plan(AdvisorProfile(**profile, preferred_country="巴西"), everything, [], [], [])
    assert {item.country for item in absent.recommendations} == {"德國", "日本", "大陸地區"}
    assert absent.summary.startswith("目前沒有「巴西」的開放合作校")


def test_advisor_country_choice_ignores_spacing_and_still_ranks_by_eligibility_inside_the_country():
    dashboard = _dashboard()
    japan = dashboard.partner_universities[0]
    blocked = japan.model_copy(update={"school_name": "日本醫學大學", "colleges": "醫學院"})
    fine = japan.model_copy(update={"school_name": "日本工學大學"})
    plan = build_advisor_plan(
        AdvisorProfile(college="工學院", degree_level="大學部", academic_score=88, preferred_country=" 日本 "),
        dashboard.model_copy(update={"partner_universities": [blocked, fine]}), [], [], [],
    )
    assert [item.school_name for item in plan.recommendations] == ["日本工學大學", "日本醫學大學"]
    assert plan.recommendations[1].eligibility_status == "does_not_match"


def test_advisor_home_language_requirement_cannot_be_offset_by_preferences():
    dashboard = _dashboard()
    partner = dashboard.partner_universities[0].model_copy(update={"eligibility":"TOEIC 600", "student_summary_zh":PartnerStudentSummary()})
    plan = build_advisor_plan(AdvisorProfile(college="工學院", degree_level="大學部", academic_score=88,
                                           language_test="toeic", language_score=650, preferred_country="日本"),
                              dashboard.model_copy(update={"partner_universities":[partner]}), [], [], [])
    assert plan.recommendations[0].eligibility_status == "does_not_match"
    assert any("本校公開門檻" in item for item in plan.recommendations[0].blocking_reasons)


def test_advisor_unknown_language_remains_explorable_and_target_outside_top_twenty_is_compared():
    dashboard = _dashboard()
    partner = dashboard.partner_universities[0]
    partners = [partner.model_copy(update={"school_name":f"學校 {index:02}"}) for index in range(25)]
    plan = build_advisor_plan(AdvisorProfile(college="工學院", degree_level="大學部", academic_score=88,
                                           target_school=partners[-1].school_name),
                              dashboard.model_copy(update={"partner_universities":partners}), [], [], [])
    assert len(plan.recommendations) == 20
    assert plan.target_match.school_name == partners[-1].school_name
    assert plan.target_match.eligibility_status == "needs_confirmation"
    assert plan.target_match.blocking_reasons == []


def test_grounded_answer_refuses_when_no_relevant_public_source(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    answer = answer_question("火星寵物檢疫與火箭燃料規定", _dashboard(), [], [], [])

    assert answer.evidence_status == "not_found"
    assert answer.sources == []
    assert "資料不足" in answer.answer


def test_initial_application_document_question_does_not_mix_incoming_or_post_acceptance_files(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    application_page = _page().model_copy(
        update={
            "text": "申請學生須填寫申請表並繳送紙本文件至國際處，完整清單依當期公告。",
            "summary": "校內交換申請入口。",
        }
    )
    incoming_page = _page().model_copy(
        update={
            "title": "簽證資訊",
            "source_url": "https://www.cgu.edu.tw/oia-ch/incoming-visa",
            "text": "交換學生在台就學須檢附護照、健康檢查、財力證明與入學許可。",
            "summary": "來台交換學生簽證文件。",
        }
    )

    answer = answer_question(
        "交換申請需要準備哪些文件？",
        StudentDashboard(school_name="長庚大學"),
        [application_page, incoming_page],
        [],
        [],
    )

    assert answer.evidence_status == "limited"
    assert "尚未列出完整" in answer.answer
    assert "護照" not in "".join(answer.key_points)
    assert all("incoming-visa" not in source.source_url for source in answer.sources)


def test_document_checker_reads_docx_and_flags_missing_study_plan_sections(monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    document = Document()
    document.add_heading("交換讀書計畫", level=1)
    document.add_paragraph("我希望到日本交換，申請動機是拓展國際視野。" * 12)
    output = BytesIO()
    document.save(output)

    result = check_document(
        file_name="讀書計畫.docx",
        content=output.getvalue(),
        document_type="auto",
        require_signature=True,
    )

    assert result.document_type == "study_plan"
    assert result.character_count > 100
    assert any(issue.title == "未找到修課規劃" for issue in result.issues)
    assert any(issue.code == "signature_missing" for issue in result.issues)
    assert result.ai_enhanced is False
    assert "不會保存原檔" in result.privacy_note


def test_advisor_and_document_check_api(tmp_path, monkeypatch) -> None:
    dashboard = _dashboard()
    monkeypatch.setattr(
        "exchange_info.api.build_student_dashboard",
        lambda pages, info, **kwargs: dashboard,
    )
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    client = _authenticated_client(tmp_path / "data")

    plan = client.post(
        "/advisor/plan",
        json={
            "college": "工學院",
            "degree_level": "大學部",
            "academic_score": 82,
            "language_test": "toeic",
            "language_score": 850,
            "preferred_country": "日本",
        },
    )
    assert plan.status_code == 200
    assert plan.json()["recommendations"][0]["school_name"] == "測試日本大學"

    checked = client.post(
        "/document-check",
        files={"file": ("履歷.txt", "教育背景：長庚大學。\n技能：英文。\n經歷：交換社團。" * 8, "text/plain")},
        data={
            "document_type": "resume",
            "require_signature": "false",
            "consent": "true",
        },
    )
    assert checked.status_code == 200
    assert checked.json()["file_name"] == "履歷.txt"
    assert checked.json()["document_type"] == "resume"
