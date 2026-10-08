"""Fetch only the reviewed public CCU source manifest into a candidate directory.

This script does not overwrite reviewed evidence, publish, contact schools, or
authenticate. A changed file remains a candidate until its contents/year/scope
are reviewed. Download success is deliberately distinct from current applicability.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import urljoin, urlsplit
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import requests
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from exchange_info.ccu_sources import BUNDLE


def public_spreadsheet_pages(body: bytes) -> list[str] | None:
    """Read public XLSX cell values; never evaluate formulas or external links."""
    archive = zipfile.ZipFile(io.BytesIO(body))
    if "xl/workbook.xml" not in archive.namelist():
        return None
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    strings = []
    if "xl/sharedStrings.xml" in archive.namelist():
        root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
        strings = ["".join(item.itertext()) for item in root.findall("x:si", ns)]
    pages = []
    sheets = sorted(
        name
        for name in archive.namelist()
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
    )
    for name in sheets:
        lines = [name]
        root = ElementTree.fromstring(archive.read(name))
        for row in root.findall(".//x:sheetData/x:row", ns):
            values = []
            for cell in row.findall("x:c", ns):
                value = cell.find("x:v", ns)
                inline = cell.find("x:is", ns)
                formula = cell.find("x:f", ns)
                content = (
                    (value.text or "")
                    if value is not None
                    else "".join(inline.itertext()) if inline is not None else ""
                )
                if cell.get("t") == "s" and content:
                    content = strings[int(content)]
                if formula is not None:
                    content = "FORMULA (not evaluated): " + (formula.text or "")
                if content:
                    values.append(str(cell.get("r", "")) + ": " + content)
            if values:
                lines.append(" | ".join(values))
        pages.append("\n".join(lines))
    return pages


def public_content_blocks(soup):
    """Read FHNW public display data without evaluating website JavaScript."""

    def text_nodes(value):
        if isinstance(value, dict):
            if isinstance(value.get("text"), str):
                yield value["text"]
            for key, child in value.items():
                if key not in {"text", "plaintext"}:
                    yield from text_nodes(child)
        elif isinstance(value, list):
            for child in value:
                yield from text_nodes(child)

    for script in soup.find_all("script"):
        raw = script.string or ""
        if not raw.startswith("window.__data="):
            continue
        for match in re.finditer(r'"content"\s*:', raw):
            try:
                content, _ = json.JSONDecoder().raw_decode(raw[match.end() :].lstrip())
            except json.JSONDecodeError:
                continue
            if not isinstance(content, dict) or not isinstance(
                content.get("data"), dict
            ):
                continue
            blocks = content["data"].get("blocks", {})
            if not isinstance(blocks, dict):
                continue
            for block in blocks.values():
                if not isinstance(block, dict) or block.get("@type") != "accordion":
                    continue
                data = block.get("data")
                if not isinstance(data, dict):
                    continue
                panels = data.get("blocks", {})
                if not isinstance(panels, dict):
                    continue
                for panel in panels.values():
                    if isinstance(panel, dict):
                        yield (
                            str(panel.get("title", ""))
                            + "\n"
                            + " ".join(text_nodes(panel))
                        )


def fetch_public_source(item: dict, output: Path, session=None) -> dict:
    url = item["url"]
    if urlsplit(url).scheme not in {"http", "https"}:
        raise ValueError("CCU source URLs must be public HTTP(S)")
    result = {
        **item,
        "status": "fetch_failed",
        "content_applicability": "needs_confirmation",
        "checked_at": datetime.now(UTC).isoformat(),
    }
    target = url
    match = re.search(r"drive\.google\.com/file/d/([^/?]+)", url)
    if match and item["authority"] == "partner_issued_document_linked_by_ccu":
        target = f"https://drive.usercontent.google.com/download?id={match[1]}&export=download"
    client = session or requests.Session()
    try:
        response = client.get(target, timeout=(10, 35))
        result.update(
            requested_url=target,
            final_url=response.url,
            http_status=response.status_code,
            content_type=response.headers.get("Content-Type"),
            bytes=len(response.content),
        )
        response.raise_for_status()
        body = response.content
        if len(body) > 32 * 1024 * 1024:
            raise ValueError(
                "Source exceeds 32 MiB; inspect separately before extraction"
            )
        result["sha256"] = hashlib.sha256(body).hexdigest()
        if body.startswith(b"%PDF-"):
            result.update(
                status="pdf",
                text_pages=[
                    page.extract_text() for page in PdfReader(io.BytesIO(body)).pages
                ],
            )
            (output / (item["source_id"] + ".pdf")).write_bytes(body)
        elif body.startswith(b"PK"):
            spreadsheet = public_spreadsheet_pages(body)
            if spreadsheet is not None:
                result.update(status="xlsx", text_pages=spreadsheet)
                (output / (item["source_id"] + ".xlsx")).write_bytes(body)
            else:
                doc = Document(io.BytesIO(body))
                result.update(
                    status="docx",
                    text_pages=[
                        "\n".join(
                            [paragraph.text for paragraph in doc.paragraphs]
                            + [
                                " | ".join(cell.text for cell in row.cells)
                                for table in doc.tables
                                for row in table.rows
                            ]
                        )
                    ],
                )
                (output / (item["source_id"] + ".docx")).write_bytes(body)
        elif body.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            if urlsplit(response.url).path.lower().endswith(".xls") or (
                response.headers.get("Content-Type", "").split(";")[0]
                == "application/vnd.ms-excel"
            ):
                (output / (item["source_id"] + ".xls")).write_bytes(body)
                result.update(status="downloaded_xls_unextracted", text_pages=[])
                try:
                    import xlrd
                except ImportError:
                    result["extraction_error"] = (
                        "Install scripts/requirements-ccu-research.txt to read legacy XLS"
                    )
                else:
                    book = xlrd.open_workbook(file_contents=body)
                    pages = []
                    for sheet in book.sheets():
                        lines = [sheet.name]
                        for row in range(sheet.nrows):
                            values = [
                                f"R{row + 1}C{column + 1}: {sheet.cell_value(row, column)}"
                                for column in range(sheet.ncols)
                                if sheet.cell_value(row, column) != ""
                            ]
                            if values:
                                lines.append(" | ".join(values))
                        pages.append("\n".join(lines))
                    result.update(
                        status="xls",
                        text_pages=pages,
                        spreadsheet_extraction_note=(
                            "xlrd reads stored cell values, including cached formula results; "
                            "it does not execute formulas, macros or external links"
                        ),
                    )
            else:
                # A legacy Word binary is not HTML. Preserve it for a real DOC
                # converter rather than treating binary strings as evidence.
                result.update(status="downloaded_doc_unextracted", text_pages=[])
                (output / (item["source_id"] + ".doc")).write_bytes(body)
        else:
            soup = BeautifulSoup(body, "html.parser")
            public_blocks = list(public_content_blocks(soup))
            result["links"] = [
                {
                    "label": a.get_text(" ", strip=True),
                    "url": urljoin(response.url, a["href"]),
                }
                for a in soup.select("a[href]")
                if urlsplit(urljoin(response.url, a["href"])).scheme
                in {"http", "https"}
            ]
            for node in soup(["script", "style", "nav", "footer", "header"]):
                node.decompose()
            # Google Sites can keep a public timetable in an escaped data-code
            # attribute. Read the supplied HTML, without executing its scripts.
            embeds = []
            for node in soup.select("[data-code]"):
                code = node.get("data-code", "")
                embedded = BeautifulSoup(code, "html.parser")
                for hidden in embedded(["script", "style"]):
                    hidden.decompose()
                text = embedded.get_text("\n", strip=True)
                if text and text not in embeds:
                    embeds.append(text)
            result.update(
                status="html",
                text_pages=[soup.get_text("\n", strip=True)] + embeds + public_blocks,
            )
            if embeds:
                result["embedded_content_count"] = len(embeds)
                result["extraction_method"] = (
                    "html_body_and_public_data_code_attributes_no_script_execution"
                )
            if public_blocks:
                result["public_accordion_count"] = len(public_blocks)
                result["extraction_method"] = (
                    "html_body_and_public_content_blocks_no_script_execution"
                )
        result["text_sha256"] = hashlib.sha256(
            json.dumps(result["text_pages"], ensure_ascii=False).encode()
        ).hexdigest()
        result["review_note"] = (
            "Fetch succeeded; applicable year, round, program and field semantics still require review."
        )
    except Exception as exc:  # noqa: BLE001
        # Preserve arbitrary parser/network failures as candidate acquisition
        # records; a single malformed public document must not abort the batch.
        result.update(status="fetch_failed", error=f"{type(exc).__name__}: {exc}")
    (output / (item["source_id"] + ".json")).write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    )
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="Candidate directory, outside the reviewed bundle",
    )
    parser.add_argument("--workers", type=int, default=4)
    args = parser.parse_args()
    if args.output.resolve().is_relative_to(BUNDLE.resolve()):
        parser.error(
            "Output must be outside the reviewed bundle; review candidates before integration"
        )
    args.output.mkdir(parents=True, exist_ok=True)
    items = json.loads((BUNDLE / "overseas-source-requests.json").read_text())
    results = []
    with ThreadPoolExecutor(max_workers=max(1, min(args.workers, 6))) as pool:
        pending = [
            pool.submit(fetch_public_source, item, args.output) for item in items
        ]
        for future in as_completed(pending):
            result = future.result()
            results.append(result)
            print(
                result["source_id"],
                result["status"],
                result.get("http_status"),
                flush=True,
            )
    (args.output / "acquisition.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2) + "\n"
    )
    print(f"Fetched {len(results)} sources; no reviewed catalogue was modified.")


if __name__ == "__main__":
    main()
