"""Reproduce the seventh CCU public-source reading batch, without network calls.

An index, outgoing rule, dated reference and current incoming rule have separate
reading dispositions. Reading a wrapper does not establish its linked policy.
"""
from __future__ import annotations

import gzip
import json
import re
from copy import deepcopy

from bs4 import BeautifulSoup
from docx import Document
from fetch_ccu_completion_sources import BUNDLE as B
from pypdf import PdfReader
from review_ccu_followup_batch import evidence, fact, receipt

PRIORITY_NOTES = {
    "p21-reference-folder-3": ("unversioned_template", "已解析完整OLE Word體檢表：醫師簽名、日期、檢查單位章、胸部X光等空格；無適用學年，不據此要求當期全體必繳。"),
    "p34-campus_contact-1": ("campus_index", "Higashiyama／Tsurumai樓層圖與2026/27 prospectus入口；不是提名資格或申請期限。"),
    "p34-life-1": ("service_index", "Student Life住宿／生活／校園入口；子頁規則未由此索引核實。"),
    "p43-life-1": ("undated_incoming_reference", "逐讀各校區性別房型、每期宿舍申請與TB證明；每年1月／7月窗口無2027確切日，不能同一校區房型套全校。"),
    "p43-life-2": ("undated_insurance_reference", "NHIS＋私人保險段及Student Medical Mutual Aid對象已讀；Exchange/Visiting不適用互助費，不拿3000KRW當交換必繳。"),
    "p47-campus_contact-1": ("campus_index", "只取得Campus Tour導覽入口，未取得Flash VR政策。"),
    "p47-campus_contact-2": ("campus_reference", "Seoul校園設施正文已讀；OIAM房號101／102內文不同，沒有以樓棟位置推出接待資格。"),
    "p47-host_entry-1": ("programme_index", "Pre-Arrival Guide導航索引，未從頁首推出當期全套政策。"),
    "p47-host_entry-2": ("programme_index", "Overview索引，年度及下層政策須由各實際正文核實。"),
    "p47-life-2": ("undated_incoming_reference", "Globee及International Houses由OIAM管理，入住分配通常學期前約一月email；無名額／房型保證或確切2027期限。"),
    "p48-host_entry-2": ("outgoing_not_ccu_incoming", "SeoulTech本校海外派出GPA／学期／學分及1/31、7/31條件，方向相反；不填CCU接待資格。"),
    "p49-life-1": ("superseded_undated_reference", "舊頁NHIS或私人保險及舊保費與Spring2027簡章兩種保險不同；現行PDF保留優先，舊頁不降級當期政策。"),
    "p53-host_entry-1": ("outgoing_not_ccu_incoming", "忠南本校海外派出3GPA與甄選配分，非CCU incoming。"),
    "p53-host_entry-2": ("outgoing_not_ccu_incoming", "Global Talent本校派出獎學金、學費80%條件，非CCU入學或獎金。"),
    "p53-life-2": ("historical_notice_index", "2023減費、2022 NHIS兩筆公告索引；未據索引查核2027保費。"),
    "p57-life-1": ("fee_index_distinct_population", "宿費PDF分研究生與本科／預科degree新生，不能直接套交換。實際兩PDF新抓取502，未取得費率正文。"),
    "p61-host_entry-1": ("undated_incoming_and_outgoing_separated", "全頁已讀，incoming文件／至少3課或9學分與outgoing流程分開；SoA+D portfolio、Engineering課程描述已在接待文件欄保留。"),
    "p63-life-1": ("undated_kl_housing_reference", "KL Blocks D/E雙人及有限E單人房、Angkasa Condos外宿，不能套Sarawak／Springhill或保證床位。"),
    "p66-official-academic-calendar": ("sheet_visible_grid", "公開Sheet顯示2026/27及已記錄異常年份；未解除第三學期矛盾，沒有用檔名推日期。"),
    "p66-official-accommodation-list": ("wrapper_body_not_document", "Google Docs外殼可讀但未取得文件正文；既有公開export正文另有来源。"),
    "p66-official-factsheet-shortlink": ("wrapper_body_not_document", "Drive短鏈外殼，未在此外殼讀到Fact Sheet規則；PDF另有已讀來源。"),
    "p69-reference-link-2": ("folder_listing_not_document", "資料夾列FQA PDF修改2023-08-15；本次讀列表而非PDF正文，檔名及修改日不是適用年度證明。"),
    "p79-host_entry-1": ("conference_not_exchange", "IEEE ICIDeA2026會議10/29–31及投稿資訊，不能當交換期限或資格。"),
    "p89-host_entry-1": ("exchange_index", "Exchange／Guest與Bachelor／Master入口，頁面更新2026-01-21；個別條款由實際子頁核查。"),
    "p89-life-1": ("undated_incoming_housing", "讀完整住房政策：exchange五城市保障、11/1 23:59CET、一offer、6日接受、14日cooling；頁未標2027，不生成當期提醒。"),
    "p91-host_entry-1": ("erasmus_charter_not_bilateral_rules", "三頁Erasmus Charter2021–2027是政策框架及電子收件頁，不是CCU雙邊2027接待簡章；免費條款不擴及任意交換。"),
    "p92-public-factsheet": ("factsheet_index", "列2026/27英西文PDF，實際PDF已有獨立來源；索引本身無完整申請規則。"),
    "p92-public-incoming": ("undated_incoming_overview", "一般一學期或一年、先聯絡母校國際處；全學位與交換分開，沒有新增數字GPA。"),
    "p92-public-mobility": ("mobility_index", "incoming／outgoing入口分開，西文母校派出介紹不作CCU接待規則。"),
    "p99-host_entry-2": ("event_not_exchange", "2026-10-17語言tandem單場活動，非2027提名／申請期限。"),
    "p104-host_entry-1": ("undated_business_exchange_index", "School of Business英文入口，跨校區／本科及一至兩學期介绍；日期標題不是申請規則正文。"),
    "p104-host_entry-2": ("undated_business_exchange_index", "德文同學院交換介绍與事實表入口；未據下層標題補住房期限／學費。"),
    "p107-life-2": ("service_index", "住宿／餐飲短索引、地址與服務入口，非完整租約或個人住房申請條件。"),
    "p108-life-2": ("undated_housing_services", "單雙三人及無障礙房、步行10分／網路一般服務；無2027床位、期限或費率保證。"),
    "p112-public-incoming-erasmus": ("current_incoming_index", "正文Spring2026/27提名2026-10-15、B1/B2及申請入口，與既有當期來源對照；不把提名日當校內期限。"),
    "p115-campus_contact-2": ("current_housing_index", "正文2026/27候補、3500床、shared occupancy；實際申請／費率／calendar子頁另讀，不能保交換床位。"),
}


def load(name):
    return json.loads((B / name).read_text())


def visible(record):
    path = B / record["bundled_file"]
    raw = path.read_bytes()
    raw = gzip.decompress(raw) if path.suffix == ".gz" else raw
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer"]):
        tag.decompose()
    return re.sub(r"\s+", " ", soup.get_text(" ", strip=True)), soup


def main():
    manual = load("completion-manual-fields.json")
    before = load("completion-reviewed-fields.json")["schools"]

    def row(page):
        existing = next((r for r in manual["schools"] if r["ccu_page"] == page), None)
        if existing is None:
            existing = {"ccu_page": page, "host_fields": {}}
            manual["schools"].append(existing)
        return existing

    def html(page, label, quote, year):
        text, _ = visible(receipt(page, label))
        assert quote in text, (page, label, quote)
        ev = evidence(page, label, quote, year=year)
        ev["application_round"] = year
        return ev

    def pdf(label, page, quote):
        record = receipt(84, "seventh-" + label)
        pages = [re.sub(r"\s+", " ", p.extract_text() or "") for p in PdfReader(B / record["bundled_file"]).pages]
        assert quote in pages[page - 1], (label, page, quote)
        ev = evidence(84, "seventh-" + label, quote, pages=[page], year="2026–2027官方incoming頁所連附件；Erasmus／雙邊適用分開，中正本期未開放")
        ev["application_round"] = "Spring2027接待參考；非中正可申請名額"
        return ev

    # UCSI's actual incoming timetable stops at January 2025. Its current portal
    # does not establish a 2027 bilateral agreement or update these dates.
    u = row(63)["host_fields"]
    uy = "頁內入學表僅Jan2024至Jan2025；歷史incoming參考，2027新版尚未核實"
    un = "明載舊年度的接待參考，不是2027當期規則；中正本期未開放。"
    def ue(q):
        return html(63, "followup-inbound", q, uy)
    for key, value, quote in [
        ("program_scope", "UCSI歷史incoming exchange由母校聯絡、甄選及申請；一般課程英文。現行partner頁沒有列中正，未因此判協議失效，2027有效協議待核。", "run entirely by the students’ home universities"),
        ("eligibility", "歷史incoming須母校在學、至少完成一學期、良好成績及身心健康；未載數字GPA或自加年級限制。", "Must have completed at least one semester of studies"),
        ("language", "歷史incoming IELTS5.5／TOEFL500，母校全英授課者可豁免測驗；未列新版iBT換算，不代替中正語言要求。", "IELTS 5.5 and TOEFL 500"),
        ("nomination_deadline", "歷史Jan2025入學：母校提名2024-09-12；表只到Jan2025，2027提名日尚未核實。", "12 September 2024"),
        ("application_deadline", "歷史Jan2025入學：文件2024-10-03；不是2027校外期限，不產生當期提醒。", "3 October 2024"),
        ("course_restrictions", "歷史incoming課程選擇於每學期前三週；具體2027開課、校區／專業先修及名額待核，不以全英授課保證所有課可選。", "First three (3) weeks of each semester"),
        ("exchange_duration", "歷史incoming可交換一或兩學期；有效協議／當期接待期仍需核實。", "One semester or two semesters"),
    ]:
        u[key] = fact(value, [ue(quote)], status="reference_only", reason=un)
    u["housing"] = fact("KL校區Blocks D/E一般校內宿舍主要雙人，E有限單人套房；外宿Angkasa Condominiums由UCSI Peterson管理。Sarawak／Bandar Springhill另校區，沒有2027交換床位、房型或租金保證。", [html(63, "life-1", "Angkasa", "未標學年KL一般住宿頁，2027交換適用尚未核實")], status="needs_confirmation", reason="實際住宿正文已讀，限定KL一般服務；沒有當期交換分配證明。")

    # Current UCY text includes non-Erasmus exchanges, but eligibility still
    # depends on an active subject-specific bilateral agreement.
    c = row(84)["host_fields"]
    cy = "2026–2027 incoming正文，Last Updated 2026-10-08；中正本期未開放參考"
    def ce(q):
        return html(84, "seventh-discovered-target-13", q, cy)
    cn = "接待校當期正文已讀；不證明中正有效當期學科協議或校內可申請。"
    c["program_scope"] = fact("2026–2027接待Erasmus與其他有效交換協議均須母校正式提名及相關學科協議；非Erasmus有独立申請／Learning Agreement。Erasmus International KA171的生活交通津貼是另一有協議計畫，不視為中正交換保證補助。", [ce("μέσω άλλων συμφωνιών συνεργασίας"), html(84, "seventh-international-english", "with which an Erasmus International agreement exists", "Erasmus International頁更新2026-09-15；個別協議適用另核")], status="needs_confirmation", reason=cn)
    c["eligibility"] = fact("接待端須母校提名及2026–2027相關學科有效協議，不可自行提名；Computer Science才要求GPA6.5/10。未由接待學科選課彈性撤銷中正歷史不可跨系限制；學制與協議須另核。", [ce("ΔΕΝ ΔΥΝΑΤΑΙ"), pdf("current-language", 1, "A GPA of 6.5/10")], status="needs_confirmation", reason=cn)
    c["language"] = fact("2026/27 Fact Sheet及language附件要求English B2，主要教學為希臘文、其他學系英語課有限。PDF拒一般母校／OLS證明，但可全英讀滿一年學程主任證明或英語母語護照；2026-10-08网页又接受母校B2修課證明，兩來源相反，保留conflict，不自行選較寬或換算數值測驗。", [pdf("current-factsheet", 2, "English B2"), pdf("current-language", 1, "OLS certificates will not be accepted"), ce("παρακολούθησαν μαθήματα γλωσσών μέχρι το επίπεδο Β2")], status="conflict", reason="當期公開頁與當期所連PDF對母校語言證明接受範圍相反，須UCY確認；不同於英文授課學程主任的明示例外。")
    c["nomination_deadline"] = fact("Spring2026/27母校提名2026-10-01，秋及全年2026-05-01；partners使用校方email提供提名平台，學生不可自行提名。這是已過的接待校日，不是中正校內梯次。", [pdf("current-factsheet", 2, "1st October 2026"), ce("1 Οκτωβρίου 2026")], reason=cn)
    c["application_deadline"] = fact("Spring2026/27文件2026-10-20（頁面Spring欄10/20未重複年，2026/27 Fact Sheet明寫2026）；秋及全年2026-05-20。個人平台链接由提名後email提供，不能自提名或推下一年同日。", [pdf("current-factsheet", 2, "20th October 2026")], reason=cn)
    c["documents"] = fact("2026/27 Erasmus列申請表、Learning Agreement、B2/希臘語證明、母校成績、身分證／護照、研究生以往學位；非Erasmus另列Exchange申請表、LA、最近成績、官方英語證明及研究生學位。醫療／簽證文件另階段、依計畫國籍確認，不混全體初申；B2母校證明接受矛盾另存語言欄。", [ce("ΑΝΤΑΛΛΑΓΕΣ ΦΟΙΤΗΤΩΝ ΕΚΤΟΣ ERASMUS+"), ce("Πρόσφατη Αναλυτική βαθμολογία")], reason=cn)
    form = receipt(84, "seventh-current-bilateral-form")
    doc = Document(B / form["bundled_file"])
    dt = " ".join(p.text for p in doc.paragraphs) + " ".join(cell.text for table in doc.tables for rr in table.rows for cell in rr.cells)
    assert "ACADEMIC YEAR 20__ - 20__" in dt and "Signature of the Student" in dt
    fe = evidence(84, "seventh-current-bilateral-form", "ACADEMIC YEAR 20__ - 20__", year="2026/27公開頁所連非Erasmus DOCX；模板學年留白、版本日期未標")
    fe["application_round"] = "非Erasmus Exchange模板參考；不推定中正協議適用"
    fe["location_in_source"] = "DOCX paragraph heading and SIGNATURES section; no fixed pagination asserted"
    c["application_forms"] = fact("已打開現行2026/27頁所連非Erasmus Exchange DOCX：學年20__–20__留白、個人／母校／學制1st/2nd/3rd／接待學科／期間／語言／工作經驗；學生、母校department與International Office、接待department與Mobility Office簽名日期。模板版本與CCU當期適用未核；Erasmus個人表則提名後寄链接，兩路徑分開。", [fe, ce("ο σύνδεσμος θα αποστέλλεται")], status="needs_confirmation", reason="無年度模板已讀且可開啟；不能把URL2025/03視為申請學年，也不證明中正本期開放。")
    c["academic_calendar"] = fact("2026/27 Spring：註冊2027-01-13至15、上課01-18至04-16、期末05-10至24；語言中心考試另04-19至22。選課01-22、移除02-05、withdraw03-05；Summer另期，不擴入春季交換。", [pdf("current-calendar", 1, "ACADEMIC CALENDAR 2026 - 2027"), pdf("current-factsheet", 2, "13-15/01/2027"), pdf("current-calendar", 1, "April 19-22, 2027")], reason=cn)
    c["course_restrictions"] = fact("當期Fact Sheet平均30ECTS、無re-sit。language附件的Erasmus可同School組合並跨School1–2課但多數須協議department；非Erasmus未證同等跨系權，中正歷史禁止跨系／本科修碩士保留另核。2026/27冬春課程列表標TBA僅此子項，不宣稱所有接待規定尚未公告。", [pdf("current-factsheet", 2, "30 ECTS"), pdf("current-language", 2, "1 -2 courses from other Schools"), ce("Εαρινό εξάμηνο 2026-2027 (TBA)")], status="needs_confirmation", reason="學科協議／Erasmus與CCU歷史限制不同範圍；具體課程未核實。")
    c["housing"] = fact("2026/27 incoming頁：床位有限，申請超額依先後順序；沒有交換生必得房或特定宿舍保證。租金未見此頁載明。", [ce("περιορισμένος αριθμός δωματίων")], reason=cn)
    c["housing_deadline"] = fact("Spring2026/27住房申請2026-11-15，冬／秋2026-06-15；獨立於提名10/1及文件10/20。名额先後不等保證床位，不推出2027秋新版日。", [ce("15η Νοεμβρίου 2026")], reason=cn)
    c["costs"] = fact("2026/27 Fact Sheet估生活EUR700–800/月，已含餐／住宿／misc；不是實際宿舍報價，不能再把宿租重複加上。Erasmus段Any other fees No不等CCU任何雙邊及visa費全免；個人許可／醫療等另程序待核。", [pdf("current-factsheet", 2, "Approx. 700-800 Euro per month")], reason=cn)
    c["insurance"] = fact("2026/27公開頁所連Erasmus醫療附件：EU/EEA可EHIC；第三國或無EHIC者須母國保險涵蓋Cyprus。非Erasmus及個人居留程序適用、保額與費率待校方核；不把列出的保險公司視為推薦／唯一選擇。", [pdf("current-insurance", 1, "with coverage in Cyprus")], status="needs_confirmation", reason="正文Erasmus條件已讀，CCU個人雙邊／居留保險要求未獨立核實。")
    c["visa"] = fact("現行所連8頁entry附件分Non-EU但在EU就讀、Non-EU且在Non-EU就讀、到校後及續留四階段；台灣／CCU雙邊個案不可直接套Erasmus。附件正文含校方送Migration審查、原件認證／醫療與前後不同費項；無承諾核准、沒有用此未標版本PDF當最新主管機關法律結論。", [pdf("current-entry", 3, "WHO STUDY IN NON-EU COUNTRIES")], status="needs_confirmation", reason="已讀附件全部8頁但Erasmus計畫、國籍認證與主管機關最新要求需個人適用查核；不代送件。")
    c["arrival"] = fact("當期Spring迎新／註冊2027-01-13至15；entry附件說Non-EU Erasmus arrival後7–10日提交after-arrival documents，CCU非Erasmus適用未核，不能稱免簽到校或保證接機。", [pdf("current-factsheet", 2, "13-15/01/2027"), pdf("current-entry", 5, "Within 7-10 days")], status="needs_confirmation", reason=cn)

    # General CSUN housing is a current annual rate, not a Spring-only contract
    # or evidence of free tuition for the separately paid SAC programme.
    s = row(115)["host_fields"]
    sy = "CSUN一般Housing Academic Year2026–2027；SAC/CCU個人交換租約適用另核"
    def se(label, quote):
        return html(115, "seventh-housing-" + label, quote, sy)
    s["housing"] = fact("2026/27 CSUN Park Apartments/Suites一般宿舍候補，約3500床、先到先處理、無交換床位保證；免費加入waitlist，住宿正式申請需CSUN Portal帳號及USD50費。Family Housing為有配偶／小孩等另一資格與USD500押金，不能套所有交換生。", [se("application", "only approximately 3,500 bed spaces"), se("application", "$50 housing application fee")], status="needs_confirmation", reason="一般年度住宿正文已核；CCU歷史自費SAC個人資格、Spring契約與床位尚未核實。")
    s["housing_deadline"] = fact("2026/27新住戶住宿申請2026-03-24開放；頁面建議儘速、先到先處理，現已有候補，沒有查得Spring2027獨立截止。住房calendar的2026-07-17 17:00是Fall首期付款／financial aid核驗，不當Spring申請截止。", [se("calendar", "New Residents: 2026-2027 Academic Year Applications Open"), se("application", "first-come, first-served basis")], status="needs_confirmation", reason="取得開放與處理方式；未見春季獨立締約日，不等官方明示尚未公告。")
    s["costs"] = fact("Park2026/27學年合約2026-08-22至2027-05-15，雙人有廚房總USD11351、無廚房9153、Suite8812（後兩另meal plan）；初始USD50 included表總額，RHA15另收。表USD1256/1011/973是分期款而非月租；晚搬入pro-rate。Spring單期租約／SAC學費及全部生活成本未核，不用全年價格除二作報價。", [se("rates", "Total Housing Costs"), se("rates", "This is an installment payment plan, not monthly rent.")], status="needs_confirmation", reason="一般2026/27全年費表已讀；交換／SAC春季實際租約與program tuition分開待核。")

    d = row(89)
    d["expose_host_overrides"] = True
    dy = "現行SDU一般exchange住房正文；未標Spring2027適用學年"
    de = html(89, "life-1", "Vejle", dy)
    old = deepcopy(next(r for r in before if r["ccu_page"] == 89)["host_fields"])
    for key in ["housing", "housing_deadline"]:
        field = d["host_fields"].setdefault(key, old[key])
        field["value"] = field["value"].split(" 本次完整住房頁補核：")[0] + (
            " 本次完整住房頁補核：Exchange保障城市為Odense／Kolding／Esbjerg／Vejle／Sønderborg，限按期申請；Sønderborg full-degree不保障。僅一offer、非首選保證；中央頁一般6日接受及14日cooling，部分城市welcome letter／7日另制不統一。"
            if key == "housing" else
            " 本次完整住房頁補核：中央頁住房5/1及11/1為23:59CET，未標2027；此為保住房申請門檻，不等每校區個人合約／付款期限或EU學生申請日。原城市個別6／7日接受差異保留。")
        field["evidence"] = [e for e in field["evidence"] if e.get("source_id") != de["source_id"]] + [{**de, "status": "needs_confirmation"}]

    # Resolve nomination aliases by actual form id and visible academic term.
    sh = row(49)["host_fields"]
    years = {8: "Spring2027", 9: "資料夾僅Spring2026檔案列表，未讀其中各PDF正文", 10: "Fall2026", 11: "Spring2027", 12: "Spring2026 closed form"}
    form_evs = []
    for number, year in years.items():
        label = f"seventh-discovered-target-{number}"
        quote = "Spring 2027 Exchange Student Nomination Form" if number in {8, 11} else "Sungshin-Spring 2026" if number == 9 else "Fall 2026 Exchange Student Nomination Form" if number == 10 else "Spring 2026 Exchange Student Nomination Form"
        form_evs.append(html(49, label, quote, year))
    r8, r11 = [receipt(49, f"seventh-discovered-target-{n}") for n in [8, 11]]
    assert r8["resolved_url"].split("?")[0] == r11["resolved_url"].split("?")[0]
    text, _ = visible(r8)
    assert "Page 1 of 2" in text and "National Chung Cheng University" in text
    sh["nomination_deadline"]["evidence"] = [e for e in sh["nomination_deadline"]["evidence"] if e.get("source_id") not in {r8["source_id"], r11["source_id"]}] + form_evs[:1] + form_evs[3:4]
    sh["application_forms"]["value"] = sh["application_forms"]["value"].split("；PDF及checklist短提名網址不同")[0].split("。本批已實讀5入口")[0] + "。本批已實讀5入口：公告長網址與PDFp3 Qfs短鏈解析為同一Spring2027 Google form ID、日期一致；只讀公開Page1/2、未填Next或送出，Google登入提示是保存進度而非已證強制登入。Page1需overall GPA（範例非門檻）、母校英文名、標為Nationality但選項實為大學名、coordinator及email；含中正選項不代表本期校內開放。PDFp1 smW短鏈實為Fall2026表，checklist zWD短鏈實為已關閉Spring2026表，不能拿舊表提名。公告Drive folder實列Spring2026材料／2025修改日，此為列表核對，不假定每附件正文已讀；當期直接公告PDF仍保留。學生TIS signup/login私有表與公開coordinator提名分開。"
    sh["application_forms"]["evidence"] = [e for e in sh["application_forms"]["evidence"] if e.get("source_id") not in {e["source_id"] for e in form_evs}] + [{**e, "status": "conflict"} for e in form_evs]
    sh["application_forms"]["reason"] = "已解公開提名同表身份與舊年度入口；XLSX舊到件年、舊表／folder链接及個人TIS完整表適用仍待核。"
    docs = load("completion-sungshin-documents.json")
    docs["residual_conflicts"] = [x for x in docs["residual_conflicts"] if "短鏈與長鏈不同" not in x]
    docs["residual_conflicts"] += [x for x in ["PDFp1 website短鏈為Fall2026；checklist短鏈為closed Spring2026", "Spring2027公告Drive folder實列Spring2026材料；未核每檔正文"] if x not in docs["residual_conflicts"]]
    docs["public_nomination"] = {"academic_year": "Spring2027", "official_url": r8["resolved_url"], "same_form_id_source_ids": [r8["source_id"], r11["source_id"]], "read_extent": "Public Page1 of2 and instructions; no answers/Next/submission", "evidence": form_evs[:1] + form_evs[3:4], "student_tis_access_unchanged": True}
    (B / "completion-sungshin-documents.json").write_text(json.dumps(docs, ensure_ascii=False, indent=2) + "\n")

    # Canonical institution identity is supported by a first-party link. Main
    # host availability remains a separate 503 acquisition, not a false 200.
    l = row(96)
    lt, soup = visible(receipt(96, "seventh-official-rafu"))
    assert "Saint Petersburg Electrotechnical University" in lt
    anchor = soup.find("a", string=re.compile("University web site"))
    assert anchor and anchor["href"] == "https://etu.ru/en/university/"
    we = html(96, "seventh-official-rafu", "University web site", "官方LETI RAfU頁校名＋University web site實際anchor；一般身份，非CCU交換年度")
    we["location_in_source"] = "University web site anchor href=https://etu.ru/en/university/; LETI institution introduction"
    wf = fact(anchor["href"], [we], reason="第一方官網子站正文校名及主站anchor已核；主站/en/university/本次503另有receipt，不宣稱可讀。非洲預科課程／免費條件不套CCU。")
    l.setdefault("general_campus_fields", {})["website"] = wf
    l.setdefault("profile_fields", {})["website"] = wf

    (B / "completion-manual-fields.json").write_text(json.dumps(manual, ensure_ascii=False, indent=2) + "\n")

    # Store honest reading scope independently of whether a useful field was
    # found. Aliases retain their acquisition URLs without counting new bodies.
    ledger = load("completion-public-body-reviews.json")
    inputs = load("completion-seventh-reading-input.json")
    records = []
    for item in inputs:
        sid = item["source_id"]
        if sid == "p21-reference-link-1-file-3":
            continue
        disposition, note = PRIORITY_NOTES[sid]
        records.append({"source_id": sid, "ccu_page": item["ccu_page"], "source_url": item["source_url"], "checked_at": item["checked_at"], "content_hash": item["sha256"], "http_status": item["http_status"], "body_read": True, "academic_year": "按閱讀註記限制；查核2026-10-11，不由文件名／copyright推申請年度", "review_note": note, "review_extent": "main visible body; folders/wrappers/indexes only where explicitly stated", "disposition": disposition, "batch": 7})
    extra = {"p63-followup-inbound": "完整incoming正文，僅2024–25資料已按reference接欄", "p63-host_entry-1": "international degree／co-op介紹，不是exchange资格", "p63-host_entry-2": "merit scholarship限面試，不保交換獎金", "p63-seventh-current-student-exchange": "完整Student Exchange介紹及partner名冊；無2027申請政策，不列中正也不推失效", "p96-seventh-official-rafu": "LETI校名與主站anchor身份段已讀；非洲預科不作CCU規則", "p115-seventh-visiting-student-housing": "正文主要行車／停車及地圖，Visiting標題未給接待資格", "p115-seventh-housing-application": "完整2026/27一般住房申請／候補／Family資格，分對象", "p115-seventh-housing-rates": "完整2026/27全年Park與Family費表；分期非月租、非Spring報價", "p115-seventh-housing-calendar": "完整Spring2026／Summer2026與AY26/27 Fall事件，非Spring2027截止", "p84-seventh-discovered-target-13": "完整2026/27 incoming含non-Erasmus、2026-10-08更新；非完整課表", "p84-seventh-discovered-target-14": "希臘語頁只有English指引，實际英語子頁另讀", "p84-seventh-international-english": "International Erasmus 1–2semester／3–12month及有協議津貼，不能套CCU雙邊", "p84-seventh-current-factsheet": "全部2頁PDF實讀、2026/27及日期；PDFp1連結文字仍2025/26，不拿旧label覆新頁", "p84-seventh-current-calendar": "全部1頁PDF實讀，Spring及LanguageCentre考試分開", "p84-seventh-current-language": "全部2頁PDF實讀、B2及母校證明矛盾、Erasmus跨School限制", "p84-seventh-current-bilateral-form": "完整DOCX所有段落與表格，學年留白及各簽名角色", "p84-seventh-current-insurance": "全部2頁PDF實讀，EU/EHIC及第三國，列公司非推薦", "p84-seventh-current-entry": "全部8頁PDF實讀，四階段／原件認證；未標版本與CCU個案法律適用未核"}
    extra["p84-seventh-discovered-target-15"] = "完整2026/27正文與target13 normalized visible body相同、原始HTTP SHA不同；保留兩URL，不宣稱samebytes或兩份不同政策"
    extra["p71-seventh-official-360-bus"] = "HTML容器僅Javascript not activated；有公開tour.js入口但無機構marker及重用圖授權，不稱已讀tour內容"
    extra.update({f"p49-seventh-discovered-target-{n}": ("公開Page1/2說明與問題讀完、未填Next；同Spring2027 form ID" if n in {8, 11} else "完整公開Fall2026 Page1/2，不讀未經Next的後頁" if n == 10 else "完整Drive Spring2026列表，沒有閱讀所列每PDF正文" if n == 9 else "完整closed Spring2026公告，非2027可用表") for n in [8, 9, 10, 11, 12]})
    for sid, note in extra.items():
        record = receipt(int(sid.split("-")[0][1:]), sid.split("-", 1)[1])
        records.append({"source_id": sid, "ccu_page": record["ccu_page"], "source_url": record["source_url"], "checked_at": record["checked_at"], "content_hash": record["sha256"], "http_status": 200, "body_read": True, "academic_year": note, "review_note": note, "review_extent": note, "batch": 7})
    ids = {r["source_id"] for r in records}
    ledger = [r for r in ledger if r.get("source_id") not in ids] + records
    (B / "completion-public-body-reviews.json").write_text(json.dumps(ledger, ensure_ascii=False, indent=2) + "\n")
    print(f"Seventh reading: {len(records)} scoped records; 37 priority URLs / 36 distinct bodies; six schools' fields mapped")


if __name__ == "__main__":
    main()
