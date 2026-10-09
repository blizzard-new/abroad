"""Capture TKU official evidence to a staging directory; never overwrite reviewed data.

Run: .venv/bin/python scripts/refresh_tku_sources.py --output /tmp/tku-evidence
Use --manifest to re-fetch the exact reviewed source IDs, including host sources.
Review content, applicability and conflicts before promoting any new source/catalogue.
No sister-school list is converted into a current application list by this script.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urljoin, urlsplit

import requests
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

BASE = "https://oieie.tku.edu.tw"


def xlsx_sheets(blob: bytes) -> list[dict]:
    """Read OOXML values, preserving sheet/cell locators and merged-country blanks."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        strings = ["".join(el.itertext()) for el in ET.fromstring(archive.read("xl/sharedStrings.xml")).findall("m:si", ns)]
        sheet_names = [el.get("name") for el in ET.fromstring(archive.read("xl/workbook.xml")).findall("m:sheets/m:sheet", ns)]
        result = []
        for index, sheet_name in enumerate(sheet_names, 1):
            root = ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            rows = []
            for row in root.findall("m:sheetData/m:row", ns):
                cells = {}
                for cell in row.findall("m:c", ns):
                    value = cell.find("m:v", ns)
                    text = value.text if value is not None else None
                    if text is not None:
                        cells[cell.get("r")] = strings[int(text)] if cell.get("t") == "s" else text
                rows.append(cells)
            result.append({"name": sheet_name, "rows": rows,
                           "merged_cells": [cell.get("ref") for cell in root.findall("m:mergeCells/m:mergeCell", ns)]})
        return result


def read_source(url: str, source_id: str, output: Path, checked_at: str) -> dict:
    response = requests.get(url, timeout=40)
    response.raise_for_status()
    blob = response.content
    result = {
        "id": source_id, "source_url": url, "final_url": response.url,
        "checked_at": checked_at, "http_status": response.status_code,
        "content_type": response.headers.get("Content-Type", ""),
        "sha256": hashlib.sha256(blob).hexdigest(), "bytes": len(blob),
    }
    if blob.startswith(b"%PDF"):
        reader = PdfReader(io.BytesIO(blob))
        result["pages"] = [{"page": i + 1, "text": page.extract_text() or ""}
                           for i, page in enumerate(reader.pages)]
        (output / f"{source_id}.pdf").write_bytes(blob)
        result["local_file"] = f"{source_id}.pdf"
    elif blob.startswith(b"PK") and urlsplit(url).path.endswith(".xlsx"):
        result["sheets"] = xlsx_sheets(blob)
        (output / f"{source_id}.xlsx").write_bytes(blob)
        result["local_file"] = f"{source_id}.xlsx"
    elif blob.startswith(b"PK") and urlsplit(url).path.endswith(".docx"):
        doc = Document(io.BytesIO(blob))
        result["tables"] = [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
        result["text"] = "\n".join([*(p.text for p in doc.paragraphs),
                                    *(" | ".join(c.text for c in row.cells)
                                      for table in doc.tables for row in table.rows)])
        (output / f"{source_id}.docx").write_bytes(blob)
        result["local_file"] = f"{source_id}.docx"
    else:
        soup = BeautifulSoup(blob, "html.parser")
        for el in soup.select("script, style, nav, header, footer"):
            el.decompose()
        result["text"] = soup.get_text("\n", strip=True)
        result["links"] = [{"title": a.get_text(" ", strip=True),
                            "url": urljoin(response.url, a["href"])}
                           for a in soup.select("a[href]")]
    (output / f"{source_id}.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    return result


def collect(output: Path, checked_at: str) -> list[dict]:
    output.mkdir(parents=True, exist_ok=False)
    index = read_source(BASE + "/oieie/?page_id=55", "news-index", output, checked_at)
    home = read_source(BASE, "home-index", output, checked_at)
    records = [index, home]
    targets = (
        ("jp-revision", lambda t: "更正2027-2028" in t and "第二次" in t),
        ("jp-original", lambda t: "公告2027-2028" in t and "第二次" in t),
        ("english-115", lambda t: "公告2026-2027年英歐（法" in t and "作業簡章" in t),
        ("mainland-115", lambda t: "公告檢送本校115學年度" in t and "作業簡章" in t),
        ("english-second-115", lambda t: "2026-2027年英歐語組" in t and "第2次甄選作業" in t),
    )
    for source_id, matches in targets:
        candidates = {link["url"] for page in [index, home] for link in page["links"] if matches(link["title"])}
        if len(candidates) != 1:
            raise ValueError(f"{source_id}: expected one official notice, got {candidates}")
        notice = read_source(next(iter(candidates)), source_id, output, checked_at)
        records.append(notice)
        attachments = {link["url"]: link["title"] for link in notice["links"]
                       if urlsplit(link["url"]).path.lower().endswith((".pdf", ".docx", ".xlsx"))}
        for i, (url, title) in enumerate(attachments.items(), 1):
            if "錄取" in title:  # Result lists contain student names, unnecessary for catalogue work.
                continue
            attachment_id = f"{source_id}-attachment-{i}"
            if source_id == "english-115" and urlsplit(url).path.lower().endswith(".xlsx"):
                attachment_id = "english-115-partners"
            try:
                attachment = read_source(url, attachment_id, output, checked_at)
                attachment["open_status"] = "verified_open"
            except (requests.RequestException, ValueError) as exc:
                attachment = {"id": attachment_id, "source_url": url,
                              "checked_at": checked_at, "open_status": "failed",
                              "error": str(exc)}
            attachment.update(title=title or unquote(urlsplit(url).path.rsplit("/", 1)[-1]), parent_id=source_id)
            records.append(attachment)
    for source_id, url in [
        ("outbound", BASE + "/oieie/?page_id=19024"),
        ("sister-directory", "https://sister.tku.edu.tw/sitemap.asp"),
    ]:
        records.append(read_source(url, source_id, output, checked_at))
    (output / "sources.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
    return records


def collect_manifest(manifest: Path, output: Path, checked_at: str) -> list[dict]:
    """Re-fetch evidence without treating old applicability metadata as newly reviewed."""
    output.mkdir(parents=True, exist_ok=False)
    records = []
    for previous in json.loads(manifest.read_text(encoding="utf-8")):
        try:
            record = read_source(previous["source_url"], previous["id"], output, checked_at)
            record["open_status"] = "verified_open"
        except (requests.RequestException, ValueError) as exc:
            record = {"id": previous["id"], "source_url": previous["source_url"],
                      "checked_at": checked_at, "open_status": "failed", "error": str(exc)}
        for key in ("title", "parent_id", "role"):
            if key in previous:
                record[key] = previous[key]
        record["previous_applicable_year"] = previous.get("applicable_year")
        record["applicability_review_status"] = "requires_content_review"
        records.append(record)
    (output / "sources.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
    return records


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--checked-at", default=datetime.now(timezone.utc).isoformat())
    parser.add_argument("--manifest", type=Path, help="Re-fetch registered sources; applicability still needs review")
    args = parser.parse_args()
    records = (collect_manifest(args.manifest, args.output, args.checked_at) if args.manifest
               else collect(args.output, args.checked_at))
    print(json.dumps({"source_count": len(records), "output": str(args.output)}))
