from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .ai_gateway import structured_response
from .models import DocumentCheckIssue, DocumentCheckResult


MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def check_document(
    *,
    file_name: str,
    content: bytes,
    document_type: str = "auto",
    expected_word_limit: int | None = None,
    require_signature: bool = False,
    use_ai: bool = False,
) -> DocumentCheckResult:
    if not content:
        raise ValueError("檔案是空的，請重新選擇。")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("檔案不可超過 10 MB。")
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("目前支援 PDF、DOCX 與 TXT。")

    text, page_count = _extract_text(extension, content)
    compact = _compact(text)
    resolved_type = _infer_document_type(file_name, compact) if document_type == "auto" else document_type
    if resolved_type not in {
        "study_plan",
        "recommendation",
        "resume",
        "application_form",
        "other",
    }:
        raise ValueError("不支援的文件類型。")

    word_count = _word_count(compact)
    issues: list[DocumentCheckIssue] = []
    passed: list[str] = []
    suggestions: list[str] = []

    if not compact:
        issues.append(
            _issue(
                "error",
                "no_text",
                "無法讀取文字",
                "這可能是掃描版 PDF；請改傳可選取文字的 PDF、DOCX 或 TXT。",
            )
        )
    elif len(compact) < 120:
        issues.append(
            _issue(
                "warning",
                "too_short",
                "內容可能不完整",
                "目前可讀內容太短，請確認沒有漏頁或只上傳到封面。",
            )
        )
    else:
        passed.append("文件文字可正常讀取")

    if expected_word_limit:
        if word_count > expected_word_limit:
            issues.append(
                _issue(
                    "error",
                    "word_limit",
                    "超過設定字數",
                    f"目前約 {word_count} 字，設定上限為 {expected_word_limit} 字。",
                )
            )
        else:
            passed.append(f"未超過 {expected_word_limit} 字上限")

    if re.search(r"(?:_{3,}|請填寫|尚未填|\[填寫|<填寫|待補)", compact, re.IGNORECASE):
        issues.append(
            _issue(
                "warning",
                "blank_fields",
                "仍有疑似空白欄位",
                "找到底線或待填提示，下載列印前請逐欄確認。",
            )
        )
    elif compact:
        passed.append("未發現明顯待填標記")

    _run_type_checks(resolved_type, compact, issues, passed, suggestions)

    signature_marker = _has_present_term(
        compact,
        r"簽名|簽章|親簽|signature|簽署人",
    )
    if require_signature:
        if signature_marker:
            passed.append("已辨識到簽名或簽章位置")
            suggestions.append("簽名欄請保留空白，列印後由本人親簽。")
        else:
            issues.append(
                _issue(
                    "warning",
                    "signature_missing",
                    "未辨識到簽名位置",
                    "若官方版本要求親簽或核章，請回到原始表單確認簽名欄。",
                )
            )

    ai_payload = None
    if use_ai and compact:
        ai_payload = structured_response(
            schema_name="application_document_review",
            schema={
                "type": "object",
                "properties": {
                    "summary": {"type": "string"},
                    "suggestions": {
                        "type": "array",
                        "items": {"type": "string"},
                        "maxItems": 5,
                    },
                    "concerns": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "severity": {
                                    "type": "string",
                                    "enum": ["error", "warning", "info"],
                                },
                                "title": {"type": "string"},
                                "detail": {"type": "string"},
                            },
                            "required": ["severity", "title", "detail"],
                            "additionalProperties": False,
                        },
                        "maxItems": 5,
                    },
                },
                "required": ["summary", "suggestions", "concerns"],
                "additionalProperties": False,
            },
            instructions=(
                "你是交換申請文件檢查員。使用繁體中文，只根據文件內容指出缺漏、矛盾、格式風險與可執行修改。"
                "不要假設未提供的校規，不要替學生捏造經歷，也不要輸出文件中的個資。"
            ),
            input_text=(
                f"文件類型：{resolved_type}\n字數上限：{expected_word_limit or '未設定'}\n"
                f"需要簽名：{'是' if require_signature else '否'}\n\n文件內容：\n{compact[:16_000]}"
            ),
            tier="advanced",
        )

    ai_enhanced = isinstance(ai_payload, dict)
    if ai_enhanced:
        for concern in ai_payload.get("concerns", [])[:5]:
            if not isinstance(concern, dict):
                continue
            severity = str(concern.get("severity", "info"))
            if severity not in {"error", "warning", "info"}:
                severity = "info"
            issues.append(
                _issue(
                    severity,
                    "ai_review",
                    _compact(str(concern.get("title", "內容提醒")))[:80],
                    _compact(str(concern.get("detail", "")))[:300],
                )
            )
        suggestions.extend(
            _compact(str(item))[:240]
            for item in ai_payload.get("suggestions", [])[:5]
            if str(item).strip()
        )

    issues = _unique_issues(issues)
    suggestions = list(dict.fromkeys(item for item in suggestions if item))[:8]
    error_count = sum(item.severity == "error" for item in issues)
    warning_count = sum(item.severity == "warning" for item in issues)
    info_count = sum(item.severity == "info" for item in issues)
    completeness_score = max(
        0,
        min(100, 100 - error_count * 22 - warning_count * 10 - info_count * 3),
    )
    if not compact:
        summary = "這份文件目前無法檢查。"
    elif error_count:
        summary = f"找到 {error_count} 項需要先修正、{warning_count} 項需要確認。"
    elif warning_count:
        summary = f"基本格式可讀，另有 {warning_count} 項需要確認。"
    else:
        summary = "未發現明顯缺漏，送件前仍要對照官方最新版。"
    if ai_enhanced and str(ai_payload.get("summary", "")).strip():
        summary = _compact(str(ai_payload["summary"]))[:320]

    privacy_note = (
        "原檔不會保存；本次已依你的選擇，將抽取文字送交 AI 檢查，請勿上傳不必要的敏感資料。"
        if ai_enhanced
        else "檔案只在本次請求的記憶體中檢查，不會保存原檔。"
    )
    return DocumentCheckResult(
        file_name=Path(file_name).name,
        document_type=resolved_type,
        summary=summary,
        page_count=page_count,
        character_count=len(compact),
        word_count=word_count,
        completeness_score=completeness_score,
        checks_passed=passed,
        issues=issues,
        suggestions=suggestions,
        extracted_preview=_preview(compact),
        ai_enhanced=ai_enhanced,
        privacy_note=privacy_note,
    )


def _extract_text(extension: str, content: bytes) -> tuple[str, int]:
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)
        if extension == ".docx":
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(paragraphs), 1
        return content.decode("utf-8-sig"), 1
    except (OSError, ValueError, UnicodeDecodeError) as error:
        raise ValueError("檔案內容無法解析，請確認檔案沒有損壞或加密。") from error


def _run_type_checks(
    document_type: str,
    text: str,
    issues: list[DocumentCheckIssue],
    passed: list[str],
    suggestions: list[str],
) -> None:
    if not text:
        return
    checks: dict[str, list[tuple[str, str, str]]] = {
        "study_plan": [
            (r"動機|申請原因|why|motivation", "申請動機", "補上為什麼選擇這個國家、學校與計畫。"),
            (r"課程|修課|study plan|course", "修課規劃", "列出預計修讀課程與選課理由。"),
            (r"返國|未來|生涯|future|career", "返國後規劃", "說明交換經驗如何銜接返國後的學習或生涯。"),
        ],
        "recommendation": [
            (r"推薦|recommend", "明確推薦語", "加入推薦人對申請者的明確評價與推薦結論。"),
            (r"教授|老師|職稱|系主任|title|position", "推薦人身分", "標示推薦人姓名、單位與職稱。"),
            (r"聯絡|email|電話|contact", "推薦人聯絡方式", "依官方格式補上推薦人聯絡方式。"),
        ],
        "resume": [
            (r"教育|學歷|education", "學歷", "補上學校、系所、學位與就讀期間。"),
            (r"經歷|專案|活動|experience|project", "經歷或專案", "用具體成果描述相關經歷。"),
            (r"技能|語言|skills?|languages?", "技能或語言", "列出語言證明與可驗證技能。"),
            (r"gpa|平均成績|學業成績", "GPA 或學業成績", "若申請規定要求，補上 GPA 或學業平均。"),
        ],
        "application_form": [
            (r"姓名|name", "姓名欄位", "確認中文姓名與護照英文姓名一致。"),
            (r"學號|student id", "學號欄位", "補上正確學號。"),
            (r"系所|department|major", "系所欄位", "補上目前系所或學程。"),
            (r"日期|date", "日期欄位", "確認所有日期格式與先後順序。"),
        ],
    }
    for pattern, label, suggestion in checks.get(document_type, []):
        if _has_present_term(text, pattern):
            passed.append(f"已包含{label}")
        else:
            issues.append(
                _issue("warning", f"missing_{label}", f"未找到{label}", suggestion)
            )
    if document_type == "study_plan" and len(text) < 500:
        suggestions.append("讀書計畫偏短，可補充選校原因、課程安排與返國後目標。")


def _infer_document_type(file_name: str, text: str) -> str:
    haystack = f"{file_name} {text[:1000]}".casefold()
    patterns = [
        ("study_plan", r"讀書計畫|study plan|statement of purpose|學習計畫"),
        ("recommendation", r"推薦信|recommendation|reference letter"),
        ("resume", r"履歷|簡歷|curriculum vitae|\bcv\b|resume"),
        ("application_form", r"申請表|application form|基本資料"),
    ]
    for document_type, pattern in patterns:
        if re.search(pattern, haystack, re.IGNORECASE):
            return document_type
    return "other"


def _has_present_term(text: str, pattern: str) -> bool:
    for match in re.finditer(pattern, text, re.IGNORECASE):
        context = text[max(0, match.start() - 24) : match.start()]
        if re.search(r"尚未(?:加入|填寫|提供)?|未(?:加入|填寫|提供|包含)|未\s*$|缺少|沒有|待補", context):
            continue
        return True
    return False


def _word_count(text: str) -> int:
    latin_words = re.findall(r"[A-Za-z0-9]+(?:['’-][A-Za-z0-9]+)*", text)
    chinese_characters = re.findall(r"[\u4e00-\u9fff]", text)
    return len(latin_words) + len(chinese_characters)


def _preview(text: str) -> str:
    if not text:
        return ""
    return text if len(text) <= 360 else f"{text[:359].rstrip()}…"


def _compact(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _issue(severity: str, code: str, title: str, detail: str) -> DocumentCheckIssue:
    return DocumentCheckIssue(
        severity=severity,
        code=re.sub(r"[^a-z0-9_]+", "_", code.casefold()).strip("_") or "review",
        title=title,
        detail=detail,
    )


def _unique_issues(issues: list[DocumentCheckIssue]) -> list[DocumentCheckIssue]:
    seen: set[tuple[str, str]] = set()
    output: list[DocumentCheckIssue] = []
    for issue in issues:
        key = (issue.title, issue.detail)
        if key in seen:
            continue
        seen.add(key)
        output.append(issue)
    return output
