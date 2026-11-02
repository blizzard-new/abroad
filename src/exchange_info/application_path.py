"""Account-owned application paths over existing rules and encrypted CAS storage.

The registry is an internal encrypted-draft namespace, not a school form. Reusing
the existing owner-bound store keeps self-host installations migration-free.
"""
from __future__ import annotations

import hashlib
import copy
import io
import json
import re
from typing import Literal

from fastapi import HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field

from .form_drafts import FormDraftConflictError, FormDraftError
from .document_vault import VaultError
from .models import (ApplicationFormDraftSave, PersonalApplicationPlanRequest,
                     ProductAccessTier, ProductFeature, ProductReleaseMode, utc_now)
from .school_identity import school_identity_key
from .plan_defaults import phase_rank

REGISTRY_ID = "uniway-application-paths-v1"
WRITING = [
    ("motivation", "寫一段申請動機", "你希望交換幫助你解決哪一個學習問題？"),
    ("school-reason", "寫下選校原因", "這間學校的哪一項課程或資源適合你？"),
    ("learning-goals", "設定研修目標", "交換結束時，你希望做得到什麼？"),
    ("course-plan", "記下一門想修的課", "記下實際查到的課名，以及想學到什麼。"),
    ("expected-outcomes", "描繪預期成果", "你希望帶回什麼作品或學習成果？"),
    ("return-plan", "連結返國後的學習", "回到母校後，如何運用這段經驗？"),
]


class PathContext(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    home_school: str = Field(min_length=2, max_length=80)
    target_school: str = Field(min_length=2, max_length=200)
    academic_year: str = Field(min_length=1, max_length=20)
    application_round: str = Field(default="", max_length=80)
    departure_year: str = Field(min_length=1, max_length=20)
    semester: Literal["fall", "spring", "full_year", "undecided"]
    college: str = Field(default="", max_length=120)
    department: str = Field(default="", max_length=120)
    degree_level: str = Field(default="", max_length=80)


class PathCreate(PathContext):
    revision: int = Field(ge=0)
    import_existing: bool = False


class PathStep(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    value: str = Field(default="", max_length=20000)
    confirmed: bool = False
    version_id: str | None = Field(default=None, min_length=8, max_length=80)
    applicable: bool | None = None


class PathUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    revision: int = Field(ge=1)
    step: PathStep


class PathArchive(BaseModel):
    model_config = ConfigDict(extra="forbid")
    revision: int = Field(ge=1)
    archived: bool


def identity(context: dict) -> str:
    return hashlib.sha256(json.dumps(context, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:32]


def load_registry(service, account_id, home):
    draft = service.get_internal_draft(account_id, home, REGISTRY_ID)
    if not draft:
        return {"revision": 0, "plans": []}
    try:
        plans = json.loads(draft.answers["plans"])
        if not isinstance(plans, list):
            raise ValueError()
    except (KeyError, ValueError, TypeError) as exc:
        raise FormDraftError("申請計畫內容無法讀取，請稍後重試。") from exc
    return {"revision": draft.revision, "plans": plans}


def save_registry(service, account_id, home, registry):
    content = json.dumps(registry["plans"], ensure_ascii=False, separators=(",", ":"))
    if len(content.encode()) > 2_000_000:
        raise HTTPException(422, "計畫內容超過保存上限，請縮短筆記。")
    saved = service.save_draft(account_id, REGISTRY_ID, ApplicationFormDraftSave(
        home_school=home, template_version="1", revision=registry["revision"], answers={"plans": content}))
    registry["revision"] = saved.revision


def level(level_id, title, kind, prompt, stage, **extra):
    return dict(id=level_id, title=title, kind=kind, prompt=prompt, stage_key=stage,
                icon={"confirm": "badge-check", "write": "notebook-pen", "file": "file-up",
                      "notice": "mail", "submission": "send", "assemble": "notebook-tabs"}.get(kind, "clipboard-list"),
                external=kind == "notice", source_urls=[], dependency_ids=[], **extra)


def build_levels(context, evaluated, catalog):
    """Use scoped service output; never invent institution requirements or dates."""
    choice = evaluated.choices[0]
    levels = [
        level("home", "確認你的母校", "confirm", context["home_school"], "找到你的方向"),
        level("term", "確認出國學期", "confirm", context["departure_year"] + " · " + {"fall":"上學期", "spring":"下學期", "full_year":"全學年", "undecided":"尚未決定"}[context["semester"]], "找到你的方向"),
        level("target", "確認交換目標", "confirm", context["target_school"], "找到你的方向"),
    ]
    for check in choice.checks:
        item = level("check-" + check.code, check.title, "confirm", check.detail, "釐清申請條件")
        item.update(source_urls=[check.source_url] if check.source_url else [],
                    check_status=check.status,
                    warning={"pass":"規則比對通過，仍以校方審核為準。", "fail":"目前條件不符合此項規則；個人確認不會改變資格判斷。", "needs_confirmation":"資料或個人條件待確認。"}[check.status])
        levels.append(item)
    documents = [d for d in catalog.items if not d.applicability or d.applicability.status not in {"not_applicable", "deferred"}]
    study_docs = [d for d in documents if re.search(r"研修計畫|讀書計畫|學習計畫|study.plan", d.name, re.I)]
    if study_docs:
        for i, (key, title, prompt) in enumerate(WRITING):
            item = level(key, title, "write", prompt, "寫出交換的理由" if i < 3 else "規劃交換的收穫")
            item.update(source_urls=study_docs[0].source_urls,
                        warning="分段寫作是準備輔助，正式格式依當期規定。")
            levels.append(item)
        item = level("assemble", "整理研修計畫", "assemble", "把已寫的小段落集中檢查，並下載草稿。", "整理研修計畫")
        item["dependency_ids"] = [x[0] for x in WRITING]
        levels.append(item)
    preparation_count = len(levels)
    for doc in documents:
        item = level("document-" + identity({"id":doc.id, "name":doc.name}), doc.name, "file",
                     "選擇私人文件庫中的檔案，或記下你自行保管的文件位置。", doc.phase or "備齊申請文件")
        item.update(source_urls=doc.source_urls,
                    action_url=doc.official_form_url if doc.download_available else None,
                    warning=doc.condition or ("此文件適用性待確認。" if doc.needs_confirmation or doc.conditional else ""),
                    required=not doc.conditional, conditional=doc.conditional, document_id=doc.id)
        levels.append(item)
    for task in choice.checklist:
        is_notice = bool(re.search(r"審查|甄選結果|提名通知|錄取結果|錄取通知", task.title))
        kind = "notice" if is_notice else "submission" if re.search(r"送件|提交|繳交|遞交", task.title) else "note"
        item = level(task.id, task.title, kind, task.detail, task.phase or "申請事項")
        item.update(source_urls=task.source_urls, dependency_ids=task.dependency_ids,
                    official_deadline_at=task.official_deadline_at.isoformat() if task.official_deadline_at else None,
                    deadline_confirmed=bool(task.official_deadline_at and task.task_origin == "official" and not task.needs_confirmation),
                    deadline_source_url=task.deadline_source_url,
                    source_signature=task.source_signature,
                    templates=task.document_link.template_ids if task.document_link else [],
                    warning="參考其他年度，需核對當期規定。" if task.task_origin == "reference_year" else "此項規定仍需確認。" if task.needs_confirmation else "")
        levels.append(item)
    # Keep the existing source phases together: departure files must not precede
    # the school's submission task just because documents came from another API.
    phase_order = list(dict.fromkeys([t.phase for t in choice.checklist] + [d.phase for d in documents]))
    known_order = sorted(phase_order, key=lambda p: phase_rank(p) if phase_rank(p) is not None else 99)
    levels[preparation_count:] = sorted(levels[preparation_count:], key=lambda l: known_order.index(l["stage_key"]) if l["stage_key"] in known_order else len(known_order))
    if not choice.checklist:
        levels.append(level("missing-rules", "確認當期申請流程", "note",
                            "目前沒有適用的完整流程。請向母校確認當期要求，這份地圖不代表完整申請清單。", "確認缺少的資訊",
                            warning="流程資料不足；完成個人準備不代表申請資料已齊全。"))
    # Limit each chapter to three real actions, without manufacturing dependencies.
    stages = []
    for item in levels:
        if not stages or stages[-1]["key"] != item["stage_key"] or len(stages[-1]["levels"]) >= 3:
            stages.append({"key":item["stage_key"], "title":item["stage_key"], "icon":item["icon"], "index":len(stages), "levels":[]})
        item["stage"] = stages[-1]["index"]
        item["number"] = levels.index(item) + 1
        stages[-1]["levels"].append(item)
    return stages


def progress(item, record):
    if item["kind"] == "confirm":
        return {"value":int(record.get("confirmed", False)), "total":1}
    content = bool(record.get("value") or record.get("version_id"))
    if item["external"]:
        return {"value":int(content), "total":1}
    return {"value":int(content) + int(content and record.get("confirmed", False)), "total":2}


def register_application_path(app, *, authenticate, university_for, evaluate, catalog_for, release_mode, feature_enabled):
    def service():
        if not getattr(app.state, "form_draft_service", None):
            raise HTTPException(503, "申請進度保存服務未啟用。")
        return app.state.form_draft_service

    def home_name(home):
        return university_for(home).school_name

    def read(account, home):
        try:
            return load_registry(service(), account.id, home)
        except FormDraftError as exc:
            raise HTTPException(503, "申請進度暫時無法讀取，請稍後重試。") from exc

    def commit(account, home, registry):
        try:
            save_registry(service(), account.id, home, registry)
        except FormDraftConflictError as exc:
            raise HTTPException(409, {"code":"application_path_conflict", "message":"另一個視窗已更新進度。你的輸入仍在，請重新載入後確認。"}) from exc
        except FormDraftError as exc:
            raise HTTPException(503, "這次未能確認保存成功，請重新載入進度後再試。") from exc

    def locate(registry, plan_id):
        plan = next((p for p in registry["plans"] if p["id"] == plan_id), None)
        if not plan:
            raise HTTPException(404, "找不到這份申請計畫。")
        return plan

    def check_revision(registry, revision):
        if registry["revision"] != revision:
            raise HTTPException(409, {"code":"application_path_conflict", "message":"進度已在另一個視窗更新，請重新載入後確認。"})

    def workspace(account, home):
        return next((w for w in account.application_workspaces if school_identity_key(w.home_school) == school_identity_key(home)), None)

    def catalog(plan):
        c = plan["context"]
        result = evaluate(PersonalApplicationPlanRequest(home_school=c["home_school"], academic_year=c["academic_year"],
            application_round=c["application_round"] or None, college=c["college"] or None,
            department=c["department"] or None, degree_level=c["degree_level"] or None,
            preferences=[{"rank":1, "school_name":c["target_school"]}]))
        docs = catalog_for(school=c["home_school"], target_school=c["target_school"], college=c["college"] or None,
            department=c["department"] or None, degree_level=c["degree_level"] or None, program_type="exchange",
            academic_year=c["academic_year"], application_round=c["application_round"] or None,
            host_admission_status="unknown", housing_choice="unknown", semester=c["semester"])
        return build_levels(c, result, docs)

    def summary(plan):
        return {k:plan[k] for k in ("id", "context", "archived", "updated_at")}

    def envelope(registry, plan, account):
        plan = copy.deepcopy(plan)
        stages = catalog(plan)
        items = {l["id"]:l for s in stages for l in s["levels"]}
        linked = {key for key, saved in plan["steps"].items() if saved.get("version_id")}
        if linked:
            vault = getattr(app.state, "document_vault_service", None)
            if not vault:
                raise HTTPException(503, "私人文件庫暫時無法確認，請稍後重新載入。")
            try:
                documents = vault.list_documents(account.id)
            except VaultError as exc:
                raise HTTPException(503, "私人文件狀態暫時無法確認，請稍後重試。") from exc
            valid = {v.id for d in documents.items for v in d.versions if v.expires_at > utc_now()}
            for key in linked:
                if key in items and plan["steps"][key]["version_id"] not in valid:
                    items[key]["file_unavailable"] = True
                    items[key]["warning"] = "連結的文件已過期或移除。請重新上傳，或改記錄自行保管的位置。"
        resolved = {}
        def effective(key, seen=frozenset()):
            if key in resolved:
                return resolved[key]
            if key not in items or key in seen:
                return {"value":0, "total":1}
            item = items[key]
            saved = plan["steps"].get(key, {})
            if saved and item.get("source_signature") and saved.get("source_signature") != item["source_signature"]:
                saved = {**saved, "confirmed":False}
                item["warning"] = "適用規則已更新，請重新核對；你的紀錄仍保留。"
            if item.get("file_unavailable"):
                saved = {**saved, "version_id":None, "confirmed":False}
            p = progress(item, saved)
            blocked = [dep for dep in item["dependency_ids"]
                       if (d := effective(dep, seen | {key}))["value"] < d["total"]]
            item["blocked_by"] = blocked
            if blocked and p["value"] == p["total"]:
                p["value"] = max(0, p["total"] - 1)
            resolved[key] = p
            return p
        done = total = 0
        for stage in stages:
            for item in stage["levels"]:
                item["progress"] = effective(item["id"])
                item["counted"] = not item["external"] and (not item.get("conditional") or plan["steps"].get(item["id"], {}).get("applicable") is True)
                if item["counted"]:
                    done += item["progress"]["value"]
                    total += item["progress"]["total"]
        return {"revision":registry["revision"], "plan":plan, "stages":stages,
                "progress":{"value":done, "total":total}, "recorded_by":"student", "platform_submitted":False}

    def quota(account, home, registry, target):
        if release_mode() == ProductReleaseMode.PUBLIC and account.product_tier != ProductAccessTier.SUBSCRIPTION:
            old = workspace(account, home)
            targets = set(old.application_schools if old else [])
            targets.update(p["context"]["target_school"] for p in registry["plans"] if not p["archived"])
            targets.add(target)
            if len(targets) > 1:
                raise HTTPException(403, "免費方案同時只能準備一所交換學校，請先調整既有計畫或確認方案權限。")

    @app.get("/account/application-paths")
    def list_paths(request: Request, home_school: str = Query(min_length=2, max_length=80)):
        account, _ = authenticate(request)
        home = home_name(home_school)
        registry = read(account, home)
        old = workspace(account, home)
        return {"revision":registry["revision"], "items":[summary(p) for p in registry["plans"]],
                "legacy": {"context":old.application_context.model_dump(mode="json"),
                           "schools":old.application_schools, "preparation_count":len(old.preparation_tasks)} if old else None}

    @app.post("/account/application-paths", status_code=201)
    def create_path(payload: PathCreate, request: Request):
        account, _ = authenticate(request)
        context = PathContext.model_validate(payload.model_dump(exclude={"revision", "import_existing"})).model_dump()
        context["home_school"] = home = home_name(context["home_school"])
        registry = read(account, home)
        plan_id = identity(context)
        existing = next((p for p in registry["plans"] if p["id"] == plan_id), None)
        if existing:
            return envelope(registry, existing, account)
        check_revision(registry, payload.revision)
        quota(account, home, registry, context["target_school"])
        if len(registry["plans"]) >= 30:
            raise HTTPException(422, "目前最多可保存 30 份計畫。")
        plan = {"id":plan_id, "context":context, "steps":{}, "archived":False, "updated_at":utc_now().isoformat()}
        stages = catalog(plan)  # Resolves target and rules before any persistence.
        if payload.import_existing:
            old = workspace(account, home)
            c = old.application_context if old else None
            if not c or any(str(getattr(c, key, None) or "") != context[key] for key in
                            ("target_school", "academic_year", "application_round", "semester", "college", "department", "degree_level")):
                raise HTTPException(422, "舊紀錄的學校、年度、梯次、學期或院系無法對應。請建立空白計畫，舊紀錄會完整保留。")
            for stage in stages:
                for item in stage["levels"]:
                    saved = old.preparation_tasks.get(item["id"])
                    if saved and not item["external"] and saved.status in {"ready", "preparing"}:
                        plan["steps"][item["id"]] = {"value":"從既有準備紀錄帶入", "confirmed":saved.status == "ready", "imported_from":"workspace", "imported_revision":old.revision, "source_signature":item.get("source_signature")}
            plan["legacy_import"] = {"workspace_revision":old.revision, "at":utc_now().isoformat()}
        registry["plans"].append(plan)
        commit(account, home, registry)
        return envelope(registry, plan, account)

    @app.get("/account/application-paths/{plan_id}")
    def get_path(plan_id: str, request: Request, home_school: str = Query(min_length=2, max_length=80)):
        account, _ = authenticate(request)
        registry = read(account, home_name(home_school))
        return envelope(registry, locate(registry, plan_id), account)

    @app.get("/account/application-paths/{plan_id}/writing.docx")
    def export_writing(plan_id: str, request: Request, home_school: str = Query(min_length=2, max_length=80)):
        """Export this plan's saved writing, without overwriting legacy form drafts."""
        from docx import Document
        account, _ = authenticate(request)
        registry = read(account, home_name(home_school))
        plan = locate(registry, plan_id)
        writing = [item for stage in catalog(plan) for item in stage["levels"] if item["kind"] == "write"]
        if not writing:
            raise HTTPException(404, "這份計畫沒有分段研修草稿。")
        doc = Document()
        doc.add_heading("研修計畫草稿", 0)
        c = plan["context"]
        doc.add_paragraph(c["target_school"] + " · " + c["departure_year"] + " · " + c["semester"])
        doc.add_paragraph("個人準備草稿；請依校方當期格式、語言與篇幅規定調整後提交。")
        for item in writing:
            doc.add_heading(item["title"], 1)
            doc.add_paragraph(plan["steps"].get(item["id"], {}).get("value", ""))
        stream = io.BytesIO()
        doc.save(stream)
        return Response(stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition":'attachment; filename="UniWay-study-plan-draft.docx"'})

    @app.put("/account/application-paths/{plan_id}/steps/{step_id}")
    def update_step(plan_id: str, step_id: str, payload: PathUpdate, request: Request,
                    home_school: str = Query(min_length=2, max_length=80)):
        account, _ = authenticate(request)
        home = home_name(home_school)
        registry = read(account, home)
        plan = locate(registry, plan_id)
        check_revision(registry, payload.revision)
        if plan["archived"]:
            raise HTTPException(409, "請先恢復這份封存計畫。")
        levels = [l for s in envelope(registry, plan, account)["stages"] for l in s["levels"]]
        item = next((l for l in levels if l["id"] == step_id), None)
        if not item:
            raise HTTPException(422, "這個關卡已不適用，請重新載入計畫。既有紀錄仍保留。")
        record = payload.step.model_dump()
        if record["applicable"] is not None and not item.get("conditional"):
            raise HTTPException(422, "只有條件式文件可設定適用性。")
        if record["version_id"]:
            if item["kind"] != "file":
                raise HTTPException(422, "此關卡不接受文件連結。")
            if not feature_enabled(account, ProductFeature.ENCRYPTED_DOCUMENT_VAULT):
                raise HTTPException(403, "請先確認私人文件庫方案權限。")
            documents = app.state.document_vault_service.list_documents(account.id)
            version = next((v for d in documents.items for v in d.versions if v.id == record["version_id"] and v.expires_at > utc_now()), None)
            if not version:
                raise HTTPException(404, "找不到有效且屬於你的文件版本。")
            record["file_name"] = version.file_name
        if item["external"] and record["confirmed"]:
            raise HTTPException(422, "通知只能自行記錄，不能標示為校方確認。")
        if record["confirmed"]:
            if item.get("conditional") and record["applicable"] is not True:
                raise HTTPException(422, "請先確認這份文件適用於你，再標記完成。")
            by_id = {l["id"]:l for l in levels}
            missing = [key for key in item["dependency_ids"] if key not in by_id or
                       by_id[key]["progress"]["value"] < by_id[key]["progress"]["total"]]
            if missing:
                raise HTTPException(422, "尚有前置事項未完成，請先查看關卡中的前置需求。")
            if item["kind"] == "assemble":
                record["value"] = "已核對組合草稿"
            if item["kind"] != "confirm" and not record["value"] and not record["version_id"]:
                raise HTTPException(422, "請先留下內容或選擇文件，再確認完成。")
        record.update(updated_at=utc_now().isoformat(), recorded_by="student", source_signature=item.get("source_signature"))
        plan["steps"][step_id] = record
        plan["updated_at"] = record["updated_at"]
        commit(account, home, registry)
        return envelope(registry, plan, account)

    @app.put("/account/application-paths/{plan_id}/archive")
    def archive_path(plan_id: str, payload: PathArchive, request: Request,
                     home_school: str = Query(min_length=2, max_length=80)):
        account, _ = authenticate(request)
        home = home_name(home_school)
        registry = read(account, home)
        plan = locate(registry, plan_id)
        check_revision(registry, payload.revision)
        if not payload.archived:
            quota(account, home, registry, plan["context"]["target_school"])
        plan["archived"] = payload.archived
        plan["updated_at"] = utc_now().isoformat()
        commit(account, home, registry)
        return envelope(registry, plan, account)
